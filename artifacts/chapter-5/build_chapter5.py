"""Build Chapter 5 with project-grounded conclusions and a proposed payment roadmap."""
from __future__ import annotations

from collections import Counter
from pathlib import Path
from zipfile import ZipFile
import importlib.util
import json
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
OUTPUT = OUT / 'Chuong_5_Ket_luan_va_huong_phat_trien_MugenBunko.docx'
spec = importlib.util.spec_from_file_location('chapter_format', ROOT / 'artifacts/chapter-4/build_chapter4.py')
fmt = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fmt)
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt

REFERENCES = [
    ('OWASP', 'Third Party Payment Gateway Integration Cheat Sheet', 'https://cheatsheetseries.owasp.org/cheatsheets/Third_Party_Payment_Gateway_Integration_Cheat_Sheet.html'),
    ('OWASP', 'Webhook Security Cheat Sheet', 'https://cheatsheetseries.owasp.org/cheatsheets/Webhook_Security_Cheat_Sheet.html'),
    ('OWASP', 'Transaction Authorization Cheat Sheet', 'https://cheatsheetseries.owasp.org/cheatsheets/Transaction_Authorization_Cheat_Sheet.html'),
    ('SePay Developer', 'Secure your webhook endpoint', 'https://developer.sepay.vn/en/sepay-webhooks/bao-mat'),
    ('MySQL', 'MySQL 8.4 Reference Manual: Fixed-Point Types (Exact Value) – DECIMAL, NUMERIC', 'https://dev.mysql.com/doc/refman/8.4/en/fixed-point-types.html'),
    ('OWASP', 'Business Logic Security Cheat Sheet', 'https://cheatsheetseries.owasp.org/cheatsheets/Business_Logic_Security_Cheat_Sheet.html'),
]
CITED_WORDS = Counter()


def p(doc, text, refs=()):
    for ref in refs:
        CITED_WORDS[ref] += len(text.split())
    suffix = ' ' + ', '.join(f'[{ref}]' for ref in refs) if refs else ''
    return fmt.p(doc, text + suffix)


def h(doc, text, level=1):
    return fmt.h(doc, text, level)


def reference(doc, number, author, title, url):
    par = fmt.p(doc, f'[{number}] {author}. ', indent=False, size=12)
    par.paragraph_format.line_spacing = 1.1
    par.paragraph_format.space_after = Pt(6)
    par.paragraph_format.keep_together = True
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), doc.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), 'Times New Roman')
    props.append(fonts)
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '24')
    props.append(size)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.append(underline)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = title
    run.append(text)
    link.append(run)
    par._p.append(link)
    fmt.set_font(par.add_run('. Truy cập ngày 23/10/2026.'), 12)


def build():
    CITED_WORDS.clear()
    doc = Document()
    fmt.setup(doc)
    doc.core_properties.title = 'CHƯƠNG 5: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN'
    doc.core_properties.subject = 'Kết quả, hạn chế và định hướng mở rộng MugenBunko, bao gồm thương mại hóa và bảo mật giao dịch chuyển khoản'
    doc.core_properties.keywords = 'MugenBunko; kết luận; hướng phát triển; chuyển khoản; đối soát; bảo mật giao dịch'
    fmt.p(doc, 'CHƯƠNG 5: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN', style='Title', indent=False, size=14)
    p(doc, 'Qua quá trình phân tích, thiết kế, xây dựng và thử nghiệm, MugenBunko đã hình thành một hệ thống đọc, xuất bản và tương tác cộng đồng dành cho tiểu thuyết số. Chương này tổng hợp các kết quả đạt được, xác định những giới hạn còn tồn tại và đề xuất hướng phát triển tiếp theo. Trong đó, thương mại hóa nội dung và giao dịch chuyển khoản được xem là một hướng nghiên cứu mới, cần phát triển cùng với cơ chế bảo mật, ghi nhận tài chính và vận hành phù hợp.')

    h(doc, '5.1. Kết quả đạt được')
    p(doc, 'Về mục tiêu chức năng, hệ thống đã hiện thực các luồng phục vụ độc giả, tác giả, kiểm duyệt viên và quản trị viên. Người dùng có thể đăng ký, xác thực email bằng OTP, đăng nhập và sử dụng những chức năng phù hợp với vai trò. Cơ chế xác thực và phân quyền được thực hiện ở backend, tạo cơ sở quản lý quyền đối với tài khoản, nội dung và các thao tác nghiệp vụ.')
    fmt.table(doc, 'Bảng 5.1. Các nhóm chức năng MugenBunko đã hiện thực', ['Nhóm chức năng', 'Kết quả đạt được'], [
        ('Đọc và khám phá', 'Danh mục, tìm kiếm/lọc tác phẩm, mục lục, đọc và chuyển chương, tùy chỉnh giao diện đọc.'),
        ('Tủ sách và tiến độ', 'Bookmark, bộ sưu tập công khai/riêng tư, chia sẻ kệ, lưu vị trí theo đoạn và đồng bộ tiến độ qua server.'),
        ('Sáng tác và xuất bản', 'Studio quản lý truyện/tập/chương, bản nháp, tự lưu, lịch sử phiên bản, gửi duyệt và hẹn lịch phát hành.'),
        ('Tương tác cộng đồng', 'Diễn đàn, bình luận và trả lời, đánh giá, theo dõi tác giả, kết bạn, thông báo và trò chuyện thời gian thực.'),
        ('Kiểm duyệt và quản trị', 'Xử lý chương chờ duyệt, báo cáo bình luận, quản lý tài khoản, vai trò, thể loại, quy định và sự kiện.'),
        ('Tuyển chọn kiểm duyệt viên', 'Quản lý kỳ tuyển chọn, hồ sơ ứng viên, bỏ phiếu, đánh giá, công bố kết quả, khiếu nại và nhiệm kỳ.'),
    ], [4.6, 11.4])
    p(doc, 'Về kiến trúc, dự án tổ chức giao diện web React, ứng dụng React Native/Expo và backend Express theo các nhóm nghiệp vụ dùng chung MySQL. Cấu trúc này giúp dữ liệu tài khoản, tác phẩm, tủ sách và các tương tác đã lưu trên server có thể được truy xuất từ nhiều phía máy khách. Mô hình dữ liệu đã phân tách các thực thể chính, các bảng liên kết và dữ liệu theo dõi hoạt động; Socket.IO hỗ trợ chức năng giao tiếp thời gian thực.')
    p(doc, 'Về trải nghiệm sử dụng, hệ thống cung cấp các chế độ giao diện Washi, Sepia và Charcoal, điều chỉnh bố cục theo kích thước màn hình, đồng thời có xử lý trạng thái tải, lỗi và thử lại ở nhiều luồng. Trình đọc hỗ trợ tiếp tục vị trí đang đọc; Studio bổ sung cơ chế bảo toàn bản nháp và khôi phục phiên bản. Đây là những kết quả góp phần giảm thao tác lặp và hỗ trợ quá trình đọc, viết kéo dài.')
    p(doc, 'Về kiểm thử, hồ sơ ngày 22/10/2026 tại snapshot mã nguồn 2d7d95c ghi nhận 68 ca trong kế hoạch, trong đó 48 ca đã thực thi đạt và 20 ca chưa thực thi. Các nhóm đã chạy bao gồm 25 ca backend, 11 ca Node web/mobile, 11 nhóm kiểm thử giao diện web và một ca tổng hợp Expo web. Những kết quả này đã được trình bày ở Chương 4 và phản ánh đúng phạm vi của đợt kiểm thử được lưu, tạo căn cứ đánh giá các luồng cốt lõi đã được kiểm chứng.')
    p(doc, 'Từ các kết quả trên, đồ án đã hoàn thành phần nền tảng cho một hệ thống kết hợp đọc truyện, sáng tác và hoạt động cộng đồng trong môi trường cục bộ. Các thành phần hiện có là cơ sở để tiếp tục hoàn thiện chất lượng, kiểm chứng trên môi trường đích và mở rộng những nghiệp vụ mới theo nhu cầu vận hành thực tế.')

    h(doc, '5.2. Hạn chế của hệ thống')
    p(doc, 'Thứ nhất, phạm vi triển khai được ghi nhận chủ yếu là môi trường phát triển và thử nghiệm cục bộ. Dự án đã có cấu hình phục vụ đóng gói, nhưng hồ sơ hiện tại chưa chứng minh hệ thống đã được vận hành trên VPS với tên miền, HTTPS công khai, giám sát và quy trình khôi phục dữ liệu. Khả năng duy trì dịch vụ khi có sự cố hoặc nhiều người sử dụng cần được kiểm chứng trên môi trường gần với vận hành thực tế.')
    p(doc, 'Thứ hai, việc kiểm thử chưa bao phủ đầy đủ kế hoạch. Hai mươi ca chưa thực thi còn liên quan đến email thực, những tình huống biên trên giao diện, lịch phát hành theo thời gian thực, thông báo và các tương tác phụ thuộc thiết bị. Các kết quả mobile hiện có tập trung vào hàm tiện ích, API và Expo web; cần bổ sung kiểm thử Android/iOS thật cho bàn phím, bộ chọn ảnh, SecureStore, chuyển nền và phục hồi kết nối. Kết quả đóng gói JavaScript cũng chưa thay thế việc tạo, ký và nghiệm thu bản cài đặt ứng dụng.')
    p(doc, 'Thứ ba, khả năng mở rộng và hiệu năng còn cần được đánh giá thêm. Một số màn hình vẫn ghép nhiều trang dữ liệu ở client để phục vụ tìm kiếm/lọc; cấu hình vận hành hiện tại hướng đến một instance, trong khi trạng thái kết nối và một số giới hạn tốc độ nằm trong bộ nhớ tiến trình. Hồ sơ đo hiệu năng cục bộ mới cung cấp dữ liệu tham khảo, chưa đủ để kết luận số người dùng đồng thời hoặc mức cải thiện trên máy chủ thực tế. Build web còn có cảnh báo kích thước chunk cần xem xét tối ưu.')
    p(doc, 'Thứ tư, một số thiết lập vẫn mang tính cục bộ, như tùy chọn đọc và lựa chọn banner tiêu điểm của trình duyệt; việc đồng bộ cấu hình và dữ liệu danh mục còn có thể hoàn thiện. Các cơ chế bảo mật hiện có gồm băm mật khẩu, OTP, quản lý phiên, phân quyền và kiểm tra yêu cầu, nhưng hệ thống vẫn cần rà soát độc lập, kiểm thử lạm dụng nghiệp vụ và quy trình ứng phó sự cố trước khi vận hành với dữ liệu hoặc giao dịch có giá trị cao.')
    p(doc, 'Thứ năm, hệ thống chưa có phân hệ thanh toán tiền thực được kiểm chứng. Các API nạp coin và duyệt rút tiền hiện đang bị tắt; trường coins trong dữ liệu tài khoản không đủ cơ sở để xem là số dư tiền thật. Dự án còn thiếu vòng đời đơn hàng thanh toán, thông báo giao dịch ngân hàng được xác thực, sổ giao dịch, đối soát, hoàn tiền và quyết toán doanh thu tác giả. Vì vậy, thương mại hóa đòi hỏi thiết kế một phân hệ riêng thay vì chỉ mở lại những API cũ.')

    h(doc, '5.3. Hướng phát triển')
    h(doc, '5.3.1. Tối ưu hiệu năng và khả năng mở rộng', 2)
    p(doc, 'Hướng ưu tiên là chuyển dần tìm kiếm, lọc và sắp xếp sang backend, kết hợp phân trang phù hợp cho danh mục, chương, bình luận và tin nhắn. Cần đo truy vấn thực tế, rà soát chỉ mục và tránh nạp nhiều dữ liệu hơn nhu cầu hiển thị. Giao diện web có thể tiếp tục chia gói theo màn hình, tải chậm những phần ít dùng, tối ưu ảnh và tài nguyên tĩnh để giảm dung lượng tải ban đầu.')
    p(doc, 'Đối với dữ liệu được truy cập lặp lại, có thể nghiên cứu cache với chính sách hết hạn và cập nhật rõ ràng, đồng thời bảo đảm dữ liệu riêng tư được phân tách đúng tài khoản. Những tác vụ như gửi email, tổng hợp thống kê và xử lý thông báo có thể đưa vào hàng đợi nền. Khi cần nhiều instance, trạng thái kết nối, giới hạn tốc độ và phối hợp tác vụ cần sử dụng cơ chế dùng chung phù hợp.')
    p(doc, 'Việc tối ưu cần đi kèm phép đo trước và sau trên cùng dữ liệu, cấu hình và kịch bản. Các chỉ số nên theo dõi gồm độ trễ, tỷ lệ lỗi, thời gian truy vấn, dung lượng phản hồi và mức sử dụng tài nguyên. Kết quả đo sẽ làm căn cứ lựa chọn cấu hình máy chủ và giải pháp mở rộng thay cho việc xác định khả năng phục vụ chỉ từ cấu hình máy phát triển.')

    h(doc, '5.3.2. Nâng cao bảo mật và chất lượng hệ thống', 2)
    p(doc, 'Cần hoàn thiện các ca kiểm thử chưa chạy, bổ sung kiểm thử hồi quy cho quyền truy cập, vòng đời phiên, chức năng soạn thảo và các tình huống cập nhật đồng thời. Quy trình phát triển có thể tích hợp tự động kiểm tra kiểu, lint, build, rà soát dependency và kiểm thử vào mỗi lần thay đổi. Các cảnh báo và phát hiện bảo mật cần được theo dõi đến khi có biện pháp xử lý hoặc giới hạn phạm vi sử dụng rõ ràng.')
    p(doc, 'Về vận hành, cần quản lý tập trung bí mật ứng dụng, phân tách môi trường phát triển/thử nghiệm/vận hành, giới hạn quyền cơ sở dữ liệu và thiết lập nhật ký có phân quyền truy cập. Dữ liệu cá nhân, token và thông tin giao dịch cần được che phù hợp trong log. Có thể bổ sung xác thực nhiều yếu tố cho quản trị viên và quy trình theo dõi các thay đổi nhạy cảm trong quyền tài khoản, cấu hình hoặc chính sách nội dung.')

    h(doc, '5.3.3. Mở rộng triển khai và hoàn thiện ứng dụng di động', 2)
    p(doc, 'Giai đoạn tiếp theo có thể đưa một bản thử nghiệm lên VPS, cấu hình tên miền, HTTPS và reverse proxy, đồng thời giữ cơ sở dữ liệu trong phạm vi truy cập phù hợp. Cần chuẩn bị quy trình migration, sao lưu và khôi phục thử, quản lý tệp tải lên, giám sát trạng thái và kế hoạch quay lại phiên bản ổn định khi phát hành gặp lỗi. Môi trường staging sẽ được sử dụng để kiểm tra trước khi đưa thay đổi đến người dùng.')
    p(doc, 'Đối với mobile, cần tạo bản cài đặt phù hợp và kiểm thử trên nhiều thiết bị Android/iOS thật, đặc biệt ở trình đọc, editor, chọn ảnh, bàn phím và những tình huống ứng dụng chuyển nền. Sau khi hoàn thiện cấu hình ký số, quyền thiết bị và tiêu chí kiểm thử, có thể xem xét phát hành qua các kênh phân phối ứng dụng. Phản hồi từ người dùng thử sẽ hỗ trợ điều chỉnh khả năng tiếp cận và trải nghiệm trên màn hình nhỏ.')

    h(doc, '5.3.4. Nghiên cứu thương mại hóa nội dung', 2)
    p(doc, 'Hướng thương mại hóa có thể bao gồm mua quyền đọc chương hoặc tác phẩm, đăng ký gói đọc theo thời hạn và ủng hộ tác giả. Mỗi mô hình cần xác định rõ nội dung được cung cấp, giá, thời hạn truy cập, cách chia doanh thu và cách xử lý khi tác phẩm bị gỡ hoặc người dùng yêu cầu hỗ trợ. Cần khảo sát nhu cầu độc giả và tác giả trước khi lựa chọn mô hình triển khai đầu tiên.')
    p(doc, 'Ở giai đoạn đầu, có thể ưu tiên thanh toán trực tiếp theo từng đơn hàng nội dung để giới hạn phạm vi nghiệp vụ. Nếu nghiên cứu nạp tín dụng nội bộ, phải tách dữ liệu đó khỏi XP, điểm hoạt động và trường coins hiện có, đồng thời xác định rõ khả năng sử dụng, hoàn trả hoặc chuyển đổi. Mô hình lưu số dư, chuyển giá trị giữa người dùng và rút tiền cần được xem xét riêng về kỹ thuật, điều kiện đối tác và yêu cầu áp dụng trước khi quyết định triển khai.')
    p(doc, 'Thương mại hóa cũng cần thống nhất với tác giả về quyền khai thác tác phẩm, tỷ lệ phân chia, các khoản khấu trừ, thời điểm xác lập doanh thu và xử lý hoàn tiền. Những vấn đề về hợp đồng, dữ liệu cá nhân, chứng từ và nghĩa vụ liên quan cần được đánh giá cùng đơn vị thanh toán và chuyên môn phù hợp khi hoàn thiện mô hình. Nội dung ở đây là định hướng nghiên cứu, chưa phải kết luận hệ thống đã đáp ứng điều kiện vận hành thương mại.')

    h(doc, '5.3.5. Thiết kế giao dịch chuyển khoản và đối soát', 2)
    p(doc, 'Phân hệ dự kiến cần liên kết đơn hàng với yêu cầu thanh toán, giao dịch của đối tác và quyền đọc được cấp. Máy chủ tính giá từ dữ liệu tin cậy, xác minh kết quả với bên xử lý thanh toán và đối chiếu mã đơn, số tiền, loại tiền trước khi cung cấp nội dung. Trang chuyển hướng về ứng dụng hoặc ảnh chụp chuyển khoản chỉ phục vụ trải nghiệm/hỗ trợ, không được dùng làm căn cứ duy nhất để xác nhận đã nhận tiền.', refs=(1,))
    fmt.table(doc, 'Bảng 5.2. Luồng xử lý chuyển khoản dự kiến của MugenBunko', ['Bước', 'Xử lý đề xuất'], [
        ('1. Lập đơn hàng', 'Ghi người mua, nội dung, giá và điều kiện truy cập tại thời điểm mua; tạo mã đơn duy nhất.'),
        ('2. Tạo yêu cầu thanh toán', 'Tạo mã tham chiếu và thông tin chuyển khoản/QR theo đặc tả của đối tác được lựa chọn.'),
        ('3. Nhận thông báo', 'Tiếp nhận thông báo máy chủ; xác thực nguồn và dữ liệu trước khi đưa vào xử lý.'),
        ('4. Đối chiếu giao dịch', 'So khớp đơn, tài khoản nhận, số tiền và trạng thái; đưa trường hợp sai lệch vào hàng xử lý.'),
        ('5. Ghi nhận và cấp quyền', 'Ghi giao dịch, các bút toán nội bộ và quyền đọc bằng thao tác có kiểm soát chống lặp.'),
        ('6. Đối soát định kỳ', 'So sánh lịch sử đối tác với sổ nội bộ; phát hiện thiếu/trùng, hoàn tiền và các khoản chưa khớp.'),
    ], [4.6, 11.4])
    p(doc, 'Thông báo chuyển khoản có thể đến chậm hoặc bị bỏ lỡ khi endpoint gián đoạn. Cần đối chiếu định kỳ với API hoặc báo cáo giao dịch của đối tác để tìm các khoản chưa ghi nhận, đồng thời kiểm tra tài khoản nhận, mã thanh toán và số tiền thực tế. Chuyển thiếu, thừa, sai nội dung hoặc nhận tiền sau khi đơn hết hạn cần có quy trình xử lý riêng. Tài liệu webhook chuyển khoản của SePay được dùng làm ví dụ đối chiếu kỹ thuật; đối tác tích hợp cụ thể sẽ được lựa chọn ở giai đoạn nghiên cứu.', refs=(4,))
    p(doc, 'Dữ liệu tiền tệ cần quy định rõ loại tiền, đơn vị lưu, độ chính xác, giới hạn và cách làm tròn. Có thể sử dụng số nguyên theo đơn vị đã xác định hoặc DECIMAL với precision/scale phù hợp; MySQL cung cấp DECIMAL cho dữ liệu số chính xác, bao gồm dữ liệu tiền tệ. Khi truyền qua JavaScript và API, cần giữ nguyên độ chính xác và kiểm tra giới hạn giá trị.', refs=(5,))
    p(doc, 'Sổ tài chính dự kiến cần lưu từng biến động với mã tham chiếu, thời điểm, số tiền và lý do, thay vì chỉ giữ một giá trị số dư cuối. Có thể nghiên cứu ghi sổ kép để kiểm tra sự cân đối giữa các tài khoản nghiệp vụ. Điều chỉnh hoặc hoàn tiền được ghi bằng bút toán liên kết với giao dịch gốc, có dấu vết người thực hiện và trạng thái xử lý.')
    p(doc, 'Cần phân biệt trạng thái đơn hàng, trạng thái thanh toán và trạng thái cấp quyền đọc. Việc chuyển giữa các trạng thái phải được kiểm tra tại server. Giao dịch cơ sở dữ liệu và khóa phù hợp hỗ trợ các thay đổi nội bộ; thao tác qua đối tác bên ngoài còn cần khóa chống lặp, cơ chế thử lại và xử lý kết quả chưa xác định. Thiết kế này cần được kiểm chứng bằng các kịch bản đồng thời và sự cố, theo định hướng kiểm soát logic nghiệp vụ.', refs=(6,))

    h(doc, '5.3.6. Bảo mật giao dịch và chi trả tác giả', 2)
    p(doc, 'Endpoint nhận thông báo cần dùng HTTPS và xác minh chữ ký theo đúng đặc tả đối tác. Chống phát lại bằng mốc thời gian đã xác thực khi giao thức hỗ trợ; lưu mã sự kiện và mã giao dịch để chống xử lý trùng. Thông báo lặp hoặc đến sai thứ tự phải được xử lý theo trạng thái hiện hành. Khóa ký cần được bảo vệ, có kế hoạch thay đổi và giới hạn quyền truy cập.', refs=(2,))
    p(doc, 'Thao tác hoàn tiền, đổi tài khoản nhận hoặc duyệt chi trả cần kiểm tra quyền, dữ liệu và xác nhận gắn với chính giao dịch đó tại server. Nếu dữ liệu quan trọng thay đổi, xác nhận cũ cần mất hiệu lực. Có thể yêu cầu xác thực bổ sung cho các thao tác nhạy cảm; cơ chế xác thực đăng nhập thông thường chưa đủ để thay thế mọi kiểm soát ủy quyền giao dịch.', refs=(3,))
    p(doc, 'Luồng thanh toán dự kiến để người mua xác nhận chuyển khoản trên ứng dụng ngân hàng hoặc kênh chính thức của đối tác. MugenBunko tiếp nhận kết quả và thông tin tham chiếu cần thiết cho đơn hàng; hệ thống không yêu cầu người dùng cung cấp mật khẩu hay mã OTP ngân hàng. Dữ liệu xác thực ngân hàng được tách khỏi thông tin đăng nhập MugenBunko và khỏi các biểu mẫu hỗ trợ giao dịch.')
    p(doc, 'Đối với MugenBunko, cần thiết kế vai trò xử lý tài chính riêng và giới hạn người được xem, đề nghị hoặc phê duyệt khoản chi. Những khoản có mức rủi ro cao có thể áp dụng nguyên tắc hai người kiểm tra độc lập. Mọi thay đổi về tài khoản nhận, tỷ lệ chia doanh thu, điều chỉnh sổ và duyệt chi trả cần có lịch sử; thông tin tài khoản ngân hàng chỉ được thu thập, hiển thị và lưu trong phạm vi cần thiết cho nghiệp vụ.')
    p(doc, 'Doanh thu tác giả nên được phân biệt giữa khoản đang chờ, khoản đã đối soát và khoản đủ điều kiện chi trả theo chính sách. Một lần gọi API chi trả bị timeout cần được kiểm tra trạng thái tại đối tác trước khi gửi lại. Các yêu cầu chi trả cần có mã chống lặp, giới hạn tần suất và quy trình xác minh người nhận. Hoàn tiền hoặc tranh chấp cũng phải được phản ánh vào báo cáo doanh thu và các khoản còn phải trả.')
    p(doc, 'Bộ kiểm thử thanh toán dự kiến cần bao gồm sửa giá hoặc người thụ hưởng, webhook giả, chữ ký sai, phát lại, thông báo trùng/sai thứ tự, thanh toán thiếu/thừa, timeout và hai yêu cầu đồng thời. Cũng cần thử trường hợp đã nhận tiền nhưng chưa cấp quyền, hoàn tiền bị gián đoạn và chi trả chưa rõ kết quả. Các thử nghiệm được thực hiện trong môi trường kiểm thử hoặc sandbox của đối tác khi có hỗ trợ, trước khi mở giao dịch có giá trị thực.')

    h(doc, '5.3.7. Lộ trình thực hiện', 2)
    fmt.table(doc, 'Bảng 5.3. Lộ trình phát triển đề xuất', ['Giai đoạn', 'Trọng tâm', 'Căn cứ chuyển bước'], [
        ('1. Hoàn thiện nền tảng', 'Đóng các ca kiểm thử còn thiếu, rà soát bảo mật và đo hiệu năng.', 'Các luồng trọng tâm được kiểm chứng; có kết quả hồi quy và kế hoạch xử lý tồn tại.'),
        ('2. Triển khai thử nghiệm', 'VPS/staging, HTTPS, sao lưu/khôi phục, giám sát và thiết bị native.', 'Có bằng chứng vận hành, khôi phục dữ liệu và kết quả thử nghiệm trên thiết bị.'),
        ('3. Thử nghiệm thanh toán', 'Mô hình bán nội dung, đơn hàng, chuyển khoản, sổ và đối soát.', 'Đối chiếu được dữ liệu thử; các ca chống lặp, sai lệch và sự cố được kiểm chứng.'),
        ('4. Mở rộng thương mại', 'Thử nghiệm có kiểm soát, hoàn tiền và quyết toán tác giả.', 'Hoàn thiện điều kiện vận hành, chính sách và kiểm soát giao dịch theo phạm vi triển khai.'),
    ], [3.6, 6.0, 6.4])
    p(doc, 'Lộ trình được tổ chức theo mức độ sẵn sàng của chức năng và khả năng kiểm soát rủi ro. Việc mở rộng thương mại sẽ dựa trên kết quả kỹ thuật, nhu cầu sử dụng, thỏa thuận với tác giả và điều kiện tích hợp thanh toán. Với nền tảng đã xây dựng, MugenBunko có cơ sở tiếp tục phát triển thành một không gian đọc và xuất bản nội dung có khả năng vận hành thực tế, trong đó an toàn dữ liệu và sự minh bạch của giao dịch là yêu cầu xuyên suốt.')

    ref_title = fmt.p(doc, 'Tài liệu đối chiếu cho định hướng thanh toán và bảo mật', indent=False, size=12)
    fmt.set_font(ref_title.runs[0], 12, bold=True)
    ref_title.paragraph_format.keep_with_next = True
    ref_title.paragraph_format.space_before = Pt(10)
    for number, (author, title, url) in enumerate(REFERENCES, 1):
        reference(doc, number, author, title, url)
    doc.save(OUTPUT)
    validate(doc)


def validate(doc):
    text = '\n'.join([par.text for par in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
    assert text.startswith('CHƯƠNG 5: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN')
    headings = [par.text for par in doc.paragraphs if par.style.name == 'Heading 1']
    assert headings == ['5.1. Kết quả đạt được', '5.2. Hạn chế của hệ thống', '5.3. Hướng phát triển']
    assert len(doc.sections) == 1 and len(doc.inline_shapes) == 0
    assert len(doc.tables) == 3
    assert '\ufffd' not in text
    assert set(CITED_WORDS) == set(range(1,7))
    assert all(n <= 180 for n in CITED_WORDS.values()), dict(CITED_WORDS)
    with ZipFile(OUTPUT) as z:
        rels = z.read('word/_rels/document.xml.rels').decode('utf-8')
        for _, _, url in REFERENCES:
            assert url in rels
    report = {
        'document': str(OUTPUT), 'document_date': '2026-10-23',
        'title': doc.core_properties.title, 'main_sections': headings,
        'subsections_5_3': 7, 'tables': len(doc.tables),
        'paper': 'A4 portrait', 'body_font': 'Times New Roman 13 pt',
        'table_and_reference_font': 'Times New Roman 12 pt', 'line_spacing_body': 1.5,
        'words_whitespace': len(text.split()), 'characters': len(text),
        'recorded_test_results': {'date':'2026-10-22','commit':'2d7d95c','PASS':48,'NOT_RUN':20},
        'current_payment_state': 'deposit-coins and cashout approval are disabled; no verified real-money subsystem',
        'payment_content_status': 'Proposed research and design, not implemented or deployed functionality',
        'local_sources': ['artifacts/chapter-4/chapter4_text.txt','artifacts/test-report/all-test-cases.json',
                          'server/routes/auth.ts','server/routes/admin.ts','server/realtime.ts',
                          'server/middleware/security.ts','mobile/README.md'],
        'external_sources': [{'number':i,'author':a,'title':t,'url':u,'accessed':'2026-10-23'} for i,(a,t,u) in enumerate(REFERENCES,1)],
        'cited_paragraph_words_by_source': dict(CITED_WORDS),
    }
    (OUT/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'chapter5_text.txt').write_text(text,encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k not in ('local_sources','external_sources')},ensure_ascii=False,indent=2))


if __name__ == '__main__':
    build()
