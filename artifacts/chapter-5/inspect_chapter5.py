"""Check the native Word export and render representative document pages."""
from pathlib import Path
import json
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0,str(Path(os.environ['TEMP'])/'mugenbunko-test-report-libs'))
import pypdfium2 as pdfium
from docx import Document
from PIL import Image, ImageDraw, ImageFont

OUT=Path(__file__).resolve().parent
NAME='Chuong_5_Ket_luan_va_huong_phat_trien_MugenBunko'
pdf=pdfium.PdfDocument(str(OUT/(NAME+'.pdf')))
word=Document(OUT/(NAME+'.docx'))
cols,tile_w,tile_h=4,330,470
sheet=Image.new('RGB',(cols*tile_w,tile_h*((len(pdf)+cols-1)//cols)),'#DADDE0')
draw=ImageDraw.Draw(sheet)
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',15)
records,texts,selected=[],[],[]
for i in range(len(pdf)):
    page=pdf[i]
    w,h=page.get_size()
    text=page.get_textpage().get_text_range()
    texts.append(text)
    lines=[s.strip() for s in text.splitlines() if s.strip()]
    row={'page':i+1,'size_pt':[round(w,1),round(h,1)],'characters':len(text),
         'first_lines':lines[:2],'last_lines':lines[-2:]}
    records.append(row)
    print(json.dumps(row,ensure_ascii=False))
    x,y=i%cols*tile_w+15,i//cols*tile_h+30
    sheet.paste(page.render(scale=300/w).to_pil().convert('RGB'),(x,y))
    draw.text((x,y-24),f'Trang {i+1}',font=font,fill='black')
    if i in (0,len(pdf)-1) or 'Bảng 5.2.' in text or '5.3.6.' in text:
        selected.append(i+1)
        page.render(scale=1.4).to_pil().save(OUT/f'preview_page_{i+1:02d}.png')
sheet.save(OUT/'preview_contact_sheet.png')
joined='\n\f\n'.join(texts)
compact=''.join(joined.split())
expected=[p.text for p in word.paragraphs if p.style.name in ('Heading 1','Heading 2') or p.text.startswith('Bảng 5.')]
missing=[t for t in expected if ''.join(t.split()) not in compact]
assert not missing,missing
assert all(r['characters']>60 for r in records), 'Unexpected blank page'
assert all('['+str(n)+']' in joined for n in range(1,7))
(OUT/'preview_text.txt').write_text(joined,encoding='utf-8')
(OUT/'preview_pages.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
report=json.loads((OUT/'validation.json').read_text(encoding='utf-8'))
report.update({'opened_and_exported_in_microsoft_word':True,'pages_microsoft_word':len(pdf),
               'all_headings_captions_and_references_present_in_pdf':True,'preview_selected_pages':selected})
(OUT/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print('PAGES',len(pdf),'SELECTED',selected)
