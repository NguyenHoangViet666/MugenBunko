"""Inspect a native Word PDF export without changing the document or screenshots."""
from pathlib import Path
import json
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, str(Path(os.environ['TEMP']) / 'mugenbunko-test-report-libs'))
import pypdfium2 as pdfium
from docx import Document
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent
NAME = 'Chuong_4_Trien_khai_he_thong_MugenBunko'
doc = pdfium.PdfDocument(str(OUT / (NAME + '.pdf')))
word = Document(OUT / (NAME + '.docx'))
font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 15)
cols, tile_w, tile_h, thumb_w = 4, 330, 470, 300
sheet = Image.new('RGB', (cols*tile_w, tile_h*((len(doc)+cols-1)//cols)), '#DADDE0')
draw = ImageDraw.Draw(sheet)
records, texts, selected = [], [], set()
for i in range(len(doc)):
    page = doc[i]
    w, h = page.get_size()
    text = page.get_textpage().get_text_range()
    texts.append(text)
    lines = [s.strip() for s in text.splitlines() if s.strip()]
    record = {'page': i+1, 'size_pt': [round(w,1),round(h,1)], 'characters': len(text),
              'first_lines': lines[:2], 'last_lines': lines[-2:]}
    records.append(record)
    print(json.dumps(record, ensure_ascii=False))
    thumb = page.render(scale=thumb_w/w).to_pil().convert('RGB')
    x, y = i%cols*tile_w+15, i//cols*tile_h+30
    sheet.paste(thumb,(x,y))
    draw.text((x,y-24),f'Trang {i+1}',font=font,fill='black')
    if i in (0,len(doc)-1) or 'Hình 4.1.' in text or 'Hình 4.9.' in text or 'Bảng 4.5.' in text:
        selected.add(i+1)
        page.render(scale=1.35).to_pil().save(OUT/f'preview_page_{i+1:02d}.png')
sheet.save(OUT/'preview_contact_sheet.png')
joined='\n\f\n'.join(texts)
compact=''.join(joined.split())
expected=[p.text for p in word.paragraphs if p.text.startswith(('4.1. ','4.2. ','4.3. ','4.4. ','4.5. ','4.6. ','Hình 4.','Bảng 4.'))]
missing=[s for s in expected if ''.join(s.split()) not in compact]
assert not missing,missing
assert all(r['characters']>40 for r in records), 'Unexpected blank page'
report=json.loads((OUT/'validation.json').read_text(encoding='utf-8'))
report.update({'opened_and_exported_in_microsoft_word':True, 'pages_microsoft_word':len(doc),
               'section_headings_and_captions_present_in_pdf':True, 'preview_selected_pages':sorted(selected)})
(OUT/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'preview_pages.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'preview_text.txt').write_text(joined,encoding='utf-8')
print('PAGES',len(doc),'SELECTED',sorted(selected))
