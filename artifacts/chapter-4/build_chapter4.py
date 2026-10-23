"""Create the report's implementation chapter from local code and recorded evidence."""
from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, str(Path(os.environ['TEMP']) / 'mugenbunko-test-report-libs'))
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
OUTPUT = OUT / 'Chuong_4_Trien_khai_he_thong_MugenBunko.docx'
FONT = 'Times New Roman'
FIGURES = []
SOURCES = [
    'package-lock.json', 'server/package-lock.json', 'mobile/package-lock.json',
    'server/server.ts', 'server/config.ts', 'server/db.ts', 'server/createAdmin.ts',
    'server/routes/auth.ts', 'server/routes/reader.ts', 'server/routes/chapters.ts',
    'server/routes/mod.ts', 'server/routes/admin.ts', 'server/routes/elections.ts',
    'server/routes/upload.ts', 'server/realtime.ts', 'server/utils/sessions.ts',
    'src/App.tsx', 'src/hooks/useStudioEditor.ts', 'mobile/README.md',
    'mobile/src/services/storage/authToken.ts', 'mobile/src/services/api/client.ts',
    'artifacts/test-report/all-test-cases.json',
    'artifacts/test-report/evidence/quality-checks.jsonl',
    'artifacts/test-report/evidence/frontend-mobile-summary.json',
    'artifacts/test-report/evidence/web-browser.log',
    'artifacts/test-report/mobile-browser-case.json',
]


def load(path):
    return json.loads((ROOT / path).read_text(encoding='utf-8-sig'))


def dep(lockfile, name):
    return load(lockfile)['packages']['node_modules/' + name]['version']


def set_font(run, size=13, *, bold=None, italic=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    fonts.attrib.clear()
    for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + k), name)
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'vi-VN')
    rpr.append(language)


def p(doc, text, *, style=None, indent=True, size=13, italic=False):
    par = doc.add_paragraph(style=style)
    set_font(par.add_run(text), size, italic=italic)
    if not indent:
        par.paragraph_format.first_line_indent = Cm(0)
    return par


def h(doc, text, level=1, *, new_page=False):
    par = p(doc, text, style='Heading ' + str(level), indent=False)
    par.paragraph_format.page_break_before = new_page
    return par


def caption(doc, text, *, keep=False):
    par = p(doc, text, style='Caption', indent=False, size=12, italic=True)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = keep
    return par


def code(doc, lines):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.left_indent = Cm(0.25)
    pf.right_indent = Cm(0.15)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = 1.05
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.keep_together = True
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), 'F2F2F2')
    par._p.get_or_add_pPr().append(shade)
    set_font(par.add_run('\n'.join(lines)), 10, name='Consolas')


def table(doc, title, headers, rows, widths):
    caption(doc, title, keep=True)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths):
        col.width = Cm(w)
    for cell, w, text in zip(t.rows[0].cells, widths, headers):
        cell.width = Cm(w)
        cell.text = text
    for values in rows:
        for cell, w, text in zip(t.add_row().cells, widths, values):
            cell.width = Cm(w)
            cell.text = str(text)
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), 'E9EDF1')
                cell._tc.get_or_add_tcPr().append(shd)
            for par in cell.paragraphs:
                pf = par.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.line_spacing = 1.1
                pf.space_before = Pt(2)
                pf.space_after = Pt(2)
                pf.keep_with_next = i == 0
                for run in par.runs:
                    set_font(run, 12, bold=i == 0)
    return t


def figure(doc, paths, text, *, max_height=12.4):
    if isinstance(paths, str):
        paths = [paths]
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.line_spacing = 1
    pf.space_before = Pt(5)
    pf.space_after = Pt(0)
    pf.keep_with_next = True
    for i, rel in enumerate(paths):
        path = ROOT / rel
        with Image.open(path) as im:
            w, hh = im.size
        allowed_w = 16 if len(paths) == 1 else 7.1
        width = min(allowed_w, max_height*w/hh)
        pic = par.add_run().add_picture(str(path), width=Cm(width))
        pic._inline.docPr.set('descr', text + ' — ' + rel)
        if i + 1 < len(paths):
            set_font(par.add_run('    '), 13)
        FIGURES.append({'path': rel, 'pixels': [w, hh], 'caption': text,
                        'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    caption(doc, f'Hình 4.{1 + sum(1 for par in doc.paragraphs if par.text.startswith("Hình 4."))}. {text}')


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.header_distance, sec.footer_distance = Cm(0.7), Cm(0.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'Footer'):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(13)
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        fonts = OxmlElement('w:rFonts')
        for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + k), FONT)
        rpr.append(fonts)
    normal = doc.styles['Normal'].paragraph_format
    normal.first_line_indent = Cm(1)
    normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.line_spacing = 1.5
    normal.space_after = Pt(6)
    normal.widow_control = True
    for name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        st = doc.styles[name]
        st.font.bold = True
        pf = st.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.line_spacing = 1.15
        pf.space_before = Pt(10)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        pf.keep_together = True
    doc.styles['Title'].font.size = Pt(14)
    doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.styles['Caption'].paragraph_format
    cap.first_line_indent = Cm(0)
    cap.line_spacing = 1.1
    cap.space_before = Pt(5)
    cap.space_after = Pt(7)
    cap.keep_together = True
    cap.keep_with_next = False
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    footer.paragraph_format.line_spacing = 1
    footer.paragraph_format.space_after = Pt(0)
    set_font(footer.add_run('Trang '), 12)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    run = OxmlElement('w:r')
    tx = OxmlElement('w:t')
    tx.text = '1'
    run.append(tx)
    fld.append(run)
    footer._p.append(fld)
    doc.core_properties.title = 'CHƯƠNG 4: TRIỂN KHAI HỆ THỐNG'
    doc.core_properties.subject = 'Triển khai cục bộ và kết quả thực hiện hệ thống MugenBunko'
    doc.core_properties.author = 'Nguyễn Hoàng Việt'
    doc.core_properties.keywords = 'MugenBunko; chương 4; triển khai hệ thống; localhost; kiểm thử'


def build():
    FIGURES.clear()
    doc = Document()
    setup(doc)
    p(doc, 'CHƯƠNG 4: TRIỂN KHAI HỆ THỐNG', style='Title', indent=False, size=14)
    p(doc, 'Chương này trình bày quá trình hiện thực, cấu hình và vận hành hệ thống MugenBunko trong môi trường phát triển, thử nghiệm cục bộ. Từ kiến trúc và thiết kế dữ liệu ở chương trước, hệ thống được tổ chức thành giao diện web, ứng dụng di động và máy chủ xử lý nghiệp vụ dùng chung. Nội dung tập trung vào các thành phần đã có trong mã nguồn, kết quả giao diện và hồ sơ kiểm thử đã được ghi nhận.')

    h(doc, '4.1. Phạm vi và môi trường triển khai')
    h(doc, '4.1.1. Phạm vi triển khai trong đồ án', 2)
    p(doc, 'MugenBunko được triển khai theo mô hình máy khách – máy chủ. Trong môi trường cục bộ, ứng dụng web chạy trên trình duyệt; backend cung cấp REST API và kết nối thời gian thực; MySQL lưu dữ liệu dùng chung. Ứng dụng di động được xây dựng bằng React Native và Expo, truy cập cùng backend để sử dụng tài khoản, tác phẩm, tủ sách và các chức năng cộng đồng.')
    p(doc, 'Phạm vi kết quả của chương bao gồm tổ chức mã nguồn, cấu hình kết nối, khởi tạo dữ liệu, các chức năng đã hiện thực và kết quả thử nghiệm được lưu trong dự án. Giao diện mobile được minh họa bằng bản Expo web trong trình duyệt mô phỏng kích thước điện thoại. Việc đưa hệ thống lên VPS, sử dụng tên miền, HTTPS công khai và phát hành ứng dụng lên kho được xem là các bước triển khai tiếp theo, trình bày trong hướng phát triển ở Chương 5.')
    h(doc, '4.1.2. Môi trường phần cứng và phần mềm', 2)
    p(doc, 'Máy phát triển và soạn báo cáo hiện tại sử dụng Windows 11 64 bit, bộ xử lý Intel Core i3-1115G4 và khoảng 16 GB RAM. Đây là cấu hình máy làm việc thực tế, không phải cấu hình tối thiểu hoặc thông số đề xuất cho máy chủ phục vụ nhiều người dùng. Node.js đang sử dụng là v24.13.0, đi kèm npm 11.6.2.')
    table(doc, 'Bảng 4.1. Các thành phần phần mềm của hệ thống', ['Thành phần', 'Phiên bản đối chiếu', 'Vai trò trong triển khai'], [
        ('Node.js / npm', '24.13.0 / 11.6.2', 'Môi trường thực thi và quản lý gói trên máy phát triển.'),
        ('React / Vite (web)', dep('package-lock.json', 'react') + ' / ' + dep('package-lock.json', 'vite'), 'Xây dựng giao diện SPA và đóng gói tài nguyên web.'),
        ('TypeScript (web)', dep('package-lock.json', 'typescript'), 'Kiểm tra kiểu trong mã nguồn giao diện.'),
        ('Express / mysql2', dep('server/package-lock.json', 'express') + ' / ' + dep('server/package-lock.json', 'mysql2'), 'Cung cấp API và truy cập cơ sở dữ liệu.'),
        ('Socket.IO', dep('server/package-lock.json', 'socket.io'), 'Kết nối thời gian thực cho chức năng xã hội.'),
        ('Expo / React Native', dep('mobile/package-lock.json', 'expo') + ' / ' + dep('mobile/package-lock.json', 'react-native'), 'Phát triển ứng dụng di động dùng chung API.'),
        ('MySQL', 'Lược đồ MySQL 8.x', 'Lưu dữ liệu quan hệ bằng InnoDB, hỗ trợ utf8mb4.'),
    ], [4.0, 4.2, 7.8])
    p(doc, 'Phiên bản thư viện trong Bảng 4.1 được đối chiếu từ lockfile của từng phần web, backend và mobile tại thời điểm soạn chương. Các phần này quản lý dependency riêng; vì vậy phiên bản React hoặc TypeScript giữa web và mobile có thể khác nhau. Khi tái tạo môi trường, sử dụng npm ci để cài theo lockfile đã lưu thay vì tự nâng đồng loạt các thư viện.')

    h(doc, '4.2. Tổ chức hệ thống khi triển khai', new_page=True)
    h(doc, '4.2.1. Kiến trúc các thành phần chạy', 2)
    p(doc, 'Web React và ứng dụng Expo đóng vai trò máy khách, gửi yêu cầu đến backend thông qua HTTP. Express tiếp nhận yêu cầu, kiểm tra dữ liệu, xác thực người dùng, phân quyền và gọi lớp truy cập MySQL. Socket.IO được gắn vào cùng HTTP server để phục vụ các sự kiện thời gian thực. Các router nghiệp vụ nằm trong cùng một ứng dụng backend, không được triển khai như những dịch vụ độc lập.')
    p(doc, 'Trong chế độ phát triển, Vite thường cung cấp web tại cổng 5173; backend mặc định lắng nghe cổng 5000 và MySQL thường dùng cổng 3306. Cổng web thực tế cần lấy theo thông báo khi khởi động Vite. Khi có thư mục dist của bản build web, Express có khả năng phục vụ tài nguyên tĩnh cùng API; chức năng này đã có trong mã nguồn nhưng không đồng nghĩa hệ thống đã được công bố trên Internet.')
    table(doc, 'Bảng 4.2. Tổ chức các thư mục triển khai', ['Thư mục / tệp', 'Nội dung và nhiệm vụ'], [
        ('src/', 'Giao diện web, trang chức năng, component, hook và tiện ích phía trình duyệt.'),
        ('server/', 'Máy chủ Express, router, middleware, Socket.IO và truy cập MySQL.'),
        ('mobile/', 'Ứng dụng Expo, màn hình, dịch vụ API và lưu trữ trên thiết bị.'),
        ('shared/', 'Các hợp đồng dữ liệu và nhãn dùng chung giữa các phần giao diện.'),
        ('public/, dist/', 'Tài nguyên công khai và kết quả đóng gói web.'),
        ('tests/, server/tests/', 'Kiểm thử đơn vị, API, phân quyền và kiểm thử trình duyệt.'),
        ('artifacts/', 'Ảnh giao diện, log kiểm thử và các tài liệu kết quả.'),
    ], [4.5, 11.5])
    h(doc, '4.2.2. Phân chia API và dữ liệu dùng chung', 2)
    p(doc, 'Backend chia chức năng theo nhóm router để thuận tiện quản lý. Máy khách gọi các API theo tiền tố /api; dữ liệu tài khoản và quyền được kiểm tra ở máy chủ trước khi thực hiện các thao tác cần đăng nhập. Cả web và mobile sử dụng cùng cơ sở dữ liệu, nhờ đó các dữ liệu đã lưu trên server có thể được truy xuất ở nhiều thiết bị.')
    table(doc, 'Bảng 4.3. Một số nhóm API chính', ['Nhóm đường dẫn', 'Nghiệp vụ phụ trách'], [
        ('/api/auth', 'Đăng ký, OTP, đăng nhập, đăng xuất, hồ sơ và theo dõi tác giả.'),
        ('/api/novels, /api/chapters', 'Danh mục truyện, nội dung chương và nghiệp vụ sáng tác.'),
        ('/api/reader, /api/library', 'Vị trí đọc, hoàn thành chương và bộ sưu tập cá nhân.'),
        ('/api/comments, /api/reviews', 'Bình luận, phản hồi và đánh giá tác phẩm.'),
        ('/api/forum, /api/social', 'Diễn đàn, kết bạn và lịch sử trò chuyện.'),
        ('/api/mod, /api/admin', 'Kiểm duyệt nội dung và quản trị theo vai trò.'),
        ('/api/elections', 'Ứng cử, bỏ phiếu, đánh giá và quản lý kỳ tuyển chọn.'),
        ('/api/upload, /healthz', 'Tải ảnh và kiểm tra tình trạng kết nối của backend.'),
    ], [6.0, 10.0])
    p(doc, 'Đường dẫn /healthz kiểm tra truy vấn SELECT 1 để phản ánh khả năng kết nối cơ sở dữ liệu. Dữ liệu ảnh được lưu cục bộ hoặc qua Cloudinary tùy cấu hình. Các biến môi trường quyết định nơi lưu và URL truy cập ảnh, nên phải thống nhất với địa chỉ backend mà máy khách sử dụng.')

    h(doc, '4.3. Cài đặt và cấu hình môi trường cục bộ', new_page=True)
    h(doc, '4.3.1. Chuẩn bị mã nguồn và thư viện', 2)
    p(doc, 'Trước khi khởi động hệ thống, chuẩn bị Node.js, npm và một dịch vụ MySQL có thể truy cập từ máy phát triển. Mã nguồn gồm ba phần có package và lockfile riêng. Các lệnh dưới đây thực hiện từ thư mục gốc dự án; phần mobile có thể cài sau nếu chỉ thử nghiệm web và backend trước.')
    code(doc, ['npm ci', 'npm --prefix server ci', 'npm --prefix mobile ci'])
    p(doc, 'Sau khi cài thư viện, tạo các tệp cấu hình cục bộ từ tệp mẫu. Tệp cấu hình backend đặt trong thư mục server vì server/config.ts tải biến môi trường từ vị trí này. Tệp mobile nằm trong thư mục mobile và cung cấp địa chỉ API cho bundle ứng dụng.')
    code(doc, ['Copy-Item server/.env.example server/.env', 'Copy-Item mobile/.env.example mobile/.env'])
    h(doc, '4.3.2. Cấu hình backend và MySQL', 2)
    p(doc, 'Trong server/.env, khai báo địa chỉ, cổng, tài khoản và mật khẩu MySQL phù hợp với môi trường cục bộ. DB_NAME xác định cơ sở dữ liệu của ứng dụng; JWT_SECRET sử dụng một chuỗi ngẫu nhiên riêng để giữ ổn định phiên đăng nhập giữa các lần khởi động. Báo cáo chỉ mô tả tên biến và mục đích, không đưa thông tin xác thực đang sử dụng vào nội dung tài liệu.')
    table(doc, 'Bảng 4.4. Các biến môi trường cần cấu hình cho môi trường cục bộ', ['Biến / nhóm biến', 'Giá trị hoặc cách sử dụng'], [
        ('NODE_ENV, PORT', 'development; cổng backend thường là 5000.'),
        ('DB_HOST, DB_PORT', 'Địa chỉ MySQL và cổng truy cập, thường 127.0.0.1:3306.'),
        ('DB_USER, DB_PASSWORD, DB_NAME', 'Thông tin truy cập riêng của môi trường phát triển.'),
        ('JWT_SECRET', 'Chuỗi bí mật ngẫu nhiên, được giữ ngoài mã nguồn công khai.'),
        ('FRONTEND_URL', 'Origin của web, ví dụ http://localhost:5173.'),
        ('PUBLIC_URL', 'URL backend dùng để tạo liên kết ảnh, ví dụ http://localhost:5000.'),
        ('DEV_OTP', 'Có thể đặt true có chủ đích để thử xác thực trong development.'),
        ('SMTP_HOST/PORT/USER/PASS', 'Thông tin máy chủ gửi thư khi cần kiểm tra email thực.'),
        ('CLOUDINARY_*', 'Để trống cả ba biến nếu sử dụng lưu ảnh trên đĩa cục bộ.'),
    ], [6.0, 10.0])
    p(doc, 'Khi thử lưu ảnh cục bộ, cần xóa các giá trị minh họa trong CLOUDINARY_CLOUD_NAME, CLOUDINARY_API_KEY và CLOUDINARY_API_SECRET của tệp mẫu. Backend kiểm tra sự hiện diện của cả ba giá trị để chọn Cloudinary; chuỗi placeholder vẫn có thể bị hiểu là đã cấu hình dịch vụ. Nếu không cấu hình Cloudinary, ảnh được lưu trong thư mục uploads của backend.')
    p(doc, 'Ở development, backend kiểm tra cơ sở dữ liệu, tạo các bảng còn thiếu và áp dụng các migration bổ sung. Quá trình bootstrap trong mã nguồn chỉ lấy phần định nghĩa lược đồ, không thực hiện phần đặt lại dữ liệu mẫu ở cuối schema.sql. Vì vậy, quy trình chạy cục bộ sử dụng cơ chế khởi tạo của ứng dụng; dữ liệu kiểm thử được tạo riêng trong fixture khi chạy các bộ kiểm thử.')
    p(doc, 'Chức năng đăng ký sử dụng mã xác thực email. Nếu cần kiểm tra việc nhận thư, phải cấu hình SMTP và xác minh tại hộp thư thật. Với thử nghiệm cục bộ, có thể bật DEV_OTP=true để nhận mã phục vụ kiểm thử qua phản hồi API trong chế độ development. Khi SMTP chưa sẵn sàng và chế độ OTP thử nghiệm không được bật, API gửi mã có thể trả lỗi dịch vụ chưa khả dụng; đây không phải bằng chứng gửi thư thành công.')
    h(doc, '4.3.3. Khởi động backend, web và tài khoản quản trị', 2)
    p(doc, 'Sau khi MySQL hoạt động và cấu hình được lưu, có thể khởi động web cùng backend bằng lệnh npm run dev. Dự án sử dụng concurrently để chạy đồng thời hai tiến trình phát triển. Trình duyệt truy cập URL do Vite thông báo; backend dùng cổng đã khai báo trong server/.env.')
    code(doc, ['npm run dev'])
    p(doc, 'Khi cần kiểm tra riêng từng thành phần, sử dụng các lệnh sau ở hai cửa sổ terminal khác nhau. Việc chạy riêng giúp quan sát log backend, lỗi truy vấn hoặc lỗi giao diện mà không nhầm lẫn giữa các tiến trình.')
    code(doc, ['npm run dev:backend', 'npm run dev:frontend'])
    p(doc, 'Có thể kiểm tra backend qua http://localhost:5000/healthz. Phản hồi status: ok thể hiện truy vấn kiểm tra MySQL thực hiện được; trạng thái này cần kết hợp với thử nghiệm đăng nhập, đọc và ghi dữ liệu để đánh giá chức năng. Trong chế độ phát triển, web mặc định gọi http://localhost:5000/api; nếu thay địa chỉ, cần cấu hình VITE_API_BASE cho phần web và khởi động lại tiến trình liên quan.')
    p(doc, 'Tài khoản quản trị đầu tiên được tạo bằng tiện ích riêng sau khi build backend. Tiện ích yêu cầu ADMIN_USERNAME hợp lệ và ADMIN_PASSWORD có độ dài tối thiểu 16 ký tự, không quá 72 byte. Lệnh từ chối ghi đè hoặc nâng quyền một tài khoản đã tồn tại; do đó không dựa vào tài khoản seed được ghi trong các hướng dẫn cũ để mô tả môi trường khởi tạo mới.')
    code(doc, ['npm --prefix server run build', '$env:ADMIN_USERNAME = "mugen_admin"', '$env:ADMIN_PASSWORD = "<mat_khau_rieng_it_nhat_16_ky_tu>"', 'npm --prefix server run create:admin'])
    p(doc, 'Giá trị mật khẩu trong ví dụ là ký hiệu thay thế và cần được thay bằng mật khẩu riêng khi thực hiện. Sau bước tạo tài khoản, quản trị viên đăng nhập qua giao diện, kiểm tra quyền truy cập và thiết lập các dữ liệu ban đầu như thể loại, quy định và tài khoản thử nghiệm theo nhu cầu sử dụng.')
    h(doc, '4.3.4. Cấu hình ứng dụng di động trong mạng cục bộ', 2)
    p(doc, 'Ứng dụng di động cần truy cập được backend qua một địa chỉ phù hợp với thiết bị. Trong mobile/.env, EXPO_PUBLIC_API_URL trỏ đến backend và EXPO_PUBLIC_WEB_URL trỏ đến web dùng cho các liên kết chia sẻ. Khi sử dụng điện thoại thật trong cùng mạng, thay localhost bằng địa chỉ IP LAN của máy chạy backend. PUBLIC_URL ở backend cũng cần dùng địa chỉ LAN để các URL ảnh trả về có thể mở từ điện thoại.')
    code(doc, ['EXPO_PUBLIC_API_URL=http://<IP_LAN>:5000', 'EXPO_PUBLIC_WEB_URL=http://<IP_LAN>:5173'])
    p(doc, 'Sau khi thay biến môi trường, khởi động lại Expo và làm sạch cache nếu cần. Lệnh dev:mobile chạy riêng Expo; dev:all chạy đồng thời web, backend và mobile. Cổng của Metro được lấy theo thông báo thực tế khi khởi động, không được giả định cố định trong báo cáo.')
    code(doc, ['npm run dev:mobile', '# Hoac chay dong thoi ca ba thanh phan:', 'npm run dev:all', '# Khi can lam sach cache Expo:', 'npm --prefix mobile run start -- --clear'])
    p(doc, 'Các biến EXPO_PUBLIC_* được đóng vào bundle máy khách nên chỉ chứa địa chỉ công khai, không chứa mật khẩu cơ sở dữ liệu hoặc khóa bí mật. Android Emulator có cơ chế địa chỉ riêng để truy cập máy host; mã client có hỗ trợ host Metro và phương án 10.0.2.2 khi phù hợp. Kiểm thử native vẫn cần môi trường và thiết bị tương ứng; việc mở Expo web trên trình duyệt chỉ xác nhận phạm vi web của ứng dụng.')
    h(doc, '4.3.5. Kiểm tra kiểu và tạo bản đóng gói', 2)
    p(doc, 'Trước khi lưu một bản chạy thử, dự án có các lệnh kiểm tra kiểu, lint và build để phát hiện lỗi phát triển. Web được đóng gói vào dist; backend được biên dịch vào server/dist. Expo cung cấp các lệnh xuất bundle theo nền tảng, được chạy riêng trong thư mục mobile.')
    code(doc, ['npm run typecheck', 'npm run lint', 'npm run build', 'npm --prefix server run build', 'npm --prefix mobile run typecheck', 'npm --prefix mobile run export:web'])
    p(doc, 'Các lệnh export:android và export:ios tạo tài nguyên JavaScript theo nền tảng khi được thực hiện; kết quả đó chưa tương đương APK hoặc IPA đã ký để cài đặt và phát hành. Quy trình ký số, phân phối và nghiệm thu ứng dụng trên thiết bị thật cần được thực hiện ở giai đoạn phát triển tiếp theo. Dockerfile và Compose đã có trong mã nguồn phục vụ chuẩn bị đóng gói, nhưng chương này không sử dụng chúng làm bằng chứng hệ thống đã chạy trên VPS.')

    add_feature_results(doc)
    add_testing(doc)
    h(doc, '4.6. Tổng kết chương', new_page=True)
    p(doc, 'Chương 4 đã trình bày cách tổ chức và cấu hình các thành phần MugenBunko trong môi trường cục bộ, từ cài dependency, kết nối MySQL và khởi tạo tài khoản quản trị đến vận hành web và chuẩn bị ứng dụng Expo. Các chức năng được hiện thực theo các nhóm đọc truyện, sáng tác, tương tác cộng đồng, kiểm duyệt, quản trị và tuyển chọn kiểm duyệt viên; các hình minh họa phản ánh giao diện đã có trong dự án.')
    p(doc, 'Hồ sơ kiểm thử ngày 22/10/2026 ghi nhận 48 ca đã thực thi đạt trong kế hoạch 68 ca, đồng thời còn 20 ca chưa thực thi. Kết quả này cung cấp căn cứ đánh giá những luồng đã được kiểm tra, nhưng phạm vi nghiệm thu cần tiếp tục mở rộng cho email thực, thiết bị native và các điều kiện vận hành ngoài môi trường cục bộ. Các hạn chế và định hướng triển khai VPS, tên miền, HTTPS, giám sát, sao lưu và phát hành ứng dụng được tiếp tục xem xét trong Chương 5.')
    doc.save(OUTPUT)
    validate(doc)


def add_feature_results(doc):
    h(doc, '4.4. Kết quả triển khai các chức năng', new_page=True)
    p(doc, 'Các kết quả dưới đây được mô tả dựa trên mã nguồn hiện có và ảnh giao diện lưu trong các đợt phát triển, kiểm thử cục bộ. Dữ liệu hiển thị trong hình là dữ liệu minh họa hoặc dữ liệu thử nghiệm. Các ảnh không được dùng để khẳng định toàn bộ chức năng đã được nghiệm thu trên môi trường Internet hoặc thiết bị di động native.')
    h(doc, '4.4.1. Đăng ký, đăng nhập và quản lý phiên', 2)
    p(doc, 'Giao diện xác thực hỗ trợ nhập thông tin tài khoản, gửi mã OTP và hoàn tất đăng ký sau khi mã được kiểm tra. Mật khẩu được băm bằng bcrypt trước khi lưu; mã OTP được lưu ở dạng băm cùng thời hạn và số lần thử. Máy chủ kiểm tra tài khoản, mật khẩu và trạng thái hoạt động trước khi xác lập phiên đăng nhập.')
    figure(doc, 'artifacts/auth/register-washi-1440.png', 'Giao diện đăng ký tài khoản và xác thực email trên web')
    p(doc, 'Trên web, phiên được sử dụng qua cookie HttpOnly; trên native, token được quản lý bằng SecureStore và gửi qua cơ chế Bearer. Đăng xuất ghi nhận token bị thu hồi và chấm dứt phiên socket liên quan. Quyền người dùng được đọc lại từ cơ sở dữ liệu khi xác thực, nên thao tác nhạy cảm không chỉ dựa vào trạng thái nút hiển thị trên giao diện.')
    h(doc, '4.4.2. Trang chủ và khám phá tác phẩm', 2, new_page=True)
    p(doc, 'Trang chủ cung cấp khu vực giới thiệu tác phẩm, các nhóm truyện và lối vào hoạt động đọc. Người dùng có thể chuyển đến phần khám phá để tìm tác phẩm theo thông tin tên, tác giả hoặc nhãn và sử dụng các điều kiện lọc phù hợp. Giao diện sử dụng các chế độ nền Washi, Sepia và Charcoal, đồng thời điều chỉnh bố cục theo kích thước hiển thị.')
    figure(doc, 'artifacts/hero/washi-1440.png', 'Trang chủ MugenBunko với giao diện Washi')
    p(doc, 'Phần tiếp tục đọc sử dụng thông tin vị trí đang đọc để đưa người dùng về chương phù hợp. Trạng thái hoàn thành được phân biệt với việc chỉ mở chương cuối. Danh mục API đã có phân trang; ở một số màn hình, client vẫn ghép các trang để phục vụ tìm kiếm và lọc toàn bộ dữ liệu hiện có, nên khả năng đáp ứng khi dữ liệu lớn cần được đánh giá thêm.')
    h(doc, '4.4.3. Trình đọc, tiến độ và tủ sách cá nhân', 2, new_page=True)
    p(doc, 'Trình đọc hiển thị nội dung chương, hỗ trợ chuyển chương và tùy chỉnh cỡ chữ, font, nền đọc, giãn dòng. Vị trí đọc gồm chương, chỉ số đoạn và độ lệch trong đoạn; dữ liệu được lưu cục bộ trước rồi đồng bộ theo tài khoản. Máy chủ so sánh thời điểm cập nhật để hạn chế yêu cầu gửi chậm ghi đè vị trí mới hơn.')
    figure(doc, 'artifacts/upgrades/library.png', 'Tủ sách cá nhân và quản lý kệ sách trong môi trường thử nghiệm')
    p(doc, 'Bookmark và bộ sưu tập được lưu trên server, bao gồm tên kệ, ảnh bìa, danh sách truyện và quyền công khai. Liên kết chia sẻ chỉ xem được khi kệ được cho phép truy cập. Việc hoàn thành chương cần phiên đọc đủ thời gian và đạt điều kiện cuộn; khóa dữ liệu hoàn thành ngăn cùng một chương được cộng XP nhiều lần. Tùy chọn font, nền và giãn dòng hiện được lưu cục bộ, không được xem là toàn bộ thiết lập đã đồng bộ giữa thiết bị.')
    h(doc, '4.4.4. Không gian sáng tác và quản lý chương', 2, new_page=True)
    p(doc, 'Studio hỗ trợ tác giả tạo và sửa thông tin tác phẩm, tổ chức chương theo tập, thay đổi thứ tự và soạn nội dung. Bản thảo đang viết được tách theo tài khoản và chương; cơ chế tự lưu và khôi phục bản nháp giúp giảm mất dữ liệu khi rời trang. Lịch sử phiên bản cho phép xem lại nội dung đã lưu và khôi phục khi cần.')
    figure(doc, 'artifacts/studio/studio-washi-1440.png', 'Studio tác giả với mục lục chương và vùng soạn thảo')
    p(doc, 'Khi tác giả gửi chương mới hoặc chương chưa xuất bản để công bố hay hẹn lịch, backend đưa chương vào trạng thái chờ duyệt. Cập nhật chương đã xuất bản có thể giữ trạng thái published theo luồng hiện tại. Sau khi được duyệt, chương có thể công bố ngay hoặc chờ lịch phát hành. Scheduler kiểm tra mỗi 30 giây để cập nhật những chương đã đến hạn; trạng thái và thời gian được xử lý ở máy chủ.')
    h(doc, '4.4.5. Diễn đàn và tương tác cộng đồng', 2, new_page=True)
    p(doc, 'Diễn đàn cung cấp chức năng tạo bài viết, đính kèm ảnh, thích bài và trao đổi qua bình luận. Các tương tác tại trang truyện gồm bình luận, trả lời, báo cáo và đánh giá sao. Quan hệ theo dõi tác giả và kết bạn được quản lý bằng dữ liệu tài khoản, kết hợp với thông báo và chức năng trò chuyện để hỗ trợ cộng đồng độc giả.')
    figure(doc, 'artifacts/upgrades/forum-washi.png', 'Giao diện diễn đàn và các khu vực nội dung cộng đồng')
    p(doc, 'Tin nhắn được lưu trong MySQL và trao đổi thời gian thực qua Socket.IO sau khi xác thực phiên. Client có khả năng tải lịch sử và phục hồi dữ liệu khi kết nối lại. Các luồng ghi bài, bình luận và đánh giá được kiểm tra dữ liệu và quyền ở backend; báo cáo vi phạm được chuyển đến khu vực kiểm duyệt để xử lý.')
    h(doc, '4.4.6. Không gian kiểm duyệt nội dung', 2, new_page=True)
    p(doc, 'Không gian kiểm duyệt cung cấp danh sách chương chờ duyệt, xem trước nội dung và các thao tác duyệt hoặc trả chương về nháp. Kiểm duyệt viên có thể tìm, lọc và phân trang hàng đợi; việc xem nội dung được thực hiện ngay trong khu vực làm việc để thuận tiện đối chiếu trước khi quyết định.')
    figure(doc, 'artifacts/upgrades/mod-pending-washi-1440.png', 'Hàng chờ chương và thao tác kiểm duyệt trên web')
    p(doc, 'Báo cáo bình luận hiển thị thông tin cần thiết để giữ hoặc xóa đúng nội dung bị phản ánh. Quyền moderator và admin được phân biệt: khu vực tài khoản của moderator phục vụ tra cứu, còn thao tác quản trị tài khoản phải đáp ứng quyền tại API tương ứng. Giao diện có trạng thái xử lý, lỗi tải và thử lại; các hành vi này đã có kịch bản kiểm thử trình duyệt trong hồ sơ kết quả.')
    h(doc, '4.4.7. Quản trị hệ thống', 2, new_page=True)
    p(doc, 'Trang quản trị tập trung các số liệu tổng quan và những công việc cần xử lý, gồm tài khoản, tác phẩm, danh mục, sự kiện và các thiết lập liên quan. Quản trị viên có thể phân quyền, xử lý yêu cầu trở thành tác giả và quản lý dữ liệu nội dung theo các API đã được bảo vệ ở backend.')
    figure(doc, 'artifacts/upgrades/admin-overview.png', 'Trang tổng quan dành cho quản trị viên')
    p(doc, 'Các chức năng quản trị nội dung gồm quản lý thể loại, gộp nhãn, cập nhật quy định cộng đồng và xử lý sự kiện. Bản nháp nội dung sự kiện được giữ khi chuyển tab hoặc lưu thất bại. Riêng lựa chọn banner tiêu điểm hiện lưu trên trình duyệt, chưa phải cấu hình dùng chung cho toàn hệ thống; điểm này cần được phân biệt khi đánh giá kết quả triển khai.')
    h(doc, '4.4.8. Tuyển chọn kiểm duyệt viên', 2, new_page=True)
    p(doc, 'Phân hệ tuyển chọn hỗ trợ quản trị viên tạo kỳ, quy định thể lệ, tiếp nhận đề cử và hồ sơ ứng viên. Người dùng đủ điều kiện có thể bỏ phiếu theo danh sách đã chốt; thành viên hội đồng thực hiện đánh giá theo tiêu chí. Hệ thống lưu kết quả, quyết định bổ nhiệm, khiếu nại và nhật ký để theo dõi quá trình xử lý.')
    figure(doc, 'artifacts/elections/community-results.png', 'Trang kỳ tuyển chọn và công bố kết quả cho cộng đồng')
    p(doc, 'Các pha draft, nominations, voting, review và completed được quản trị viên chuyển qua thao tác nghiệp vụ; hệ thống không tự chuyển toàn bộ pha chỉ theo đồng hồ. Khi chốt, kết quả được lưu thành snapshot để giữ căn cứ công bố. Scheduler xử lý việc hết hạn nhiệm kỳ; giao dịch và khóa dữ liệu được sử dụng tại các bước ghi quan trọng để hạn chế xung đột giữa bỏ phiếu, chốt kỳ và cấp quyền.')
    h(doc, '4.4.9. Kết quả giao diện ứng dụng di động', 2, new_page=True)
    p(doc, 'Ứng dụng Expo tổ chức các màn hình đọc, tủ sách, diễn đàn và hồ sơ cho kích thước điện thoại. Các chức năng dùng lại tài khoản, API và dữ liệu backend; thông tin đã đồng bộ trên server có thể được truy xuất từ web hoặc mobile. Hình dưới minh họa trình đọc và trò chuyện từ bản Expo web, chạy bằng Edge với viewport điện thoại.')
    figure(doc, ['artifacts/mobile/reader.png', 'artifacts/mobile/chat.png'], 'Trình đọc và trò chuyện trên Expo web/Edge ở kích thước điện thoại', max_height=13.8)
    p(doc, 'Các ảnh và log trình duyệt cho phép đánh giá bố cục và luồng sử dụng trong phạm vi Expo web. Những thành phần như bàn phím, bộ chọn ảnh, SecureStore, chuyển nền và khôi phục phiên cần được kiểm tra tiếp trên Android và iOS thật. Vì vậy, kết quả này được trình bày như bước kiểm chứng giao diện và tích hợp API, không phải nghiệm thu ứng dụng native đã phát hành.')


def add_testing(doc):
    record = load('artifacts/test-report/all-test-cases.json')
    statuses = Counter(c['status'] for c in record['cases'])
    assert statuses == {'PASS': 48, 'NOT_RUN': 20}
    h(doc, '4.5. Kiểm thử và đánh giá kết quả triển khai', new_page=True)
    h(doc, '4.5.1. Căn cứ và phương pháp kiểm thử', 2)
    p(doc, 'Phần này sử dụng hồ sơ kiểm thử đã ghi nhận ngày 22/10/2026, do Nguyễn Hoàng Việt thực hiện, với snapshot mã nguồn 2d7d95c được ghi trong báo cáo kiểm thử. Các số liệu phản ánh đợt chạy đó; chúng không thay thế việc kiểm thử hồi quy đối với mọi thay đổi mã nguồn sau thời điểm ghi nhận. Trong quá trình biên soạn chương, việc đối chiếu dựa trên log, metadata mã thoát và các ca kiểm thử lưu trong artifacts/test-report.')
    p(doc, 'Các kiểm thử API sử dụng cơ sở dữ liệu MySQL tạm tách biệt và fixture riêng; SMTP được tắt, OTP thử nghiệm được bật để kiểm tra luồng xác thực có kiểm soát. Kiểm thử trình duyệt web chạy bằng Edge headless. Phần mobile có kiểm thử hàm tiện ích, API và một ca tổng hợp Expo web ở các kích thước điện thoại, máy tính bảng; không coi các thao tác trong một ca tổng hợp là nhiều ca độc lập để tăng số lượng kết quả.')
    table(doc, 'Bảng 4.5. Tổng hợp kết quả trong hồ sơ kiểm thử ngày 22/10/2026', ['Nhóm kiểm thử', 'Đạt', 'Chưa chạy', 'Phạm vi'], [
        ('Backend integration', '25', '0', '5 upgrades, 11 security và 9 hardening.'),
        ('Node web/mobile', '11', '0', '10 ca unit và 1 ca API tổng hợp.'),
        ('Trình duyệt web', '11', '0', '11 nhóm kiểm tra luồng giao diện web.'),
        ('Trình duyệt Expo web', '1', '0', '1 ca tổng hợp giao diện và tích hợp mobile.'),
        ('Kiểm thử thủ công dự kiến', '0', '20', 'Các trường hợp cần bổ sung kiểm chứng.'),
        ('Tổng cộng', '48', '20', '68 ca trong kế hoạch kiểm thử.'),
    ], [4.6, 1.5, 2.2, 7.7])
    p(doc, 'Trong 48 ca đã thực thi, cả 48 ca đạt và không có ca lỗi được ghi nhận trong snapshot này. Tỷ lệ thực thi so với kế hoạch là 48/68, xấp xỉ 70,6%; còn 20 ca giữ trạng thái chưa chạy. Tỷ lệ này thể hiện mức thực hiện kế hoạch kiểm thử, không phải tỷ lệ bao phủ mã nguồn hoặc khẳng định hệ thống không còn lỗi.')
    h(doc, '4.5.2. Một số kịch bản chức năng đã được kiểm chứng', 2)
    table(doc, 'Bảng 4.6. Các kịch bản tiêu biểu có bằng chứng trong hồ sơ kiểm thử', ['Kịch bản', 'Kết quả đã ghi nhận'], [
        ('Đồng bộ vị trí đọc', 'Vị trí theo đoạn được phục hồi giữa các context; cập nhật cũ không ghi đè dữ liệu mới.'),
        ('Hoàn thành chương và XP', 'Dữ liệu hoàn thành ngăn ghi nhận lặp; kiểm thử cập nhật đồng thời chỉ cộng điểm một lần.'),
        ('Kệ sách và chia sẻ', 'Thêm truyện, chia sẻ kệ công khai và thu hồi quyền xem khi chuyển riêng tư.'),
        ('Lịch sử chương', 'Khôi phục phiên bản và giữ lại phần bản thảo chưa lưu trong luồng đã kiểm tra.'),
        ('OTP và phiên đăng nhập', 'Kiểm tra băm mã, thời gian chờ gửi lại, giới hạn số lần thử, chống dùng lại, cookie và thu hồi token khi đăng xuất.'),
        ('Quyền và dữ liệu riêng tư', 'Các yêu cầu vượt quyền hoặc truy cập dữ liệu không thuộc phạm vi bị từ chối.'),
        ('Diễn đàn và quản trị', 'Thử lỗi mạng, thử lại, điều hướng bàn phím và các thao tác quản trị trong kịch bản tự động.'),
        ('Expo web và tích hợp API', 'Ca tổng hợp có kiểm tra dữ liệu backend cho tủ sách, chương, bài viết và tin nhắn.'),
    ], [5.2, 10.8])
    p(doc, 'Các kịch bản trong Bảng 4.6 là nhóm hành vi được đối chiếu với log, không được cộng thêm vào tổng số ca tại Bảng 4.5. Dự án còn có log hồi quy tuyển chọn và những ảnh chụp từ các lần phát triển khác; do một số bộ chạy trùng phạm vi backend hoặc giao diện, các số lượng đó không được cộng trực tiếp vào báo cáo 68 ca.')
    h(doc, '4.5.3. Kết quả kiểm tra mã nguồn và đóng gói', 2)
    p(doc, 'Hồ sơ chất lượng đã lưu ghi nhận mã thoát 0 đối với kiểm tra kiểu web, lint web, build web, kiểm tra kiểu mobile và build backend. Các kết quả này cho biết bước kiểm tra tương ứng đã hoàn thành trong đợt chạy được ghi nhận. Chúng hỗ trợ đánh giá tính nhất quán mã nguồn và khả năng đóng gói, nhưng không thay thế kiểm thử nghiệp vụ hoặc kiểm tra môi trường vận hành.')
    table(doc, 'Bảng 4.7. Các bước kiểm tra chất lượng đã được lưu kết quả', ['Lệnh kiểm tra', 'Kết quả ghi nhận'], [
        ('npm run typecheck', 'Mã thoát 0.'),
        ('npm run lint', 'Mã thoát 0.'),
        ('npm run build', 'Mã thoát 0; có cảnh báo về kích thước chunk.'),
        ('npm --prefix mobile run typecheck', 'Mã thoát 0.'),
        ('npm --prefix server run build', 'Mã thoát 0.'),
    ], [9.0, 7.0])
    p(doc, 'Build web có cảnh báo chunk lớn hơn ngưỡng 500 kB trong hồ sơ hiện có. Cảnh báo này là căn cứ xem xét cách chia gói và tải tài nguyên khi tiếp tục tối ưu. Việc ghi nhận build thành công không được sử dụng để suy ra thời gian tải hoặc số người dùng đồng thời có thể phục vụ; các kết luận hiệu năng cần phép đo và điều kiện thử nghiệm riêng.')
    h(doc, '4.5.4. Những nội dung cần tiếp tục kiểm chứng', 2)
    p(doc, 'Các ca thủ công chưa thực thi được giữ nguyên trạng thái để thể hiện trung thực phạm vi đánh giá. Những nhóm còn cần kiểm chứng gồm email thực, các tình huống biên trên giao diện, scheduler theo thời gian thực, thông báo nghiệp vụ và các tương tác phụ thuộc thiết bị. Việc đã có mã nguồn hoặc ảnh giao diện không tự chứng minh một ca kiểm thử đã đạt.')
    table(doc, 'Bảng 4.8. Phạm vi cần bổ sung kiểm chứng sau đợt kiểm thử', ['Nhóm nội dung', 'Công việc còn cần thực hiện'], [
        ('Email thực', 'Kiểm tra SMTP, nhận thư trong hộp thư, gửi lại và hết hạn mã trên luồng sử dụng thực.'),
        ('Các tình huống biên UI', 'Đối chiếu đầy đủ tìm kiếm/lọc, mục lục, đánh giá, bài có ảnh và thay đổi quyền bình luận.'),
        ('Nghiệp vụ theo thời gian', 'Kiểm tra lịch phát hành và thông báo theo các mốc thời gian thực tế.'),
        ('Android và iOS native', 'Kiểm tra bàn phím, editor, picker ảnh, kho token, chuyển nền, safe area và kết nối lại.'),
        ('Môi trường vận hành online', 'Kiểm chứng VPS, HTTPS, tên miền, giám sát, sao lưu và tải thực tế ở giai đoạn triển khai tiếp theo.'),
    ], [5.0, 11.0])
    p(doc, 'Trong phạm vi đã kiểm tra, các luồng trọng tâm như đọc và đồng bộ, tủ sách, bảo toàn bản thảo, quyền truy cập và thao tác quản trị có bằng chứng hoạt động. Hệ thống vẫn cần hoàn thiện kiểm thử chưa chạy và kiểm chứng trên môi trường đích trước khi đánh giá khả năng phục vụ công khai. Kết quả triển khai cục bộ là cơ sở cho bước hoàn thiện và định hướng vận hành ở chương tiếp theo.')


def validate(doc):
    text = '\n'.join([par.text for par in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert text.startswith('CHƯƠNG 4: TRIỂN KHAI HỆ THỐNG')
    assert '\ufffd' not in text
    assert len(doc.sections) == 1
    assert len(doc.inline_shapes) == 10
    assert len(doc.tables) == 8
    assert all(f'4.{i}. ' in text for i in range(1, 7))
    assert len(re.findall(r'^Hình 4\.\d+\.', text, re.M)) == 9
    with ZipFile(OUTPUT) as z:
        media = [n for n in z.namelist() if n.startswith('word/media/')]
        assert len(media) == 10
        hashes = {hashlib.sha256(z.read(n)).hexdigest() for n in media}
        assert hashes == {f['sha256'] for f in FIGURES}
    report = {
        'document': str(OUTPUT), 'title': doc.core_properties.title,
        'body_font': 'Times New Roman 13 pt', 'table_and_caption_font': 'Times New Roman 12 pt',
        'code_font': 'Consolas 10 pt', 'line_spacing_body': 1.5,
        'paper': 'A4 portrait', 'margins_cm': {'left': 3, 'right': 2, 'top': 2, 'bottom': 2},
        'main_sections': 6, 'word_tables': len(doc.tables), 'figure_captions': 9,
        'embedded_images': 10, 'images_byte_identical_to_sources': True,
        'words_whitespace': len(text.split()), 'characters': len(text),
        'test_snapshot': {'date': '2026-10-22', 'commit': '2d7d95c', 'PASS': 48, 'NOT_RUN': 20},
        'scope': 'Local implementation; mobile screenshots are Expo web; no VPS/native-release claim',
        'source_files': SOURCES, 'figures': FIGURES,
    }
    (OUT / 'validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    (OUT / 'chapter4_text.txt').write_text(text, encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k not in ('figures','source_files')}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    build()
