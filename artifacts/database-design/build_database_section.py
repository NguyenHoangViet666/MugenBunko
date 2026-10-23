"""Build section 3.6 from the project's DDL, retaining the supplied ERD pixels."""
from __future__ import annotations

from collections import OrderedDict
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import os
import re
import shutil
import sys

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, str(Path(os.environ['TEMP']) / 'mugenbunko-test-report-libs'))
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
OUTPUT = OUT / 'Muc_3_6_Thiet_ke_co_so_du_lieu_MugenBunko.docx'
ERD = OUT / 'ERD_MugenBunko_original.png'
ERD_WORD = OUT / 'ERD_MugenBunko_word.png'
FONT = 'Times New Roman'


def split_sql(text):
    parts, start, depth, quote = [], 0, 0, None
    for i, c in enumerate(text):
        if quote:
            if c == quote and (i == 0 or text[i-1] != '\\'):
                quote = None
        elif c in "'\"`":
            quote = c
        elif c == '(':
            depth += 1
        elif c == ')':
            depth -= 1
        elif c == ',' and depth == 0:
            parts.append(text[start:i].strip())
            start = i+1
    parts.append(text[start:].strip())
    return [p for p in parts if p]


def names(text):
    return [x.strip().strip('`').split()[0] for x in split_sql(text)]


def read_model():
    model = OrderedDict()
    for source in ('server/schema.sql', 'server/db.ts', 'server/readerSchema.ts', 'server/electionSchema.ts'):
        text = (ROOT / source).read_text(encoding='utf-8').replace('\\`', '`')
        text = re.sub(r'--[^\n]*', '', text)
        for m in re.finditer(r'CREATE TABLE IF NOT EXISTS\s+`?(\w+)`?\s*\((.*?)\)\s*ENGINE\s*=\s*InnoDB', text, re.S | re.I):
            table, body = m.groups()
            if table in model:
                continue
            fields, primary, unique, fks, indexes = [], [], [], {}, []
            for part in split_sql(body):
                pk = re.match(r'PRIMARY KEY\s*\((.*?)\)', part, re.S | re.I)
                uq = re.match(r'UNIQUE(?:\s+KEY)?(?:\s+`?\w+`?)?\s*\((.*?)\)', part, re.S | re.I)
                fk = re.match(r'FOREIGN KEY\s*\((.*?)\)\s*REFERENCES\s*`?(\w+)`?\s*\((.*?)\)(.*)', part, re.S | re.I)
                ix = re.match(r'(?:INDEX|KEY)\s+`?(\w+)`?\s*\((.*?)\)', part, re.S | re.I)
                if pk:
                    primary = names(pk.group(1))
                elif uq:
                    unique.append(names(uq.group(1)))
                elif fk:
                    local, target, remote, policy = fk.groups()
                    for a, b in zip(names(local), names(remote)):
                        fks[a] = {'target': f'{target}.{b}', 'delete': 'CASCADE' if 'CASCADE' in policy else 'SET NULL' if 'SET NULL' in policy else 'RESTRICT'}
                elif ix:
                    indexes.append({'name': ix.group(1), 'fields': ix.group(2).replace('`', '')})
                else:
                    col = re.match(r'`?(\w+)`?\s+([A-Za-z]+(?:\([^)]*\))?)(.*)', part, re.S)
                    if not col:
                        raise ValueError((table, part))
                    name, typ, tail = col.groups()
                    if name.upper() in ('CHECK', 'CONSTRAINT'):
                        continue
                    if 'PRIMARY KEY' in tail.upper():
                        primary.append(name)
                    if re.search(r'\bUNIQUE\b', tail, re.I):
                        unique.append([name])
                    fields.append({'name': name, 'type': typ.upper(), 'tail': tail.strip()})
            model[table] = {'fields': fields, 'primary': primary, 'unique': unique, 'fks': fks, 'indexes': indexes, 'source': source}
    # Additive migrations in initializeDbConnection are part of the implemented model.
    users = model['users']['fields']
    pos = next(i for i, c in enumerate(users) if c['name'] == 'displayname') + 1
    users.insert(pos, {'name': 'email', 'type': 'VARCHAR(150)', 'tail': 'NULL UNIQUE'})
    model['users']['unique'].append(['email'])
    model['events']['fields'].extend([
        {'name': 'status', 'type': 'VARCHAR(50)', 'tail': "NOT NULL DEFAULT 'draft'"},
        {'name': 'content', 'type': 'TEXT', 'tail': 'NULL DEFAULT NULL'},
    ])
    assert len(model) == 37, list(model)
    return model


PURPOSE = {
    'users': 'Lưu tài khoản, hồ sơ cá nhân, thông tin đăng nhập và trạng thái người dùng.',
    'user_roles': 'Gán một hoặc nhiều vai trò cho tài khoản; vai trò được tách khỏi hồ sơ người dùng.',
    'email_otps': 'Lưu trạng thái xác thực email bằng mã dùng một lần, thời hạn và số lần nhập sai.',
    'revoked_tokens': 'Lưu dấu băm token đã thu hồi để từ chối sử dụng lại token còn hạn.',
    'novels': 'Lưu thông tin chung của tác phẩm, tác giả và các số liệu phục vụ hiển thị danh mục.',
    'chapters': 'Lưu nội dung và thứ tự chương, trạng thái xuất bản cùng thời điểm phát hành theo lịch.',
    'chapter_versions': 'Lưu các bản chụp nội dung chương, phục vụ tra cứu lịch sử và khôi phục bản nháp.',
    'novel_tags': 'Lưu tập nhãn gắn với truyện; một nhãn chỉ xuất hiện một lần trong cùng tác phẩm.',
    'genres': 'Lưu danh mục tên thể loại do quản trị viên quản lý.',
    'deleted_novels': 'Giữ mã truyện, tiêu đề và lý do gỡ bỏ sau khi bản ghi truyện gốc bị xóa.',
    'bookmarks': 'Ghi nhận truyện được người dùng đánh dấu vào tủ sách.',
    'library_collections': 'Lưu kệ sách do người dùng tạo, chế độ công khai và mã chia sẻ.',
    'collection_novels': 'Liên kết các truyện với kệ sách; một truyện có thể nằm trên nhiều kệ.',
    'reading_progress': 'Lưu vị trí đọc mới nhất của một tài khoản đối với từng truyện để tiếp tục đọc trên các thiết bị.',
    'chapter_reads': 'Ghi nhận chương đã hoàn thành; khóa ghép ngăn ghi nhận lặp cùng một chương cho một người dùng.',
    'reading_sessions': 'Lưu phiên đọc hiện hành theo người dùng và chương để kiểm tra điều kiện hoàn thành.',
    'read_cooldowns': 'Lưu lần tính lượt đọc gần nhất theo địa chỉ IP và truyện để hạn chế đếm lặp.',
    'comments': 'Lưu bình luận ở trang truyện và quan hệ trả lời giữa các bình luận.',
    'reviews': 'Lưu đánh giá bằng số sao và nội dung nhận xét; mỗi người dùng có tối đa một đánh giá cho mỗi truyện.',
    'reports': 'Lưu báo cáo vi phạm đối với bình luận truyện, người báo cáo và lý do.',
    'forum_posts': 'Lưu bài viết diễn đàn, chuyên mục, ảnh đính kèm và thiết lập hạn chế bình luận.',
    'forum_comments': 'Lưu bình luận và phản hồi trong bài viết diễn đàn.',
    'forum_post_likes': 'Ghi nhận lượt thích bài viết của người dùng; mỗi cặp người dùng và bài viết là duy nhất.',
    'follows': 'Lưu quan hệ độc giả theo dõi tác giả; cả hai phía đều tham chiếu đến tài khoản.',
    'friends': 'Lưu quan hệ kết bạn giữa hai tài khoản và trạng thái xử lý lời mời.',
    'messages': 'Lưu tin nhắn trực tiếp giữa người gửi và người nhận cùng trạng thái đã đọc.',
    'notifications': 'Lưu thông báo gửi đến từng tài khoản và trạng thái đã đọc.',
    'announcements': 'Lưu thông báo riêng của một tác phẩm.',
    'events': 'Lưu sự kiện chung của hệ thống, nội dung chi tiết và trạng thái phát hành.',
    'rules': 'Lưu quy định cộng đồng và thứ tự hiển thị các điều khoản.',
    'elections': 'Lưu đợt tuyển chọn kiểm duyệt viên, thể lệ, thời hạn và kết quả đã chốt.',
    'election_candidates': 'Lưu hồ sơ ứng viên, kết quả xét duyệt và các mốc thử việc, nhiệm kỳ.',
    'election_voters': 'Chốt danh sách tài khoản đủ điều kiện bỏ phiếu hoặc đánh giá trong từng kỳ.',
    'election_ballots': 'Lưu lựa chọn của cử tri đối với ứng viên trong một kỳ tuyển chọn.',
    'election_reviews': 'Lưu nhận xét của thành viên hội đồng theo tiêu chí cho từng ứng viên.',
    'election_complaints': 'Lưu khiếu nại liên quan đến kỳ tuyển chọn và kết quả xử lý.',
    'election_audit': 'Lưu nhật ký thao tác của phân hệ tuyển chọn để truy vết quá trình xử lý.',
}

COMMON = {
    'id': 'Mã định danh bản ghi.', 'user_id': 'Mã người dùng.', 'novel_id': 'Mã truyện.',
    'chapter_id': 'Mã chương truyện.', 'author_id': 'Mã tác giả.', 'post_id': 'Mã bài viết diễn đàn.',
    'election_id': 'Mã kỳ tuyển chọn.', 'candidate_id': 'Mã hồ sơ ứng viên.',
    'created_at': 'Thời điểm tạo bản ghi.', 'updated_at': 'Thời điểm cập nhật gần nhất.',
    'title': 'Tiêu đề.', 'description': 'Thông tin mô tả.', 'content': 'Nội dung chi tiết.',
    'status': 'Trạng thái xử lý của bản ghi.', 'text': 'Nội dung văn bản.',
    'parent_id': 'Mã bình luận cha; NULL đối với bình luận gốc.',
    'is_read': 'Đã đọc hay chưa (0/1).',
}

DETAIL = {
    'users': {
        'username': 'Tên đăng nhập duy nhất.', 'password': 'Chuỗi băm mật khẩu bằng bcrypt.',
        'displayname': 'Tên hiển thị.', 'email': 'Địa chỉ email; cho phép NULL đối với tài khoản chưa khai báo.',
        'coins': 'Số xu của tài khoản.', 'level': 'Cấp độ tài khoản.', 'xp': 'Điểm kinh nghiệm tích lũy.',
        'bio': 'Giới thiệu cá nhân.', 'avatar_seed': 'Giá trị tạo hoặc chọn ảnh đại diện.',
        'status': 'Trạng thái tài khoản, ví dụ active hoặc suspended.',
        'author_request': 'Trạng thái yêu cầu trở thành tác giả.', 'date_joined': 'Thời điểm tham gia hệ thống.',
    },
    'user_roles': {'role': 'Vai trò reader, author, moderator hoặc admin.'},
    'email_otps': {
        'email': 'Email đang được xác thực.', 'code_hash': 'Dấu băm HMAC-SHA256 của email và mã OTP.',
        'expires_at': 'Thời điểm hết hạn, tính bằng mili giây Unix.',
        'sent_at': 'Thời điểm gửi mã gần nhất, tính bằng mili giây Unix.',
        'attempts': 'Số lần nhập sai mã xác thực.',
    },
    'revoked_tokens': {'token_hash': 'Dấu băm của token bị thu hồi.', 'expires_at': 'Thời điểm token hết hạn, tính bằng mili giây Unix.'},
    'novels': {
        'cover': 'Đường dẫn hoặc URL ảnh bìa.', 'summary': 'Tóm tắt nội dung truyện.',
        'genre': 'Chuỗi thể loại của truyện; chưa dùng khóa ngoại tới genres.',
        'status': 'Trạng thái tác phẩm, ví dụ active hoặc suspended.',
        'type': 'Định dạng truyện, ví dụ series hoặc oneshot.',
        'reads': 'Số lượt đọc đã ghi nhận.', 'rating': 'Điểm đánh giá tổng hợp.',
        'bookmarks_count': 'Số lượt đánh dấu, được lưu để phục vụ hiển thị.',
    },
    'chapters': {
        'volume_name': 'Tên tập hoặc quyển.', 'content': 'Toàn bộ nội dung chương.',
        'status': 'Trạng thái vòng đời chương: bản nháp, chờ duyệt, công bố hoặc hẹn lịch tùy thao tác nghiệp vụ.',
        'scheduled_release': 'Thời điểm phát hành theo lịch; NULL nếu không hẹn lịch.',
        'word_count': 'Số từ của chương.', 'sort_order': 'Thứ tự chương trong tác phẩm.',
    },
    'chapter_versions': {'editor_id': 'Mã người chỉnh sửa; cho phép NULL khi tài khoản bị xóa.', 'volume_name': 'Tên tập tại thời điểm lưu phiên bản.', 'content': 'Bản chụp nội dung chương.', 'word_count': 'Số từ của phiên bản.', 'source': 'Nguồn tạo phiên bản, ví dụ save hoặc thao tác khôi phục.'},
    'novel_tags': {'tag': 'Nhãn mô tả tác phẩm.'},
    'genres': {'name': 'Tên thể loại duy nhất.'},
    'deleted_novels': {'id': 'Mã truyện trước khi xóa; không liên kết khóa ngoại với novels.', 'reason': 'Lý do gỡ tác phẩm.', 'deleted_at': 'Thời điểm gỡ tác phẩm.'},
    'library_collections': {'name': 'Tên kệ sách.', 'description': 'Mô tả ngắn về kệ sách.', 'cover_novel_id': 'Truyện được chọn làm ảnh bìa; NULL nếu chưa chọn hoặc truyện bị xóa.', 'is_public': 'Cho phép truy cập công khai (0/1).', 'share_id': 'Mã chia sẻ duy nhất; quyền xem còn phụ thuộc is_public.'},
    'collection_novels': {'collection_id': 'Mã kệ sách.'},
    'reading_progress': {'scroll_ratio': 'Tỷ lệ vị trí cuộn trong chương, từ 0 đến 1.', 'paragraph_index': 'Chỉ số đoạn đang đọc, bắt đầu từ 0.', 'paragraph_offset': 'Tỷ lệ vị trí trong đoạn, từ 0 đến 1.', 'client_updated_at': 'Thời gian cập nhật ở thiết bị, tính bằng mili giây Unix.', 'updated_at': 'Thời điểm máy chủ cập nhật, độ chính xác mili giây.'},
    'chapter_reads': {'completed_at': 'Thời điểm hoàn thành chương.'},
    'reading_sessions': {'session_id': 'Mã phiên đọc ngẫu nhiên gồm 32 ký tự.', 'started_at': 'Thời điểm bắt đầu phiên đọc.'},
    'read_cooldowns': {'ip_address': 'Địa chỉ IPv4 hoặc IPv6 của lượt truy cập.', 'last_read_at': 'Thời điểm tính lượt đọc gần nhất.'},
    'reviews': {'stars': 'Số sao đánh giá; ràng buộc CHECK từ 1 đến 5.', 'text': 'Nội dung nhận xét tác phẩm.'},
    'reports': {'reporter_id': 'Mã người gửi báo cáo.', 'comment_id': 'Mã bình luận bị báo cáo.', 'reason': 'Lý do báo cáo vi phạm.'},
    'forum_posts': {'category': 'Chuyên mục bài viết.', 'restrict_comments': 'Cờ hạn chế bình luận (0/1), được API kiểm tra.', 'image_url': 'URL ảnh đính kèm; có thể để trống.'},
    'follows': {'follower_id': 'Mã người theo dõi.', 'author_id': 'Mã tài khoản được theo dõi.'},
    'friends': {'user_id_1': 'Tài khoản thứ nhất trong quan hệ kết bạn.', 'user_id_2': 'Tài khoản thứ hai trong quan hệ kết bạn.', 'status': 'Trạng thái lời mời, ví dụ pending hoặc accepted.'},
    'messages': {'sender_id': 'Mã người gửi.', 'receiver_id': 'Mã người nhận.', 'message_text': 'Nội dung tin nhắn.'},
    'events': {'status': 'Trạng thái sự kiện; mặc định draft.'},
    'rules': {'order_index': 'Thứ tự hiển thị điều khoản.', 'updated_at': 'Thời điểm sửa điều khoản gần nhất.'},
    'elections': {'rules': 'Thể lệ và tham số kỳ tuyển chọn ở dạng JSON.', 'nomination_closes_at': 'Hạn kết thúc đề cử, tính bằng mili giây Unix.', 'voting_closes_at': 'Hạn kết thúc bỏ phiếu, tính bằng mili giây Unix.', 'result_snapshot': 'Kết quả được chốt tại thời điểm đóng kỳ, dạng JSON.', 'created_by': 'Mã người tạo kỳ tuyển chọn.', 'published_at': 'Thời điểm công bố kết quả.'},
    'election_candidates': {
        'user_id': 'Mã tài khoản ứng viên.', 'nominated_by': 'Mã người đề cử.',
        'motivation': 'Động lực tham gia ứng tuyển.', 'experience': 'Kinh nghiệm của ứng viên.',
        'availability': 'Khả năng bố trí thời gian tham gia.', 'scenario_answer': 'Câu trả lời tình huống kiểm duyệt.',
        'supplement': 'Nội dung bổ sung hồ sơ.', 'verification_reason': 'Lý do kiểm tra hoặc xác minh hồ sơ.',
        'decision': 'Kết quả xét duyệt hoặc bổ nhiệm; mặc định pending.', 'public_reason': 'Giải trình có thể công khai.',
        'internal_note': 'Ghi chú nội bộ của quản trị viên.', 'override_reason': 'Lý do quyết định ngoại lệ.',
        'decided_by': 'Mã người quyết định; tham chiếu nghiệp vụ, chưa khai báo FK.',
        'probation_ends_at': 'Hạn kết thúc thử việc, tính bằng mili giây Unix.',
        'term_ends_at': 'Hạn kết thúc nhiệm kỳ, tính bằng mili giây Unix.',
        'confirmed_at': 'Thời điểm xác nhận chính thức, tính bằng mili giây Unix.',
        'ended_at': 'Thời điểm kết thúc đảm nhiệm, tính bằng mili giây Unix.', 'end_reason': 'Lý do kết thúc đảm nhiệm.',
    },
    'election_voters': {'can_vote': 'Đủ điều kiện bỏ phiếu trong kỳ (0/1).', 'can_review': 'Đủ điều kiện tham gia hội đồng đánh giá (0/1).'},
    'election_ballots': {'voter_id': 'Mã cử tri bỏ phiếu.'},
    'election_reviews': {'reviewer_id': 'Mã thành viên đánh giá.', 'verdict': 'Kết luận đánh giá của thành viên hội đồng.', 'criteria': 'Kết quả theo từng tiêu chí ở dạng JSON.', 'comment': 'Nhận xét chi tiết.'},
    'election_complaints': {'message': 'Nội dung khiếu nại.', 'resolution': 'Kết quả và giải trình xử lý khiếu nại.', 'resolved_by': 'Mã người xử lý; tham chiếu nghiệp vụ, chưa khai báo FK.'},
    'election_audit': {'actor_id': 'Mã chủ thể thao tác; tham chiếu nghiệp vụ, chưa khai báo FK.', 'action': 'Tên hoặc mã hành động được ghi nhận.', 'details': 'Chi tiết thao tác ở dạng JSON.'},
}

GROUPS = [
    ('a) Tài khoản và xác thực', ['users', 'user_roles', 'email_otps', 'revoked_tokens']),
    ('b) Tác phẩm và xuất bản', ['novels', 'chapters', 'chapter_versions', 'novel_tags', 'genres', 'deleted_novels']),
    ('c) Tủ sách và hoạt động đọc', ['bookmarks', 'library_collections', 'collection_novels', 'reading_progress', 'chapter_reads', 'reading_sessions', 'read_cooldowns']),
    ('d) Tương tác và cộng đồng', ['comments', 'reviews', 'reports', 'forum_posts', 'forum_comments', 'forum_post_likes', 'follows', 'friends', 'messages']),
    ('e) Thông báo và quản trị nội dung', ['notifications', 'announcements', 'events', 'rules']),
    ('f) Tuyển chọn kiểm duyệt viên', ['elections', 'election_candidates', 'election_voters', 'election_ballots', 'election_reviews', 'election_complaints', 'election_audit']),
]


def font(run, size=13, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    fonts.attrib.clear()
    for k in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        fonts.set(qn(f'w:{k}'), FONT)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'vi-VN')
    rpr.append(lang)


def para(doc, text, style=None, *, indent=True, size=13, italic=False):
    p = doc.add_paragraph(style=style)
    font(p.add_run(text), size, italic=italic)
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    return p


def heading(doc, text, level=2):
    p = para(doc, text, f'Heading {level}', indent=False)
    return p


def caption(doc, text):
    p = para(doc, text, 'Caption', indent=False, size=12, italic=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths):
        col.width = Cm(w)
    for c, w, s in zip(t.rows[0].cells, widths, headers):
        c.width = Cm(w)
        c.text = s
    for row in rows:
        cells = t.add_row().cells
        for c, w, s in zip(cells, widths, row):
            c.width = Cm(w)
            c.text = str(s)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        for key, value in [('val', 'single'), ('sz', '4'), ('color', '8C8C8C')]:
            el.set(qn(f'w:{key}'), value)
        borders.append(el)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for key, value in [('top', '55'), ('bottom', '55'), ('left', '75'), ('right', '75')]:
        el = OxmlElement(f'w:{key}')
        el.set(qn('w:w'), value)
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    props.append(margins)
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E9EDF1')
                c._tc.get_or_add_tcPr().append(shade)
            for p in c.paragraphs:
                pf = p.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.line_spacing = 1.05
                pf.space_before = Pt(0)
                pf.space_after = Pt(1)
                pf.widow_control = True
                pf.keep_with_next = i == 0
                for run in p.runs:
                    font(run, 12, bold=i == 0)
    return t


def section_page(sec, landscape=False):
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width = Cm(29.7 if landscape else 21)
    sec.page_height = Cm(21 if landscape else 29.7)
    sec.left_margin = Cm(2 if landscape else 3)
    sec.right_margin = Cm(1.5 if landscape else 2)
    sec.top_margin = Cm(1.35 if landscape else 2)
    sec.bottom_margin = Cm(1.35 if landscape else 2)
    sec.header_distance = Cm(0.5 if landscape else 0.7)
    sec.footer_distance = Cm(0.55 if landscape else 0.8)


def setup(doc):
    section_page(doc.sections[0])
    for name in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4', 'Caption', 'Footer'):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(13)
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        rf = OxmlElement('w:rFonts')
        for k in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
            rf.set(qn(f'w:{k}'), FONT)
        rpr.append(rf)
    normal = doc.styles['Normal'].paragraph_format
    normal.first_line_indent = Cm(1)
    normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.line_spacing = 1.5
    normal.space_before = Pt(0)
    normal.space_after = Pt(6)
    normal.widow_control = True
    for name in ('Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'):
        style = doc.styles[name]
        style.font.bold = True
        pf = style.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.line_spacing = 1.15
        pf.space_before = Pt(10)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        pf.keep_together = True
    doc.styles['Heading 1'].font.size = Pt(14)
    cap = doc.styles['Caption'].paragraph_format
    cap.first_line_indent = Cm(0)
    cap.line_spacing = 1.1
    cap.space_before = Pt(5)
    cap.space_after = Pt(7)
    cap.keep_with_next = False
    cap.keep_together = True
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    footer.paragraph_format.line_spacing = 1
    footer.paragraph_format.space_after = Pt(0)
    font(footer.add_run('Trang '), 12)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    doc.settings.element.append(update)
    # Preserve the screenshot at its original quality in the Word package.
    no_compress = OxmlElement('w:doNotAutoCompressPictures')
    doc.settings.element.append(no_compress)
    doc.core_properties.title = '3.6. Thiết kế cơ sở dữ liệu — MugenBunko'
    doc.core_properties.subject = 'Thiết kế cơ sở dữ liệu quan hệ cho hệ thống đọc và xuất bản truyện MugenBunko'
    doc.core_properties.author = 'Nguyễn Hoàng Việt'
    doc.core_properties.keywords = 'MugenBunko; thiết kế cơ sở dữ liệu; ERD; MySQL; từ điển dữ liệu'


def field_explanation(tname, field, t):
    name, tail = field['name'], field['tail']
    markers = []
    if name in t['primary']:
        markers.append('PK')
    if 'AUTO_INCREMENT' in tail.upper():
        markers.append('AI')
    if 'NOT NULL' in tail.upper() or name in t['primary']:
        markers.append('NN')
    if [name] in t['unique']:
        markers.append('UQ')
    if name in t['fks']:
        markers.append('FK → ' + t['fks'][name]['target'])
    elif 'NOT NULL' not in tail.upper() and name not in t['primary']:
        markers.append('NULL')
    desc = DETAIL.get(tname, {}).get(name, COMMON.get(name))
    if desc is None:
        raise ValueError(f'Missing description: {tname}.{name}')
    return '; '.join(markers) + '. ' + desc


def build():
    model = read_model()
    doc = Document()
    setup(doc)
    heading(doc, '3.6. Thiết kế cơ sở dữ liệu', 1)
    heading(doc, '3.6.1. Tổng quan và nguyên tắc thiết kế', 2)
    para(doc, 'Cơ sở dữ liệu của MugenBunko được tổ chức theo mô hình quan hệ, phục vụ đồng thời ứng dụng web và ứng dụng di động thông qua hệ thống API. Dữ liệu cần quản lý bao gồm tài khoản, tác phẩm, chương truyện, tiến độ đọc, tủ sách cá nhân, hoạt động cộng đồng và nghiệp vụ tuyển chọn kiểm duyệt viên.')
    para(doc, 'Hệ thống sử dụng MySQL với bộ máy lưu trữ InnoDB. Bộ ký tự utf8mb4 và đối chiếu utf8mb4_unicode_ci hỗ trợ lưu trữ tiếng Việt, tiếng Nhật và các ký tự Unicode. Các thực thể được nhận diện bằng khóa chính; quan hệ một–nhiều được biểu diễn bằng khóa ngoại, còn quan hệ nhiều–nhiều được tách thành bảng liên kết có khóa ghép.')
    para(doc, 'Lược đồ hiện tại gồm 37 bảng, chia thành sáu nhóm nghiệp vụ. Thiết kế ưu tiên tách dữ liệu theo chức năng, hạn chế bản ghi trùng và bảo đảm khả năng truy vết những thay đổi quan trọng. Một số số liệu tổng hợp được lưu trực tiếp để phục vụ truy vấn hiển thị.')
    caption(doc, 'Bảng 3.6.1. Các nhóm dữ liệu của hệ thống MugenBunko')
    table(doc, ['Nhóm nghiệp vụ', 'Số bảng', 'Dữ liệu quản lý'], [
        ('Tài khoản và xác thực', '4', 'Tài khoản, vai trò, OTP email, token thu hồi.'),
        ('Tác phẩm và xuất bản', '6', 'Truyện, chương, phiên bản, nhãn, thể loại, truyện đã gỡ.'),
        ('Tủ sách và hoạt động đọc', '7', 'Đánh dấu, kệ sách, vị trí đọc, phiên đọc, chương hoàn thành.'),
        ('Tương tác và cộng đồng', '9', 'Bình luận, đánh giá, diễn đàn, theo dõi, kết bạn, tin nhắn.'),
        ('Thông báo và quản trị', '4', 'Thông báo cá nhân, thông báo truyện, sự kiện, quy định.'),
        ('Tuyển chọn kiểm duyệt viên', '7', 'Kỳ tuyển chọn, ứng viên, cử tri, phiếu, đánh giá, khiếu nại, nhật ký.'),
    ], [4.4, 1.5, 10.1])

    # Keep a uniform A4 portrait layout for predictable Word pagination.
    erd_heading = heading(doc, '3.6.2. Sơ đồ thực thể – liên kết (ERD)', 2)
    erd_heading.paragraph_format.space_before = Pt(0)
    erd_heading.paragraph_format.page_break_before = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    p._p.get_or_add_pPr().append(snap)
    pic = p.add_run().add_picture(str(ERD_WORD), width=Cm(16))
    pic._inline.docPr.set('descr', 'Sơ đồ ERD MugenBunko do người dùng cung cấp; giữ nguyên ảnh 1451 × 888 pixel.')
    erd_caption = caption(doc, 'Hình 3.6.1. Sơ đồ thực thể – liên kết của cơ sở dữ liệu MugenBunko')
    erd_caption.paragraph_format.space_after = Pt(0)

    relations_heading = heading(doc, '3.6.3. Mô hình quan hệ và các liên kết chính', 2)
    relations_heading.paragraph_format.page_break_before = True
    para(doc, 'Trong mô hình, users và novels là hai thực thể trung tâm. Một tài khoản có thể sáng tác nhiều truyện; mỗi truyện thuộc một tác giả và có nhiều chương. Các bảng phục vụ đọc, tương tác và quản trị liên kết với những thực thể này để bảo đảm dữ liệu được xác định đúng đối tượng.')
    caption(doc, 'Bảng 3.6.2. Các quan hệ tiêu biểu trong cơ sở dữ liệu')
    table(doc, ['Quan hệ', 'Bội số', 'Cách biểu diễn'], [
        ('Tài khoản – vai trò', '1–N', 'user_roles.user_id → users.id; khóa ghép (user_id, role).'),
        ('Tác giả – truyện', '1–N', 'novels.author_id → users.id.'),
        ('Truyện – chương', '1–N', 'chapters.novel_id → novels.id.'),
        ('Chương – phiên bản', '1–N', 'chapter_versions.chapter_id → chapters.id.'),
        ('Người dùng – truyện đánh dấu', 'N–N', 'Bảng bookmarks có khóa ghép (user_id, novel_id).'),
        ('Kệ sách – truyện', 'N–N', 'Bảng collection_novels có khóa ghép (collection_id, novel_id).'),
        ('Người dùng – vị trí đọc', '1–N', 'reading_progress có khóa ghép (user_id, novel_id); tham chiếu thêm chapter_id.'),
        ('Truyện – bình luận/đánh giá', '1–N', 'comments.novel_id và reviews.novel_id → novels.id.'),
        ('Bình luận – phản hồi', '1–N', 'comments.parent_id tự tham chiếu comments.id; diễn đàn dùng cơ chế tương tự.'),
        ('Bài diễn đàn – bình luận', '1–N', 'forum_comments.post_id → forum_posts.id.'),
        ('Người dùng – người dùng', 'N–N', 'follows và friends chứa hai khóa ngoại cùng tham chiếu users.id.'),
        ('Kỳ tuyển chọn – ứng viên', '1–N', 'election_candidates.election_id → elections.id.'),
        ('Cử tri – ứng viên theo kỳ', 'N–N', 'election_ballots liên kết kỳ tuyển chọn, cử tri và hồ sơ ứng viên.'),
    ], [4.4, 1.5, 10.1])
    para(doc, 'Các danh mục hoặc dữ liệu lưu vết không nhất thiết liên kết trực tiếp với bảng trung tâm. Cụ thể, genres lưu danh mục thể loại nhưng novels.genre hiện vẫn là chuỗi; deleted_novels giữ thông tin truyện đã gỡ và không có khóa ngoại tới novels. Đây là những đặc điểm của lược đồ hiện hành, cần được phân biệt với các quan hệ đã được cơ sở dữ liệu ràng buộc.')

    heading(doc, '3.6.4. Từ điển dữ liệu', 2)
    para(doc, 'Các bảng dưới đây mô tả thuộc tính, kiểu dữ liệu và ràng buộc theo các tệp khởi tạo lược đồ và migration hiện có của dự án. Ký hiệu sử dụng: PK là thành phần khóa chính; FK là khóa ngoại; UQ là duy nhất; NN là không được NULL; AI là tự tăng. Với khóa ghép, tính duy nhất được xét trên toàn bộ tổ hợp cột, không xét riêng từng thành phần.')
    para(doc, 'MySQL biểu diễn BOOLEAN tương đương TINYINT(1). Các trường TIMESTAMP dùng cho thời điểm tạo hoặc cập nhật; TIMESTAMP(3) có độ chính xác mili giây. Những trường thời hạn dạng BIGINT trong xác thực, đồng bộ tiến độ và bầu chọn lưu thời gian Unix theo mili giây. Các cột trạng thái dạng VARCHAR được kiểm tra giá trị và chuyển trạng thái ở tầng nghiệp vụ.')
    n = 3
    for title, tnames in GROUPS:
        heading(doc, title, 3)
        for tname in tnames:
            t = model[tname]
            p = heading(doc, f'Bảng {tname}', 4)
            p.paragraph_format.space_before = Pt(8)
            p = para(doc, PURPOSE[tname], indent=False)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            p = caption(doc, f'Bảng 3.6.{n}. Cấu trúc bảng {tname}')
            p.paragraph_format.keep_with_next = True
            table(doc, ['Thuộc tính', 'Kiểu dữ liệu', 'Ràng buộc và ý nghĩa'], [
                (f['name'], f['type'], field_explanation(tname, f, t)) for f in t['fields']
            ], [4.4, 3.5, 8.1])
            notes = []
            if len(t['primary']) > 1:
                notes.append('Khóa chính ghép: (' + ', '.join(t['primary']) + ').')
            for uq in t['unique']:
                if len(uq) > 1:
                    notes.append('Ràng buộc duy nhất: (' + ', '.join(uq) + ').')
            if tname == 'election_ballots':
                notes.append('Một cử tri có thể chọn nhiều ứng viên theo thể lệ; khóa chính ngăn lặp cùng lựa chọn trong một kỳ.')
            if tname == 'reading_sessions':
                notes.append('Chỉ lưu phiên hiện hành của mỗi cặp người dùng và chương; không phải bảng lưu toàn bộ lịch sử phiên.')
            if tname == 'novel_tags':
                notes.append('Nhãn được lưu trực tiếp bằng chuỗi; lược đồ chưa có bảng danh mục tags riêng.')
            if tname == 'friends':
                notes.append('Khóa chính phân biệt thứ tự hai tài khoản; kiểm tra cặp đảo chiều và tự kết bạn thuộc tầng ứng dụng.')
            if notes:
                p = para(doc, ' '.join(notes), indent=False, size=12, italic=True)
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(5)
            n += 1

    heading(doc, '3.6.5. Ràng buộc toàn vẹn và xử lý dữ liệu', 2)
    heading(doc, 'a) Toàn vẹn thực thể và chống trùng lặp', 3)
    para(doc, 'Mỗi bảng có khóa chính để nhận diện bản ghi. Các bảng thực thể sử dụng mã số nguyên tự tăng; các bảng liên kết thường sử dụng khóa chính ghép. Ví dụ, bookmarks dùng (user_id, novel_id), chapter_reads dùng (user_id, chapter_id) và collection_novels dùng (collection_id, novel_id). Cách tổ chức này ngăn cùng một quan hệ được ghi lặp trong cơ sở dữ liệu.')
    para(doc, 'Các thuộc tính username, email, genres.name và library_collections.share_id có ràng buộc duy nhất theo lược đồ. Đối với reviews, tổ hợp (user_id, novel_id) giới hạn mỗi người dùng có một đánh giá cho một truyện. Trong phân hệ tuyển chọn, (election_id, user_id) ở election_candidates ngăn một tài khoản có nhiều hồ sơ trong cùng kỳ.')
    heading(doc, 'b) Toàn vẹn tham chiếu và chính sách xóa', 3)
    para(doc, 'Các khóa ngoại bảo đảm bản ghi con tham chiếu đến bản ghi cha tồn tại. Nhiều bảng nội dung và tương tác sử dụng ON DELETE CASCADE để xóa dữ liệu phụ thuộc khi tài khoản, truyện, chương hoặc bài viết bị xóa. Việc xóa cần được kiểm soát ở tầng nghiệp vụ vì có thể ảnh hưởng đến nhiều bản ghi liên quan.')
    para(doc, 'Một số liên kết sử dụng ON DELETE SET NULL: khi truyện được chọn làm ảnh bìa kệ bị xóa, library_collections.cover_novel_id chuyển về NULL; khi tài khoản người chỉnh sửa bị xóa, chapter_versions.editor_id chuyển về NULL để giữ phiên bản chương nếu chương còn tồn tại. Các khóa ngoại trong phân hệ tuyển chọn không khai báo xóa dây chuyền nên việc xóa bản ghi đang được tham chiếu bị hạn chế.')
    para(doc, 'Không phải mọi cột có hậu tố _id hoặc _by đều đã có khóa ngoại. Các trường election_candidates.decided_by, election_complaints.resolved_by và election_audit.actor_id đang lưu mã tham chiếu nghiệp vụ nhưng chưa khai báo FK trong DDL. Vì vậy, tầng ứng dụng phải kiểm tra chủ thể thao tác và không được coi các cột này đã được MySQL bảo vệ bằng ràng buộc tham chiếu.')
    heading(doc, 'c) Ràng buộc miền giá trị và tính nhất quán nghiệp vụ', 3)
    para(doc, 'Điểm đánh giá reviews.stars có CHECK từ 1 đến 5. API lưu tiến độ kiểm tra scroll_ratio và paragraph_offset trong khoảng 0 đến 1, paragraph_index không âm, thời gian cập nhật hợp lệ và chương thực sự thuộc truyện đang đọc. Các khóa ngoại riêng lẻ của reading_progress không tự bảo đảm quan hệ chapter_id thuộc novel_id, nên kiểm tra này được thực hiện trước khi ghi dữ liệu.')
    para(doc, 'Tương tự, API tuyển chọn kiểm tra ứng viên thuộc đúng kỳ, tư cách cử tri, thời hạn bỏ phiếu và giới hạn số lựa chọn theo thể lệ. Khóa ghép của election_ballots chỉ ngăn phiếu trùng đối với cùng một ứng viên, không tự áp đặt số ứng viên tối đa mà cử tri được chọn. Các trạng thái xuất bản, kết bạn và xét duyệt được kiểm tra theo luồng nghiệp vụ thay vì khai báo bằng ENUM trong cơ sở dữ liệu.')
    heading(doc, 'd) Giao dịch và cập nhật đồng thời', 3)
    para(doc, 'Các thao tác cần thay đổi nhiều bảng được tổ chức bằng giao dịch để dữ liệu chỉ được chốt khi toàn bộ bước xử lý thành công. Trong tuyển chọn, thao tác ghi khóa dòng kỳ bầu chọn; việc quyết định bổ nhiệm, cập nhật vai trò, ghi nhật ký và tạo thông báo được thực hiện trong cùng giao dịch nhằm hạn chế xung đột giữa bỏ phiếu, đóng kỳ và phân quyền.')
    para(doc, 'Khi đồng bộ tiến độ đọc, máy chủ so sánh client_updated_at để tránh yêu cầu gửi chậm ghi đè vị trí mới hơn. Khi ghi nhận hoàn thành chương, phiên đọc và điều kiện thời gian được kiểm tra; khóa chính của chapter_reads ngăn cộng điểm lặp cho cùng một chương. Các trường novels.reads, novels.rating và novels.bookmarks_count là dữ liệu tổng hợp được lưu sẵn; khi bổ sung hoặc sửa luồng cập nhật, cần bảo đảm chúng nhất quán với dữ liệu chi tiết.')

    heading(doc, '3.6.6. Chỉ mục, bảo mật và đối chiếu lược đồ', 2)
    heading(doc, 'a) Thiết kế chỉ mục theo truy vấn', 3)
    para(doc, 'Ngoài chỉ mục khóa chính, duy nhất và khóa ngoại, lược đồ bổ sung chỉ mục theo những truy vấn thường dùng. Chỉ mục ghép được lựa chọn dựa trên các cột lọc và thứ tự sắp xếp của từng luồng chức năng.')
    caption(doc, 'Bảng 3.6.40. Một số chỉ mục phục vụ truy vấn thường dùng')
    table(doc, ['Bảng và cột lập chỉ mục', 'Mục đích'], [
        ('chapters(novel_id, status)', 'Lấy chương của một truyện theo trạng thái xuất bản.'),
        ('chapters(scheduled_release)', 'Tìm chương đến thời điểm phát hành.'),
        ('chapter_versions(chapter_id, id)', 'Tra cứu lịch sử phiên bản theo chương.'),
        ('forum_posts(created_at DESC), forum_posts(author_id)', 'Hiển thị bài mới và lọc theo người viết.'),
        ('notifications(user_id, created_at DESC)', 'Lấy thông báo gần nhất của một tài khoản.'),
        ('messages(receiver_id, is_read)', 'Tìm tin nhắn chưa đọc của người nhận.'),
        ('messages(sender_id, receiver_id, id)', 'Tra cứu hội thoại theo hai tài khoản và mã tin nhắn.'),
        ('novels(reads DESC, id DESC)', 'Sắp xếp truyện theo lượt đọc, kèm mã để ổn định thứ tự.'),
        ('election_candidates(decision, ended_at, term_ends_at, election_id)', 'Tìm nhiệm kỳ đến hạn và xử lý kết thúc quyền.'),
        ('election_audit(election_id, id)', 'Tra cứu nhật ký trong từng kỳ tuyển chọn.'),
    ], [8.5, 7.5])
    heading(doc, 'b) Bảo vệ dữ liệu xác thực và dữ liệu riêng tư', 3)
    para(doc, 'Mật khẩu được lưu dưới dạng chuỗi băm bcrypt. Bảng email_otps lưu dấu băm mã OTP thay cho mã gốc; thời hạn mã, khoảng cách gửi lại và số lần nhập sai được kiểm soát khi xác thực. Bảng revoked_tokens lưu dấu băm token đã thu hồi cùng thời điểm hết hạn để ngăn tiếp tục sử dụng phiên đã đăng xuất.')
    para(doc, 'Quyền truy cập được xác định từ tài khoản và user_roles tại máy chủ. Nội dung tin nhắn, kệ sách riêng tư, nhận xét nội bộ, khiếu nại và thông tin phiếu bầu chỉ được cung cấp theo quyền của từng đối tượng. Mã share_id của kệ sách không thay thế việc kiểm tra is_public. Tài khoản cơ sở dữ liệu dùng khi vận hành cần được giới hạn quyền phù hợp; các giá trị do người dùng nhập được truyền qua tham số truy vấn.')
    heading(doc, 'c) Đối chiếu ERD và lược đồ cài đặt', 3)
    para(doc, 'Ảnh ERD tại Hình 3.6.1 được giữ nguyên theo sơ đồ cung cấp. Phần từ điển dữ liệu mô tả lược đồ cài đặt trong server/schema.sql, các migration tại server/db.ts, server/readerSchema.ts và server/electionSchema.ts. Những vị trí khác phiên bản được ghi nhận để tránh nhầm lẫn khi triển khai hoặc đối chiếu báo cáo.')
    para(doc, 'Cụ thể, ảnh ERD thể hiện users.email có độ dài 255 ký tự và có các trường google_id, is_email_verified; migration hiện tại bổ sung email VARCHAR(150) và chưa khai báo hai trường còn lại. Bảng email_otps trên ảnh thể hiện trường mã OTP và thời điểm tạo; mã nguồn hiện tại sử dụng email, code_hash, expires_at, sent_at và attempts, trong đó các mốc thời gian là BIGINT. Các khác biệt này không làm thay đổi cách phân chia các nhóm nghiệp vụ nhưng cần đồng bộ khi cập nhật sơ đồ ở những phiên bản sau.')
    para(doc, 'Thiết kế được tổ chức theo hướng tách dữ liệu thực thể và dữ liệu quan hệ để giảm trùng lặp. Tuy nhiên, novels.genre vẫn lưu chuỗi thể loại và một số bộ đếm được lưu sẵn, nên không khẳng định toàn bộ lược đồ đã chuẩn hóa hoàn toàn. Khi mở rộng, có thể tách quan hệ truyện–thể loại thành bảng liên kết và bổ sung ràng buộc tham chiếu cho các mã chủ thể nghiệp vụ còn thiếu, sau khi rà soát dữ liệu và chính sách lưu lịch sử.')

    doc.save(OUTPUT)
    validate(OUTPUT, model)


def validate(path, model):
    doc = Document(path)
    full = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
    assert '\ufffd' not in full
    assert full.startswith('3.6. Thiết kế cơ sở dữ liệu')
    assert len(doc._element.xpath('.//a:blip')) == 1
    assert len(doc.tables) == 40
    with ZipFile(path) as z:
        media = [n for n in z.namelist() if n.startswith('word/media/')]
        assert len(media) == 1
        from io import BytesIO
        from PIL import Image
        embedded = Image.open(BytesIO(z.read(media[0]))).convert('RGBA')
        original = Image.open(ERD).convert('RGBA')
        assert embedded.size == original.size
        assert embedded.tobytes() == original.tobytes(), 'ERD pixels changed'
    for name in model:
        assert f'Bảng {name}' in full, name
        for f in model[name]['fields']:
            assert f['name'] in full
    report = {
        'document': str(path),
        'tables_in_model': len(model),
        'fields_in_dictionary': sum(len(t['fields']) for t in model.values()),
        'word_tables': len(doc.tables),
        'embedded_images': len(doc._element.xpath('.//a:blip')),
        'erd_sha256': hashlib.sha256(ERD.read_bytes()).hexdigest(),
        'erd_embedded_pixel_identical': True,
        'body_font': 'Times New Roman 13 pt',
        'dictionary_font': 'Times New Roman 12 pt',
        'paper': 'A4 portrait',
        'words_whitespace': len(full.split()),
        'schema_sources': sorted({t['source'] for t in model.values()}),
    }
    (OUT / 'validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    (OUT / 'database_model.json').write_text(json.dumps(model, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    build()
