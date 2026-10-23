"""Create an editable Vietnamese theory chapter with native Word formatting."""

from __future__ import annotations

import json
import os
from pathlib import Path
import re
import sys
from zipfile import ZipFile
from xml.etree import ElementTree as ET

LIBS = Path(os.environ['TEMP']) / 'mugenbunko-test-report-libs'
sys.path.insert(0, str(LIBS))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent
OUTPUT = OUT / 'Chuong_2_Co_so_ly_thuyet_MugenBunko.docx'
FONT = 'Times New Roman'
SIZE = 13


def set_font(run):
    run.font.name = FONT
    run.font.size = Pt(SIZE)
    run.font.color.rgb = RGBColor(0, 0, 0)
    props = run._element.get_or_add_rPr()
    fonts = props.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        props.insert(0, fonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn(f'w:{key}'), FONT)
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn(f'w:{key}'), None)
    complex_size = props.find(qn('w:szCs'))
    if complex_size is None:
        complex_size = OxmlElement('w:szCs')
        props.append(complex_size)
    complex_size.set(qn('w:val'), str(SIZE * 2))
    language = props.find(qn('w:lang'))
    if language is None:
        language = OxmlElement('w:lang')
        props.append(language)
    language.set(qn('w:val'), 'vi-VN')


def build():
    document = Document()
    sec = document.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    sec.header_distance = Cm(0.7)
    sec.footer_distance = Cm(0.8)

    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Header', 'Footer'):
        style = document.styles[name]
        style.font.name = FONT
        style.font.size = Pt(SIZE)
        style.font.color.rgb = RGBColor(0, 0, 0)
        props = style.element.get_or_add_rPr()
        fonts = props.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            props.insert(0, fonts)
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn(f'w:{key}'), FONT)
        for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn(f'w:{key}'), None)

    normal = document.styles['Normal'].paragraph_format
    normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.first_line_indent = Cm(1)
    normal.line_spacing = 1.5
    normal.space_before = Pt(0)
    normal.space_after = Pt(6)
    normal.widow_control = True

    for name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = document.styles[name]
        style.font.bold = True
        fmt = style.paragraph_format
        fmt.first_line_indent = Cm(0)
        fmt.left_indent = Cm(0)
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.line_spacing = 1.5
        fmt.space_before = Pt(10 if name == 'Heading 1' else 6)
        fmt.space_after = Pt(6)
        fmt.keep_with_next = True
        fmt.keep_together = True
    document.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.core_properties.title = 'Chương 2. Cơ sở lý thuyết'
    document.core_properties.subject = 'Hệ thống đọc, xuất bản và mạng xã hội tiểu thuyết đa nền tảng MugenBunko'
    document.core_properties.author = 'Nguyễn Hoàng Việt'
    document.core_properties.keywords = 'MugenBunko; cơ sở lý thuyết; React; Expo; Express; MySQL'
    document.core_properties.comments = 'Chương báo cáo độc lập; toàn bộ văn bản Times New Roman 13 pt.'

    settings = document.settings.element
    auto_update = OxmlElement('w:updateFields')
    auto_update.set(qn('w:val'), 'true')
    settings.append(auto_update)
    default_run_props = document.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    for tag, value in (('w:sz', '26'), ('w:szCs', '26')):
        element = default_run_props.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            default_run_props.append(element)
        element.set(qn('w:val'), value)
    default_fonts = default_run_props.find(qn('w:rFonts'))
    if default_fonts is not None:
        default_fonts.attrib.clear()
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            default_fonts.set(qn(f'w:{key}'), FONT)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    footer.paragraph_format.line_spacing = 1
    footer.paragraph_format.space_after = Pt(0)
    set_font(footer.add_run('Trang '))
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    field_run = OxmlElement('w:r')
    field_props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn(f'w:{key}'), FONT)
    field_props.append(fonts)
    for tag in ('w:sz', 'w:szCs'):
        size = OxmlElement(tag)
        size.set(qn('w:val'), '26')
        field_props.append(size)
    field_run.append(field_props)
    field_text = OxmlElement('w:t')
    field_text.text = '1'
    field_run.append(field_text)
    field.append(field_run)
    footer._p.append(field)

    content = (OUT / 'chapter2_content.txt').read_text(encoding='utf-8')
    paragraphs = [p.strip() for p in content.split('\n\n') if p.strip()]
    refs = False
    for text in paragraphs:
        if text.startswith('CHƯƠNG 2.'):
            paragraph = document.add_paragraph(text, 'Title')
        elif text == 'TÀI LIỆU THAM KHẢO':
            refs = True
            paragraph = document.add_paragraph(text, 'Heading 1')
            paragraph.paragraph_format.page_break_before = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif re.match(r'^2\.\d+\.\d+\. ', text):
            paragraph = document.add_paragraph(text, 'Heading 2')
        elif re.match(r'^2\.\d+\. ', text):
            paragraph = document.add_paragraph(text, 'Heading 1')
        else:
            paragraph = document.add_paragraph(text)
            if refs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.space_after = Pt(7)
                paragraph.paragraph_format.keep_together = True
        for run in paragraph.runs:
            set_font(run)

    document.save(OUTPUT)
    validate(OUTPUT)


def validate(path):
    document = Document(path)
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(path) as archive:
        for name in ('word/document.xml', 'word/footer1.xml'):
            root = ET.fromstring(archive.read(name))
            for run in root.findall('.//w:r', ns):
                if run.find('w:t', ns) is None:
                    continue
                props = run.find('w:rPr', ns)
                assert props is not None, f'Unformatted run in {name}'
                size = props.find('w:sz', ns)
                fonts = props.find('w:rFonts', ns)
                assert size is not None and size.get(qn('w:val')) == '26', f'Incorrect font size in {name}'
                assert fonts is not None and fonts.get(qn('w:ascii')) == FONT, f'Incorrect font in {name}'

    body = '\n'.join(p.text for p in document.paragraphs)
    assert body.startswith('CHƯƠNG 2. CƠ SỞ LÝ THUYẾT')
    assert len(re.findall(r'^2\.[1-8]\. ', body, re.M)) == 8
    assert '```' not in body and '**' not in body
    assert not re.search(r'^\s*[#*]\s', body, re.M)
    assert '\ufffd' not in body
    assert len(body.split()) >= 4000
    references = set(re.findall(r'^\[(\d+)\]', body, re.M))
    body_without_refs = body.split('TÀI LIỆU THAM KHẢO')[0]
    citations = set(re.findall(r'\[(\d+)\]', body_without_refs))
    assert citations <= references, f'Missing references: {citations - references}'
    assert references <= citations, f'Uncited references: {references - citations}'
    report = {
        'document': str(path),
        'font': FONT,
        'size_pt': SIZE,
        'paper': 'A4',
        'line_spacing': 1.5,
        'margins_cm': {'left': 3, 'right': 2, 'top': 2, 'bottom': 2},
        'paragraphs': len(document.paragraphs),
        'words_whitespace': len(body.split()),
        'characters': len(body),
        'major_sections': 8,
        'subsections': len(re.findall(r'^2\.\d+\.\d+\. ', body, re.M)),
        'references': len(references),
        'all_text_runs_font_checked': True,
        'markdown_syntax_found': False,
    }
    (OUT / 'validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=True, indent=2))


if __name__ == '__main__':
    build()
