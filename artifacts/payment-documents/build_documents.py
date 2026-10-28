"""Create the payment documentation set from the approved plan and specifications.

Run with the Python executable returned by load_workspace_dependencies.
Word previews and QA evidence are kept in .cache/payment-documents-qa.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from document_content import documents, SOURCES

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
QA = ROOT / '.cache/payment-documents-qa'
PLAN = (ROOT / 'docs/PAYMENT_INTEGRATION_PLAN.md').read_text(encoding='utf-8-sig')
FONT = 'Calibri'
PAGE_WIDTH = 15.5


def el(tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def set_font(run, name=FONT, size=None, bold=None, color='000000'):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = el('rFonts')
        rpr.insert(0, rf)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + attr), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = el('lang', val='vi-VN')
        rpr.append(lang)


def add_link(paragraph, label, url=None, anchor=None):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    h = el('hyperlink')
    if anchor:
        h.set(qn('w:anchor'), anchor)
    else:
        h.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = el('r')
    rpr = el('rPr')
    rpr.append(el('rFonts', ascii=FONT, hAnsi=FONT))
    rpr.append(el('color', val='000000' if anchor else '1B4F72'))
    rpr.append(el('sz', val=22))
    r.append(rpr)
    t = el('t')
    t.text = label
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def inline(p, text, size=11, color='000000', bold=False):
    pattern = re.compile(r'(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*)')
    for token in pattern.split(text):
        if not token:
            continue
        match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', token)
        if match:
            add_link(p, match[1], match[2])
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            set_font(run, 'Consolas', min(size, 10), color=color)
        else:
            is_bold = token.startswith('**') and token.endswith('**')
            run = p.add_run(token[2:-2] if is_bold else token)
            set_font(run, size=size, color=color, bold=(bold or is_bold))


def clean_heading(s):
    s = s.replace('_', ' ')
    s = re.sub(r'^\d+(?:\.\d+)*\.?\s*', '', s)
    s = re.sub(r'[`*#]', '', s)
    s = re.sub(r'[^\w\s]', ' ', s, flags=re.UNICODE)
    return re.sub(r'\s+', ' ', s).strip()


class Writer:
    def __init__(self, spec):
        self.spec = spec
        self.doc = Document()
        self.headings = []
        self.section_count = 0
        self.subsection_count = 0
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.9)
        sec.left_margin, sec.right_margin = Cm(2.8), Cm(2.2)
        sec.footer_distance = Cm(.9)
        sec.header_distance = Cm(.8)
        styles = self.doc.styles
        for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'List Paragraph'):
            style = styles[name]
            style.font.name = FONT
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.underline = False
            for border in style._element.xpath('./w:pPr/w:pBdr'):
                border.getparent().remove(border)
            style._element.get_or_add_rPr().append(el('rFonts', ascii=FONT, hAnsi=FONT, eastAsia=FONT, cs=FONT))
            pf = style.paragraph_format
            pf.widow_control = True
            pf.space_after = Pt(6)
            pf.line_spacing = 1.12
        styles['Normal'].font.size = Pt(11)
        styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for name, size in [('Title', 23), ('Subtitle', 13), ('Heading 1', 15), ('Heading 2', 13), ('Heading 3', 12)]:
            styles[name].font.size = Pt(size)
            styles[name].font.bold = name != 'Subtitle'
            styles[name].paragraph_format.keep_with_next = True
            styles[name].paragraph_format.space_before = Pt(14 if name.startswith('Heading') else 0)
            styles[name].paragraph_format.space_after = Pt(8)
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        styles['Caption'].font.size = Pt(10.5)
        styles['Caption'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        styles['Caption'].paragraph_format.space_after = Pt(9)
        # Page fields help navigation in each separate technical document.
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(f'MB PAY {spec["number"]:02d}   |   Trang ')
        set_font(run, size=9)
        field = el('fldSimple', instr='PAGE')
        r = el('r'); t = el('t'); t.text = '1'; r.append(t); field.append(r); p._p.append(field)
        props = self.doc.core_properties
        props.title = spec['title']
        props.subject = 'Tích hợp thanh toán điện tử trong MugenBunko'
        props.author = 'MugenBunko'
        props.keywords = 'MugenBunko, VNPAY, sandbox, đồ án tốt nghiệp'
        props.language = 'vi-VN'
        self.doc.add_paragraph(spec['title'], 'Title')
        self.p(f'Tài liệu MB PAY {spec["number"]:02d} | Phiên bản 1.0 | Ngày {spec.get("date", "27/10/2026")}', size=10.5)
        self.p(spec['intro'])
        self.p('Trạng thái: kế hoạch kỹ thuật cho nhánh codex/payment-integration. Module thanh toán chưa được triển khai; các API và bảng mới bên dưới là hợp đồng dự kiến.', size=10)
        if spec.get('reference_in_intro'):
            self.p('Căn cứ: docs/PAYMENT_INTEGRATION_PLAN.md. Đọc cùng: ' + '; '.join(spec['related']) + '.', size=9.5)
        self.doc.add_paragraph('Mục lục', 'Heading 1')
        self.toc_placeholder = self.doc.add_paragraph()
        self.doc.add_paragraph().paragraph_format.space_after = Pt(3)

    def p(self, text, size=11, style=None):
        p = self.doc.add_paragraph(style=style)
        inline(p, str(text), size)
        return p

    def h(self, text, level=1):
        title = clean_heading(text)
        if level == 1:
            self.section_count += 1
            self.subsection_count = 0
            prefix = str(self.section_count)
        else:
            self.subsection_count += 1
            prefix = f'{self.section_count} {self.subsection_count}'
        p = self.doc.add_paragraph(f'{prefix} {title}', f'Heading {min(level, 3)}')
        anchor = f'mb_{self.spec["number"]}_{len(self.headings)+1}'
        start = el('bookmarkStart', id=len(self.headings)+1, name=anchor)
        end = el('bookmarkEnd', id=len(self.headings)+1)
        p._p.insert(0, start); p._p.append(end)
        if level == 1:
            self.headings.append((f'{prefix} {title}', anchor))
        else:
            # Use unique bookmark IDs even for secondary headings.
            self.headings.append((None, anchor))
        return p

    def table(self, headers, rows, widths=None):
        cols = len(headers)
        if headers == ['Method', 'Endpoint', 'Quyền', 'Mục đích']:
            widths = [1.6, 4.8, 2.9, 6.2]
        if widths is None:
            weights = []
            for i, header in enumerate(headers):
                lengths = [len(re.sub(r'[`*]', '', str(row[i]))) for row in rows if i < len(row)]
                avg = sum(lengths) / max(len(lengths), 1)
                weights.append(max(8, min(55, avg ** .75 * 3 + len(header) * .4)))
            widths = [PAGE_WIDTH * w / sum(weights) for w in weights]
        table = self.doc.add_table(rows=1, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        pr = table._tbl.tblPr
        borders = el('tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(el(side, val='single', sz='5', color='D9D9D9'))
        pr.append(borders)
        margins = el('tblCellMar')
        for side, value in [('top', 80), ('bottom', 80), ('left', 90), ('right', 90)]:
            margins.append(el(side, w=value, type='dxa'))
        pr.append(margins)
        for row_index, values in enumerate([headers, *rows]):
            row = table.rows[0] if row_index == 0 else table.add_row()
            trpr = row._tr.get_or_add_trPr()
            trpr.append(el('cantSplit'))
            if row_index == 0:
                trpr.append(el('tblHeader'))
            for i, value in enumerate(values):
                cell = row.cells[i]
                cell.width = Cm(widths[i])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                fill = '344256' if row_index == 0 else ('F3F5F7' if row_index % 2 == 0 else 'FFFFFF')
                cell._tc.get_or_add_tcPr().append(el('shd', fill=fill))
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                if row_index in (0, 1) or (len(rows) > 2 and row_index == len(rows)-1):
                    p.paragraph_format.keep_with_next = True
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if (widths[i] < 2.8 or (len(str(value)) < 17 and i > 0)) else WD_ALIGN_PARAGRAPH.LEFT
                inline(p, str(value), 10.5 if cols <= 4 else 10, 'FFFFFF' if row_index == 0 else '000000', row_index == 0)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(3)
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(3)
        spacer.add_run().font.size = Pt(3)

    def figure(self, name, caption):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(QA / 'figures' / f'{name}.png'), width=Cm(15.5))
        docpr = run._r.xpath('.//wp:docPr')[0]
        docpr.set('descr', caption)
        p.paragraph_format.keep_with_next = True
        self.doc.add_paragraph(caption, 'Caption')

    def code(self, text):
        for line in text.strip('\n').splitlines():
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            # Preserve all characters and allow ordinary Word wrapping.
            set_font(p.add_run(line or ' '), 'Consolas', 8.8)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(3)

    def markdown(self, text):
        lines = text.strip().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line or line.startswith('<!--'):
                i += 1; continue
            if line.startswith('```'):
                language = line[3:].strip(); block = []; i += 1
                while i < len(lines) and not lines[i].strip().startswith('```'):
                    block.append(lines[i]); i += 1
                if language == 'mermaid':
                    self.figure('sequence', 'Hình 1 Luồng xác nhận thanh toán giữa client backend cơ sở dữ liệu và cổng')
                else:
                    self.code('\n'.join(block))
                i += 1; continue
            if line.startswith('|'):
                records = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    raw = lines[i].strip()[1:-1]
                    record = [part.strip() for part in re.split(r'(?<!\\)\|', raw)]
                    if not all(re.fullmatch(r':?-+:?', x.replace(' ', '')) for x in record):
                        records.append(record)
                    i += 1
                # Markdown may contain code unions with pipes. Normalize defensive parsing.
                n = len(records[0])
                rows = [r[:n-1] + [' | '.join(r[n-1:])] if len(r) > n else r + ['']*(n-len(r)) for r in records[1:]]
                self.table(records[0], rows)
                continue
            m = re.match(r'^(#{2,4})\s+(.*)', line)
            if m:
                self.h(m[2], min(len(m[1])-1, 3)); i += 1; continue
            if line.startswith('- '):
                p = self.p(line[2:], style='List Bullet')
                p.paragraph_format.space_after = Pt(4)
                i += 1; continue
            if re.match(r'^\d+\.\s', line):
                # Keep explicit sequence numbers rather than continuing unrelated Word lists.
                p = self.p(line)
                p.paragraph_format.left_indent = Cm(.4)
                p.paragraph_format.first_line_indent = Cm(-.4)
                i += 1; continue
            parts = [line]; i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r'^(#{2,4}\s|\||```|- |\d+\.\s)', lines[i].strip()):
                parts.append(lines[i].strip()); i += 1
            self.p(' '.join(parts))

    def finish(self):
        # Static bookmarked TOC remains functional without updating fields.
        placeholder = self.toc_placeholder
        for title, anchor in self.headings:
            if title is None:
                continue
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_link(p, title, anchor=anchor)
            placeholder._p.addprevious(p._p)
        placeholder._p.getparent().remove(placeholder._p)
        if self.spec.get('sources'):
            self.h('Tài liệu liên quan và nguồn tham khảo')
            self.p('Căn cứ triển khai: docs/PAYMENT_INTEGRATION_PLAN.md, phiên bản ngày 27/10/2026. Mã nguồn hiện tại quyết định phần đã tồn tại; đặc tả trong bộ tài liệu xác định phần cần xây dựng.')
            for index in self.spec['sources']:
                label, url = SOURCES[index]
                p = self.doc.add_paragraph()
                add_link(p, label, url)
            self.p('Đọc cùng các tài liệu: ' + '; '.join(self.spec['related']) + '.', size=10.5)
        elif not self.spec.get('reference_in_intro'):
            self.p('Căn cứ: docs/PAYMENT_INTEGRATION_PLAN.md. Đọc cùng: ' + '; '.join(self.spec['related']) + '.', size=10.5)
        # The source heading is intentionally excluded from the static main-content TOC.
        path = OUT / self.spec['filename']
        self.doc.save(path)
        return {'file': str(path), 'title': self.spec['title'], 'sections': self.section_count,
                'paragraphs': len(self.doc.paragraphs), 'tables': len(self.doc.tables)}


def source_section(number):
    pattern = rf'(?ms)^## {number}\. .*?(?=^## \d+\. |\Z)'
    match = re.search(pattern, PLAN)
    if not match:
        raise ValueError(f'Missing source section {number}')
    return match[0].strip()


def source_subsection(number, sub):
    text = source_section(number)
    pattern = rf'(?ms)^### {number}\.{sub}\. .*?(?=^### {number}\.\d+\. |\Z)'
    match = re.search(pattern, text)
    if not match:
        raise ValueError(f'Missing source subsection {number}.{sub}')
    return re.sub(r'^### ', '## ', match[0].strip(), count=1)


def figures():
    folder = QA / 'figures'; folder.mkdir(parents=True, exist_ok=True)
    regular = 'C:/Windows/Fonts/times.ttf'
    bold = 'C:/Windows/Fonts/timesbd.ttf'
    def font(size, heavy=False): return ImageFont.truetype(bold if heavy else regular, size)
    def box(draw, rect, title, lines):
        draw.rounded_rectangle(rect, radius=14, fill='#F4F6F8', outline='#344256', width=3)
        x1,y1,x2,y2=rect
        draw.text(((x1+x2)/2,y1+18),title,font=font(40,True),fill='black',anchor='mt')
        y=y1+67
        for line in lines:
            draw.text(((x1+x2)/2,y),line,font=font(34),fill='black',anchor='mt'); y+=44
    def arrow(draw,a,b,label=None):
        import math
        draw.line((a,b),fill='#344256',width=4)
        angle=math.atan2(b[1]-a[1],b[0]-a[0]); side=16
        pts=[b,(b[0]-side*math.cos(angle-.5),b[1]-side*math.sin(angle-.5)),(b[0]-side*math.cos(angle+.5),b[1]-side*math.sin(angle+.5))]
        draw.polygon(pts,fill='#344256')
        if label: draw.text(((a[0]+b[0])/2,(a[1]+b[1])/2-20),label,font=font(32),fill='black',anchor='mm')
    im=Image.new('RGB',(1800,1000),'white'); d=ImageDraw.Draw(im)
    box(d,(70,40,660,190),'Web React',['Ví   Lịch sử   Kết quả'])
    box(d,(1130,40,1730,190),'Expo Mobile',['Browser sandbox   Deep link'])
    box(d,(390,320,1410,525),'Express API',['Auth và kiểm tra đầu vào','PaymentService   WalletService','Adapter VNPAY   Worker đối soát'])
    box(d,(70,730,790,940),'MySQL',['Đơn nạp   Ledger   Số dư','Sự kiện   Audit   Lease'])
    box(d,(1080,730,1730,940),'VNPAY Sandbox',['Checkout   IPN   Truy vấn'])
    arrow(d,(365,190),(650,320),'API có phiên'); arrow(d,(1430,190),(1150,320),'API có phiên')
    arrow(d,(650,525),(440,730),'Transaction'); arrow(d,(1200,525),(1400,730),'Xác minh giao thức')
    im.save(folder/'architecture.png')
    im=Image.new('RGB',(1800,1130),'white'); d=ImageDraw.Draw(im)
    xs=[200,660,1110,1590]
    for x,label in zip(xs,['Web hoặc Mobile','Express API','MySQL','Cổng Sandbox']):
        d.rectangle((x-170,35,x+170,115),fill='#F3F5F7',outline='#344256',width=3)
        d.text((x,74),label,font=font(36,True),anchor='mm',fill='black')
        d.line((x,115,x,1080),fill='#ABB3BD',width=2)
    for a,b,y,label in [(0,1,205,'1  Tạo đơn và idempotency key'),(1,2,330,'2  Lưu snapshot đơn rồi commit'),(1,0,455,'3  Trả checkout URL và mã đơn'),(0,3,580,'4  Người dùng thanh toán sandbox'),(3,1,705,'5  IPN có chữ ký'),(1,2,830,'6  Khóa đơn và cấp xu trong transaction'),(0,1,1000,'7  Đọc lại trạng thái và số dư')]:
        arrow(d,(xs[a],y),(xs[b],y))
        d.text(((xs[a]+xs[b])/2,y-35),label,font=font(34),fill='black',anchor='mm')
    im.save(folder/'sequence.png')
    im=Image.new('RGB',(1800,1100),'white'); d=ImageDraw.Draw(im)
    boxes=[((40,35,570,250),'users',['id PK','coins   số dư đọc nhanh','status   trạng thái tài khoản']),
           ((650,35,1180,250),'coin_packages',['id PK   code UNIQUE','amount_vnd   coin_amount','active']),
           ((40,440,850,705),'payment_orders',['id PK   user_id FK   package_id FK','provider và môi trường merchant','snapshot giá xu   trạng thái','idempotency   lease đối soát']),
           ((1020,440,1760,705),'wallet_entries',['id PK   user_id FK','payment_order_id UNIQUE','delta_coins   balance_after','migration_key UNIQUE']),
           ((40,865,850,1060),'payment_events',['order_id FK cho phép NULL','nguồn và fingerprint','kết quả xác minh và xử lý']),
           ((1020,865,1760,1060),'payment_audit',['order_id FK   actor_user_id FK','hành động   lý do','request_id   created_at'])]
    for args in boxes: box(d,*args)
    arrow(d,(305,250),(305,440),'1 tới nhiều'); arrow(d,(900,250),(690,440),'1 tới nhiều')
    arrow(d,(850,570),(1020,570),'0 hoặc 1 topup'); arrow(d,(390,705),(390,865),'1 tới nhiều')
    arrow(d,(820,705),(1260,865),'1 tới nhiều')
    im.save(folder/'erd.png')


def main():
    OUT.mkdir(parents=True, exist_ok=True); QA.mkdir(parents=True, exist_ok=True)
    figures()
    manifest=[]
    for spec in documents():
        if len(sys.argv) > 1 and str(spec['number']) not in sys.argv[1:]:
            continue
        w=Writer(spec)
        for kind, value in spec['blocks']:
            if kind == 'source': w.markdown(source_section(value))
            elif kind == 'subsource': w.markdown(source_subsection(*value))
            elif kind == 'md': w.markdown(value)
            elif kind == 'figure': w.figure(*value)
            elif kind == 'table': w.table(*value)
            else: raise ValueError(kind)
        item=w.finish(); manifest.append(item)
        print(json.dumps(item,ensure_ascii=False),flush=True)
    (QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')


if __name__ == '__main__': main()
