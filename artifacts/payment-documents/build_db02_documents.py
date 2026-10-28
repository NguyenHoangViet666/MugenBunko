"""Build DB02 technical documentation and a new daily report revision.

Use the bundled Python runtime from load_workspace_dependencies.
The previous daily report and all original technical plans remain unchanged.
"""
from pathlib import Path
import json
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[2]
QA = ROOT / '.cache/payment-db02-documents-qa'
COMMIT = '337c48a3ef06ad08a6dcee1cae1066e98b10d37b'
GITHUB = 'https://github.com/NguyenHoangViet666/MugenBunko'
FONT = 'Calibri'


def element(tag, **attrs):
    result = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        result.set(qn('w:' + key), str(value))
    return result


def format_run(run, size=11, bold=False, color='000000', face=FONT):
    run.font.name = face
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    props = run._element.get_or_add_rPr()
    props.append(element('rFonts', ascii=face, hAnsi=face, eastAsia=face, cs=face))
    props.append(element('lang', val='vi-VN'))


class Writer:
    def __init__(self, title, footer):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.7)
        sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)
        sec.footer_distance = Inches(.3)
        for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet'):
            style = self.doc.styles[name]
            style.font.name = FONT
            style.font.size = Pt(11)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.underline = False
            style._element.get_or_add_rPr().append(element('rFonts', ascii=FONT, hAnsi=FONT, eastAsia=FONT, cs=FONT))
            for border in style._element.xpath('./w:pPr/w:pBdr'):
                border.getparent().remove(border)
            pf = style.paragraph_format
            pf.space_after, pf.line_spacing = Pt(6), 1.08
            pf.widow_control = True
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for name, size in [('Title', 21), ('Heading 1', 14), ('Heading 2', 12)]:
            style = self.doc.styles[name]
            style.font.size, style.font.bold = Pt(size), True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10 if name.startswith('Heading') else 0)
            style.paragraph_format.space_after = Pt(6)
        self.doc.styles['List Bullet'].paragraph_format.space_after = Pt(4)
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        format_run(p.add_run(footer + '   |   Trang '), 9)
        field = element('fldSimple', instr='PAGE')
        run, text = element('r'), element('t')
        text.text = '1'
        run.append(text)
        field.append(run)
        p._p.append(field)
        props = self.doc.core_properties
        props.title, props.author, props.language = title, 'MugenBunko', 'vi-VN'
        props.subject = 'Triển khai thanh toán DB02 ngày 28 tháng 10 năm 2026'
        props.keywords = 'MugenBunko, DB02, migration, opening balance, ledger'

    def p(self, text='', style=None, size=11):
        p = self.doc.add_paragraph(style=style)
        for part in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text):
            if not part:
                continue
            if part.startswith('`'):
                format_run(p.add_run(part[1:-1]), min(size, 10), face='Consolas')
            elif part.startswith('**'):
                format_run(p.add_run(part[2:-2]), size, True)
            else:
                format_run(p.add_run(part), size, style in ('Title', 'Heading 1', 'Heading 2'))
        return p

    def h(self, title, level=1):
        assert not re.search(r'[^\w\s]', title), title
        return self.p(title, f'Heading {level}', 14 if level == 1 else 12)

    def page(self, title):
        self.doc.add_page_break()
        self.h(title)

    def bullet(self, text):
        return self.p(text, 'List Bullet')

    def code(self, text):
        lines = text.strip('\n').splitlines()
        for index, line in enumerate(lines):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            p.paragraph_format.keep_with_next = index < len(lines) - 1
            format_run(p.add_run(line or ' '), 10, face='Consolas')
        self.p('', size=4).paragraph_format.space_after = Pt(2)

    def table(self, headers, rows, widths, center=()):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
        for ri, values in enumerate([headers, *rows]):
            row = table.rows[0] if ri == 0 else table.add_row()
            row._tr.get_or_add_trPr().append(element('cantSplit'))
            if ri == 0:
                row._tr.get_or_add_trPr().append(element('tblHeader'))
            for ci, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = cell._tc.get_or_add_tcPr()
                margins, borders = element('tcMar'), element('tcBorders')
                for edge, amount in [('top', 65), ('bottom', 65), ('left', 90), ('right', 90)]:
                    margins.append(element(edge, w=amount, type='dxa'))
                for edge in ('top', 'left', 'bottom', 'right'):
                    borders.append(element(edge, val='single', sz=4, color='D9D9D9'))
                pr.append(margins)
                pr.append(borders)
                pr.append(element('shd', fill='344256' if ri == 0 else ('F1F5F8' if ri % 2 else 'FFFFFF')))
                p = cell.paragraphs[0]
                p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(0), 1.06
                if ci in center:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                format_run(p.add_run(str(value)), 10.5, ri == 0, 'FFFFFF' if ri == 0 else '000000')
        self.p('', size=4).paragraph_format.space_after = Pt(0)

    def link(self, label, url, lead=''):
        from docx.opc.constants import RELATIONSHIP_TYPE
        p = self.p(lead, size=10.5)
        hyperlink = element('hyperlink')
        hyperlink.set(qn('r:id'), p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        run, props, text = element('r'), element('rPr'), element('t')
        props.append(element('rFonts', ascii=FONT, hAnsi=FONT))
        props.append(element('sz', val=21))
        props.append(element('color', val='244F75'))
        props.append(element('u', val='single'))
        text.text = label
        run.append(props)
        run.append(text)
        hyperlink.append(run)
        p._p.append(hyperlink)

    def save(self, path, pages):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        return {'file': str(path), 'expected_pages': pages}


def technical():
    title = 'Tài liệu kỹ thuật DB02 về migration số dư và opening balance MugenBunko'
    w = Writer(title, 'MB PAY 11 DB02')
    w.p(title, 'Title', 21)
    w.p('Ngày 28 tháng 10 năm 2026 | Phiên bản 1.0 | Commit triển khai `337c48a`', size=10)
    w.p('DB02 đã được triển khai và kiểm thử: giữ số dư cũ, tạo một opening balance cho mỗi user và chuyển các đường tạo tài khoản sang ghi ledger cùng transaction. Tài liệu mô tả code đã tồn tại, điều kiện vận hành và cách tiếp tục sau gián đoạn. **Database ứng dụng đang dùng chưa chạy migration v2.**')
    w.h('1 Phạm vi triển khai')
    w.table(['Hạng mục', 'Trạng thái tại checkpoint'], [
        ['Code và kiểm thử DB02', 'Đã hoàn thành trên fixture MySQL local 9.2.0.'],
        ['Database ứng dụng', 'Chưa backfill hoặc đổi default/nullability; preflight vẫn khớp baseline.'],
        ['Triển khai staging và cổng', 'MySQL 8.4 chưa kiểm chứng; chưa có checkout hay VNPAY adapter.'],
    ], [1.8, 5.1])
    w.p('V1 giữ nguyên DDL, seed và checksum. V2 bổ sung cột xu chuẩn hóa và checkpoint dữ liệu; runner xác minh prefix version, checksum và metadata trước khi cho startup thành công. Các bảng nghiệp vụ của v1 không bị thay đổi trong DB02.')
    w.h('Các module liên quan', 2)
    w.table(['Tệp trong server', 'Trách nhiệm'], [
        ['paymentSchemaMigration.ts', 'Điều phối version, named lock, điều kiện chạy tường minh và startup probe.'],
        ['payments/schemaV2.ts', 'DDL, trigger freeze và checksum bất biến của v2.'],
        ['payments/balanceMigration.ts', 'Preflight, cài/kiểm tra freeze, batch backfill và đối chiếu.'],
        ['payments/walletOpening.ts', 'Opening cho user mới; hàm inspectWalletLedger chỉ đọc.'],
        ['payments/schemaV2Contract.ts', 'Metadata contract của checkpoint v2.'],
        ['paymentResetGuard.js', 'Chặn demo seed/reset trên DB đã có lịch sử hoặc đã bắt đầu v2.'],
    ], [2.75, 4.15])
    w.p('Đọc cùng `docs/payments/DB02_STATUS.md`, `SCHEMA_CONTRACT.md` và kế hoạch cơ sở dữ liệu trong tài liệu MB PAY 02. Tài liệu kế hoạch vẫn là baseline; code tại commit nêu trên quyết định phần đã triển khai.', size=10)

    w.page('2 Contract dữ liệu và điều kiện startup')
    w.h('Số dư và opening balance', 2)
    w.p('Sau v2, `users.coins` là `INT NOT NULL DEFAULT 0`, có CHECK từ 0 đến 2.000.000.000. Migration thay định nghĩa cột và giữ giá trị từng user; NULL, âm, vượt trần hoặc kiểu cột không phù hợp làm runner dừng, không tự cắt hoặc làm tròn dữ liệu.')
    w.p('Mỗi user có đúng một `opening_balance`, kể cả số dư 0 và ID âm. Key `opening:v1:<userId>` là định dạng định danh opening đã chốt từ T01; chữ v1 trong key không phải schema version. Entry có `delta_coins = balance_after = coins`, không gắn payment order. Opening sẵn có phải khớp số dư và key, nếu không runner dừng.')
    w.h('Checkpoint payment balance migration', 2)
    w.table(['Trường', 'Kiểu và ý nghĩa'], [
        ['version', 'INT UNSIGNED PRIMARY KEY; CHECK version = 2.'],
        ['last_user_id', 'INT NULL; NULL là chưa xử lý user nào, không mặc định bắt đầu từ ID 0.'],
        ['expected_user_count', 'BIGINT UNSIGNED; số user tại thời điểm freeze hoàn tất.'],
        ['expected_total_coins', 'DECIMAL(65,0); tổng xu ban đầu để tránh mất chính xác.'],
        ['completed_batches', 'INT UNSIGNED DEFAULT 0; số batch đã commit.'],
        ['updated_at', 'DATETIME(3); cập nhật cùng checkpoint trong session UTC.'],
    ], [2.25, 4.65])
    w.h('Các chế độ khởi động', 2)
    w.table(['Trường hợp', 'Hành vi'], [
        ['Production startup', 'Chỉ đọc metadata/journal; thiếu hoặc sai v2 thì dừng trước khi mở HTTP.'],
        ['Development với DB rỗng', 'Có thể khởi tạo v1/v2. Kiểm tra lại sau freeze nếu user vừa xuất hiện.'],
        ['Development có user và chưa đạt v2', 'Báo PAYMENT_BALANCE_EXPLICIT; không tự backfill khi mở API.'],
        ['Lệnh migrate:production', 'Cho phép chuyển số dư tường minh, sau khi chuẩn bị backup/restore và dừng writer.'],
    ], [2.25, 4.65])
    w.p('Startup probe dùng tài khoản DML và không cần quyền DDL. Với v2, probe kiểm tra thêm cột xu, CHECK có thực thi, metadata bảng checkpoint và sự hiện diện của bản ghi version 2. Đây là kiểm tra cấu trúc; không chạy đối chiếu toàn bộ ledger mỗi lần startup.', size=10.5)

    w.page('3 Thuật toán freeze và backfill')
    for text in [
        '1. Lấy named lock theo database bằng cùng namespace với v1. Connection migration phải có database được chọn và ở autocommit; không gọi runner trong transaction nghiệp vụ.',
        '2. Kiểm tra journal và schema v1. Với user hiện hữu, startup thường bị từ chối trước khi bắt đầu v2. Preflight kiểm tra số dư, kiểu cột và chặn backfill trên ledger đã có topup.',
        '3. Cài ba trigger BEFORE INSERT, BEFORE UPDATE và BEFORE DELETE trên users. Trigger UPDATE chỉ kiểm tra khi coins thay đổi. Khi resume, trigger phải còn đủ và đúng nội dung.',
        '4. Ghi journal v2 ở trạng thái applying sau khi đủ trigger. Trigger đọc trạng thái hiện tại bằng SELECT FOR SHARE, không dùng snapshot cũ của writer để quyết định cho ghi.',
        '5. Chụp lại số user và tổng xu sau freeze. Nếu user vừa xuất hiện trong luồng startup DB rỗng không cho phép chuyển dữ liệu, dừng với PAYMENT_BALANCE_EXPLICIT. Chuẩn hóa cột xu và CHECK.',
        '6. Đọc batch theo PK sau last_user_id, khóa user bằng FOR UPDATE. Kiểm tra hoặc tạo opening, cập nhật checkpoint và commit cùng transaction. Mặc định 250 user mỗi batch; batchSize hợp lệ từ 1 đến 1000.',
        '7. Đối chiếu một opening/user, tổng ledger với coins từng user và tổng hệ thống, đồng thời kiểm tra snapshot. Chỉ khi khớp mới ghi v2 applied, dọn trigger và giải phóng named lock.',
    ]:
        w.p(text)
    w.h('Vì sao freeze vẫn giữ sau gián đoạn', 2)
    w.p('Trigger dùng locking read nên thấy marker applying kể cả khi transaction writer đã tạo snapshot REPEATABLE READ trước đó. Shared next-key lock của writer đã qua guard buộc marker v2 chờ transaction đó kết thúc, rồi snapshot mới được chụp. Khi v2 applying, chỉ connection đang sở hữu named lock được ghi xu.')
    w.p('Nếu connection migration mất, named lock được giải phóng nhưng journal vẫn applying, nên writer tiếp tục bị chặn. Batch đã commit cùng checkpoint được giữ; batch lỗi rollback cả entry lẫn checkpoint. MySQL DDL có thể implicit commit, vì vậy resume không phải rollback toàn bộ DDL.')
    w.p('Thông báo trigger là `PAYMENT_BALANCE_MIGRATION_IN_PROGRESS`; driver có thể trả `ER_SIGNAL_EXCEPTION`. Freeze bảo vệ DML, không ngăn DBA chạy DDL/TRUNCATE hoặc tự vô hiệu guard. Vẫn phải dừng writer/seed cũ trong maintenance.', size=10.5)

    w.page('4 Các đường ghi mới và chính sách demo')
    w.h('Đăng ký tài khoản và tạo admin', 2)
    w.p('Route đăng ký chèn user với 0 xu rồi gọi `recordNewUserOpening` trên cùng connection trong `db.transaction`. Helper đọc journal v2 bằng FOR SHARE, khóa user bằng FOR UPDATE, yêu cầu coins bằng 0 và chèn opening 0. Chỉ sau đó transaction mới hoàn tất cùng role và notification.')
    w.p('CLI `create:admin` dùng cùng helper và transaction, chỉ định coins = 0 tường minh. Tài khoản đã tồn tại không được ghi đè hoặc tự nâng quyền. ADMIN_USERNAME và ADMIN_PASSWORD phải được cấu hình theo chính sách CLI hiện có; không ghi credentials vào tài liệu hoặc Git.')
    w.p('Nếu entry lỗi, helper chạy hai lần hoặc v2 chưa applied, transaction tạo user rollback. Payload coins từ client không tạo quyền cấp xu. Fixture mới dùng opening 0; fixture legacy tách riêng để thử dữ liệu nullable/default 1000.')
    w.h('Seed reset và dữ liệu mẫu', 2)
    w.table(['Đường chạy', 'Chính sách sau DB02'], [
        ['schema.sql', 'Chỉ DDL, không TRUNCATE, không user/admin hoặc mật khẩu mẫu.'],
        ['seed_novels.js và seed_forum_50.js', 'Chỉ dùng DB demo riêng trước v2 và chưa có lịch sử tài chính.'],
        ['drop-db.js', 'Kiểm tra tên DB và guard trước khi DROP; từ chối DB có lịch sử hoặc v2 đã bắt đầu.'],
        ['paymentResetGuard.js', 'Giữ named lock chung với migration cho đến khi caller đóng session.'],
    ], [2.5, 4.4])
    w.p('Guard từ chối khi journal có version từ 2 trở lên hoặc các bảng payment_orders, wallet_entries, payment_events, payment_audit có dữ liệu. Vì giữ chung named lock, seed/reset không chạy đua với migration. Xu demo trước backfill trở thành opening và không được coi là doanh thu.')
    w.h('Thời gian và giới hạn của kiểm tra dữ liệu', 2)
    w.p('Connection migration v2 đặt timezone UTC. Opening của tài khoản mới ghi `UTC_TIMESTAMP(3)` tường minh, kể cả khi session caller có timezone khác. Chuẩn hóa toàn bộ pool runtime và cách API đọc datetime tiếp tục ở DB03.')
    w.p('Hàm `inspectWalletLedger` chỉ đọc, trả số user, tổng coins/ledger bằng chuỗi, số user thiếu/thừa opening và số user lệch số dư. Nó hỗ trợ đối chiếu DB02; chưa có CLI ledger-audit hoặc API ví được công bố trong checkpoint này.')

    w.page('5 Quy trình vận hành và đối chiếu')
    w.p('Quy trình dưới đây dành cho lần áp dụng lên dữ liệu đang dùng. Công việc tại checkpoint chỉ chạy trên fixture; không có migration nào đã được áp dụng lên database ứng dụng.')
    for text in [
        'Có bản backup và đã thử restore vào database riêng đúng phiên bản mục tiêu, gồm kiểm tra số user và số dư; chưa coi rehearsal trên fixture là chứng nhận backup production.',
        'Dừng toàn bộ writer, seed và phiên bản ứng dụng cũ; chọn đúng DB_NAME, cấu hình tài khoản migration có quyền tạo/đổi bảng và trigger. Không đưa secret vào Git.',
        'Dùng binary/checksum đã review. MySQL local 9.2.0 đã được kiểm thử; container MySQL 8.4 phải được xác minh trước triển khai staging.',
    ]:
        w.bullet(text)
    w.h('Các lệnh đã có trong repository', 2)
    w.code('npm --prefix server run build\nnpm --prefix server run payments:preflight\nnpm --prefix server run migrate:production')
    w.p('Preflight chỉ đọc: exit 0 khi không có số dư bất thường, 2 khi có NULL/âm/vượt trần, 1 khi không hoàn thành kiểm tra. Exit 0 không chứng minh merchant, schema v2 hoặc backup đã sẵn sàng. `migrate:production` thực sự ghi schema/dữ liệu và cho phép allowBalanceMigration; không dùng lệnh này để xem thử.')
    w.h('Các truy vấn chỉ đọc sau migration', 2)
    w.code("SELECT version,status,applied_at\nFROM payment_schema_migrations ORDER BY version;\nSELECT version,last_user_id,completed_batches\nFROM payment_balance_migration;")
    w.p('Cả v1 và v2 cần applied. Dùng truy vấn dưới đây để tìm số user không có đúng một opening hoặc không khớp ledger; kết quả kỳ vọng là 0:')
    w.code("SELECT COUNT(*) AS users_with_mismatch\nFROM users u LEFT JOIN (\n  SELECT user_id,SUM(delta_coins) AS ledger_coins,\n    SUM(entry_type='opening_balance') AS opening_count\n  FROM wallet_entries GROUP BY user_id\n) w ON w.user_id=u.id\nWHERE COALESCE(w.opening_count,0)<>1\n   OR w.ledger_coins IS NULL OR w.ledger_coins<>u.coins;")
    w.p('Đối chiếu thêm tổng coins và ledger với snapshot, xác nhận startup bằng tài khoản DML rồi mới mở writer của code mới. Không chạy ứng dụng cũ song song sau khi chuyển xong; không dùng SQL sửa coins để che sai lệch.', size=10.5)

    w.page('6 Xử lý lỗi và bằng chứng kiểm thử')
    w.table(['Nhóm lỗi', 'Hướng xử lý'], [
        ['PAYMENT_BALANCE_EXPLICIT', 'Cần đường migration tường minh sau backup/restore; không sửa journal để mở startup.'],
        ['PAYMENT_BALANCE_INVALID hoặc COLUMN', 'Rà soát dữ liệu/kiểu cột, xử lý có dấu vết; runner không tự ép về 0 hoặc làm tròn.'],
        ['PAYMENT_FREEZE_MISSING hoặc DRIFT', 'Giữ maintenance; kiểm tra trigger và dữ liệu trước khi sửa có kiểm soát. Không tự gỡ freeze để mở ghi.'],
        ['PAYMENT_OPENING_CONFLICT hoặc LEDGER_ACTIVITY', 'Kiểm tra opening/key hoặc topup có sẵn; không ghi đè ledger để tiếp tục.'],
        ['PAYMENT_LEDGER_MISMATCH', 'Đối chiếu từng user và snapshot; không ép journal thành applied.'],
        ['PAYMENT_SCHEMA_VERSION hoặc DRIFT', 'Kiểm tra binary, checksum, version và schema; không sửa checksum trong journal.'],
        ['PAYMENT_SCHEMA_LOCKED', 'Xác định runner đang giữ lock. Sau gián đoạn, chỉ resume cùng bản triển khai; named lock mất không tự mở freeze.'],
    ], [2.65, 4.25])
    w.p('Tên nhóm trong bảng rút gọn hậu tố chung; mã đầy đủ gồm `PAYMENT_BALANCE_COLUMN`, `PAYMENT_FREEZE_DRIFT`, `PAYMENT_LEDGER_ACTIVITY` và `PAYMENT_SCHEMA_DRIFT`. Khi dừng giữa chừng, chạy lại lệnh migration bằng cùng binary sau khi kiểm tra nguyên nhân. Không reset database, xóa opening/checkpoint hoặc sửa số dư để vượt kiểm tra.', size=10.5)
    w.h('Kiểm chứng đã thực hiện', 2)
    w.table(['Lượt kiểm tra', 'Kết quả và phạm vi'], [
        ['Hồi quy tổng hợp', '117/117 đạt, không lỗi/skip; foundation, migration, writers, security/OTP/password, elections và reader.'],
        ['Kiểm tra startup sau chốt cuối', '3/3 đạt: giữ số dư, chặn user commit muộn và yêu cầu migration tường minh. Hai ca cũ được chạy lại, không cộng thành 120 test độc lập.'],
        ['Backup và restore', 'mysqldump --single-transaction trên fixture v1, restore vào DB ngẫu nhiên khác và giữ số dư. Không phải backup dữ liệu ứng dụng.'],
    ], [2.05, 4.85])
    w.p('Build TypeScript thành công. Các ca còn thử rollback cả batch/checkpoint, resume qua connection khác, snapshot cũ, mất lock, opening mâu thuẫn, zero/trần/ID âm, writer đăng ký/admin và timestamp UTC. Script browser/load được kiểm tra cú pháp; bộ browser/load đầy đủ chưa chạy lại.')
    w.link('Checkpoint DB02 và điều kiện vận hành', GITHUB + '/blob/' + COMMIT + '/docs/payments/DB02_STATUS.md')
    w.link('Mã nguồn triển khai 337c48a', GITHUB + '/commit/' + COMMIT)
    w.p('Log local: `.cache/payment-db02/verification.log` và `final-startup-checks.log`, không đưa vào Git. DB03 tiếp tục UTC của pool/API, index bằng truy vấn thực tế và audit vận hành. Chưa có checkout, adapter VNPAY, service cấp xu hoặc worker đối soát.', size=10)
    return w.save(ROOT/'artifacts/payment-documents/11_DB02_Migration_so_du_va_opening_balance.docx', 6)


def daily():
    title = 'Báo cáo tiến độ tích hợp thanh toán MugenBunko'
    w = Writer(title, 'Báo cáo ngày 28 tháng 10 năm 2026')
    w.p(title, 'Title', 21)
    w.p('Ngày 28 tháng 10 năm 2026 | Bản cập nhật sau DB02', size=10.5)
    w.p('Nhánh phát triển `codex/payment-integration` | Commit code `337c48a`', size=10)
    w.p('Hôm nay hoàn thành **T01, DB01 và phần code/kiểm thử DB02**. Phần DB02 được làm trước buổi tiếp theo theo điều chỉnh lịch làm việc. Code đã đẩy lên GitHub; database ứng dụng chưa chạy migration số dư.')
    w.h('Công việc đã hoàn thành')
    w.h('T01 Chuẩn bị nền tảng', 2)
    w.p('Kiểm kê đường ghi xu, chốt schema/DTO/giới hạn, tạo fixture riêng, công cụ preflight chỉ đọc và checklist VNPAY Sandbox. Bộ tài liệu nền gồm chín kế hoạch kỹ thuật cùng roadmap năm tháng đã được lưu trên nhánh.')
    w.h('DB01 Tích hợp schema và startup', 2)
    w.p('Tích hợp năm bảng nghiệp vụ; thêm journal version/checksum, named lock, khả năng chạy lại sau DDL bị ngắt và startup probe. Production đọc schema bằng tài khoản DML, dừng khi schema thiếu/sai và đóng pool lỗi.')
    w.h('DB02 Migration số dư và writer', 2)
    for text in [
        'Giữ số dư cũ, chuẩn hóa cột xu và tạo một opening/user theo batch. Entry và checkpoint cùng transaction; batch đã commit được giữ khi resume.',
        'Freeze writer bằng trigger và trạng thái applying; vẫn chặn ghi sau khi connection migration mất. Chặn cả writer dùng snapshot cũ và user vừa commit lúc startup.',
        'Đăng ký/admin tạo 0 xu và opening cùng transaction. Bỏ reset/admin mẫu khỏi schema.sql; guard chặn seed/reset trên DB có lịch sử hoặc v2 đã bắt đầu.',
        'Database có user cần lệnh migration tường minh sau backup/restore; mở API không tự backfill. Opening mới ghi UTC, còn chính sách toàn bộ pool/API tiếp tục ở DB03.',
    ]:
        w.bullet(text)
    w.h('Tài liệu bàn giao trong lần cập nhật này', 2)
    w.p('Tài liệu MB PAY 11 mô tả contract v2, thuật toán freeze/resume, writer, runbook, lỗi và bằng chứng kiểm thử. Báo cáo này là bản cập nhật sau DB02; bản báo cáo T01/DB01 đã xuất trước đó được giữ lại.')
    w.p('Phạm vi hoàn thành là code, tài liệu và kiểm thử bằng fixture. Backup/restore trên dữ liệu thật, container MySQL 8.4 và nghiệm thu VNPAY vẫn chưa hoàn tất.', size=10.5)

    w.page('Kết quả kiểm thử và dữ liệu đang dùng')
    w.p('Backend TypeScript build thành công. Lượt hồi quy sau DB02 trên MySQL local 9.2.0 đạt **117/117**, không lỗi và không bỏ qua test:')
    w.table(['Nhóm trong lượt hồi quy', 'Đạt', 'Không đạt'], [
        ['Foundation T01', '14/14', '0'],
        ['Migration và startup DB01', '20/20', '0'],
        ['Migration số dư DB02', '21/21', '0'],
        ['Đăng ký admin và opening', '9/9', '0'],
        ['Security OTP và mật khẩu', '35/35', '0'],
        ['Bầu cử moderator', '13/13', '0'],
        ['Reader collections và chapter versions', '5/5', '0'],
        ['Tổng cộng lượt hồi quy', '117/117', '0'],
    ], [4.5, 1.2, 1.2], center=(1, 2))
    w.p('Sau khi chốt kiểm tra startup DB rỗng có user commit muộn, ba ca liên quan chạy trên binary cuối đều đạt **3/3**. Lượt này chạy lại hai ca đã có và thêm một ca mới; không cộng lặp thành 120 test độc lập.')
    w.p('Các tình huống trọng yếu gồm rollback batch/checkpoint, gián đoạn rồi resume bằng connection khác, freeze sau khi mất khóa, snapshot cũ, số dư 0/trần/ID âm, opening mâu thuẫn và user mới ghi opening UTC. Backup/restore bằng mysqldump đã thử trên fixture v1 vào một DB riêng.')
    w.h('Preflight chỉ đọc cuối')
    w.table(['Chỉ số', 'Database ứng dụng'], [
        ['Số user và tổng xu', '23 user và 26.000 xu, khớp baseline.'],
        ['NULL âm hoặc vượt trần', '0 ở từng nhóm.'],
        ['Default và nullability', 'Vẫn 1000 và nullable; chưa chạy v2 trên DB này.'],
    ], [2.8, 4.1])
    w.p('Review bổ sung trước DB02 chưa phát hiện lỗi cần sửa trong T01/DB01. DB02 được kiểm tra nguồn và kiểm thử theo các lượt nêu trên. Không coi những kết quả này là nghiệm thu 32 ca PAY cho gateway, service và UI.', size=10.5)

    w.page('Bàn giao và công việc tiếp theo')
    w.h('Các checkpoint đã đẩy lên GitHub', 2)
    for short, full, label in [
        ('bde6fa2', 'bde6fa29859a2d3aa7c448e0808cc001ade9337d', 'Bộ kế hoạch kỹ thuật và roadmap năm tháng'),
        ('a18b0c4', 'a18b0c43b6d2f4979d7addb4970f4c147a585987', 'T01 contract fixture preflight và checklist'),
        ('5bf0067', '5bf0067bc9927cfd347940df8c535588275c39cd', 'DB01 migration version và startup probe'),
        ('337c48a', COMMIT, 'DB02 opening balance và guard writer'),
    ]:
        w.link(short, GITHUB + '/commit/' + full, label + '  ')
    w.h('Giới hạn còn lại', 2)
    for text in [
        'MySQL 8.4 trong Compose chưa được kiểm chứng. Kết quả local 9.2.0 không thay thế chạy thử trên staging đúng phiên bản.',
        'Chưa backup/restore hoặc chuyển số dư trên database ứng dụng thật. Khi triển khai cần maintenance, dừng writer cũ và chạy migration tường minh.',
        'Merchant code, hash secret và URL callback public chưa có trong môi trường đã kiểm tra. Chưa đăng ký merchant, gọi gateway hoặc có giao dịch sandbox/live.',
        'Chưa có checkout, service cấp xu đã xác minh, worker đối soát hoặc giao diện thanh toán. Bộ browser/load đầy đủ chưa chạy lại trong DB02.',
    ]:
        w.bullet(text)
    w.h('Mốc tiếp theo DB03', 2)
    w.p('Chốt UTC của pool runtime và các API đọc, kiểm tra index bằng truy vấn thực tế, mở rộng audit vận hành và đối chiếu ledger. Hoàn tất nền dữ liệu trước service cấp xu BE02. Tiếp tục chuẩn bị MySQL 8.4 và merchant Sandbox theo checklist; deadline roadmap vẫn là 28/03/2027.')
    w.h('Tài liệu và bằng chứng', 2)
    w.p('Trạng thái mới nhất: `docs/payments/DB02_STATUS.md`. Đọc cùng `SCHEMA_CONTRACT.md`, inventory writer và tài liệu Word MB PAY 11. Các checkpoint T01/DB01 giữ vai trò lịch sử.')
    w.p('Log local: `.cache/payment-db02/verification.log` và `final-startup-checks.log`, không đưa vào Git. Script tái hiện: test:payments:foundation, test:payments:migration, test:payments:opening, test:security, test:elections và test:upgrades trong server/package.json.', size=10)
    return w.save(ROOT/'artifacts/payment-reports/2026-10-28_Bao_cao_tien_do_thanh_toan_MugenBunko_DB02.docx', 3)


if __name__ == '__main__':
    QA.mkdir(parents=True, exist_ok=True)
    outputs = [technical(), daily()]
    (QA/'manifest.json').write_text(json.dumps(outputs, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(outputs, ensure_ascii=False))
