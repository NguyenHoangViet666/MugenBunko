"""Check deliverable structure and record QA without modifying application code."""
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree

sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
QA=ROOT/'.cache/payment-documents-qa'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def polish():
    for path in sorted(OUT.glob('0[1-9]_*.docx')):
        doc=Document(path); changed=False
        for p in doc.paragraphs:
            if p.style.name.startswith('Heading'):
                for run in p.runs:
                    if '_' in run.text:
                        run.text=run.text.replace('_',' '); changed=True
        if changed:
            doc.save(path); print('Polished '+path.name)


def validate():
    report=[]; all_errors=[]
    files=sorted(OUT.glob('0[1-9]_*.docx'))
    assert len(files)==9, len(files)
    for path in files:
        errors=[]
        with ZipFile(path) as z:
            xml=etree.fromstring(z.read('word/document.xml'))
            styles=etree.fromstring(z.read('word/styles.xml'))
            titles=xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Title"]]',namespaces=NS)
            if len(titles)!=1: errors.append('Missing unique Title paragraph')
            for style in styles.xpath('//w:style[@w:styleId="Title" or starts-with(@w:styleId,"Heading")]',namespaces=NS):
                if style.xpath('./w:pPr/w:pBdr',namespaces=NS): errors.append('Border on title or heading style')
            starts=xml.xpath('//w:bookmarkStart/@w:name',namespaces=NS)
            targets=xml.xpath('//w:hyperlink[not(@r:id)]/@w:anchor',namespaces=NS)
            if any(t not in starts for t in targets): errors.append('Broken internal TOC anchor')
            if len(starts)!=len(set(starts)): errors.append('Duplicate bookmark name')
            text=' '.join(''.join(p.xpath('.//w:t/text()',namespaces=NS))
                          for p in xml.xpath('//w:p',namespaces=NS))
            if any(token in text for token in ['turn0search','turn1view','TODO','Lorem ipsum','```']):
                errors.append('Source marker or placeholder remains')
            if 'chưa được triển khai' not in text: errors.append('Missing implementation status')
            if path.name.startswith('09_'):
                for i in range(1,33):
                    if f'PAY{i:02d}' not in text: errors.append(f'Missing PAY{i:02d}')
                if text.count('Chưa thực hiện')<32: errors.append('Test statuses are incomplete')
            doc=Document(path)
            for p in doc.paragraphs:
                if p.style.name in ['Title','Heading 1','Heading 2','Heading 3']:
                    if re.search(r'[^\w\s]',p.text,re.UNICODE) or '_' in p.text:
                        errors.append('Punctuation in heading '+p.text)
            for t in xml.xpath('//w:tbl',namespaces=NS):
                if not t.xpath('./w:tr[1]/w:trPr/w:tblHeader',namespaces=NS): errors.append('Missing repeated table header')
                if not t.xpath('./w:tblPr/w:tblBorders',namespaces=NS): errors.append('Missing explicit table borders')
            item={'file':path.name,'bytes':path.stat().st_size,
                  'words':len(text.split()),'tables':len(doc.tables),'toc_links':len(targets),'errors':errors}
            report.append(item); all_errors.extend([path.name+': '+e for e in errors])
            print(json.dumps(item,ensure_ascii=False))
    (QA/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    if all_errors: raise SystemExit('\n'.join(all_errors))


if __name__=='__main__':
    if '--polish' in sys.argv: polish()
    else: validate()
