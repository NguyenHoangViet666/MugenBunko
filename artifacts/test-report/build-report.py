"""Build an editable Vietnamese Word QA report from recorded test evidence."""
from pathlib import Path
from collections import Counter
import json
import os
import sys
import zipfile
import xml.etree.ElementTree as ET

sys.path.insert(0, str(Path(os.environ['TEMP']) / 'mugenbunko-test-report-libs'))
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'artifacts/test-report'
EVIDENCE = OUT / 'evidence'
TESTER = 'Nguyễn Hoàng Việt'
NAVY, TEAL, MUTED = '17324D', '087F8C', '627489'

def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))

def text(value):
    if isinstance(value, list):
        return '\n'.join(str(v) for v in value)
    return str(value or '')

def web_cases():
    meta = read_json(EVIDENCE / 'web-browser.json')
    log = (EVIDENCE / 'web-browser.log').read_text(encoding='utf-8-sig')
    assert meta['exit_code'] == 0 and 'Browser smoke checks passed.' in log
    records = [
        ('Khách tiếp tục chương cuối đang đọc', 'Guest; truyện 1, chương 2; cache v2 tại đoạn 15, scrollRatio=0.4 và completedChapterIds=[].', ['Mở chi tiết truyện 1 và nạp cache vị trí chương 2.', 'Tải lại trang, chọn tiếp tục đọc.'], 'Reader mở đúng “Chương 2: Khu vườn sau thư viện”; URL #/novel/1/reader/1, dù chưa đánh dấu chương hoàn tất.', 96),
        ('Đồng bộ đoạn đọc và phục hồi sau khi mất mạng', 'Tài khoản upgrade_reader; hai browser context độc lập; truyện 1 có 60 đoạn/chương.', ['Cuộn đến đoạn 23, chuyển sang Tủ sách và chờ dữ liệu MySQL.', 'Mở Reader trong context thứ hai; so sánh đoạn khôi phục.', 'Giả lập offline, cuộn tới đoạn 36; kiểm tra pending; bật mạng lại.'], 'Sai lệch đoạn giữa hai context không quá 1; cuộn được khôi phục; offline lưu pending và online đồng bộ paragraph_index >= 33.', 109),
        ('Tạo kệ, chia sẻ công khai và thu hồi quyền xem', 'Tài khoản upgrade_reader; truyện 1; tên kệ “Để dành cuối tuần”.', ['Tạo kệ công khai, thêm truyện 1.', 'Dùng context khách mở liên kết chia sẻ.', 'Chuyển kệ sang riêng tư và tải lại liên kết khách.'], 'Kệ và truyện hiển thị cho khách khi công khai; is_public=0 sau khi đổi riêng tư; khách nhận giao diện lỗi truy cập.', 140),
        ('Khôi phục phiên bản chương và giữ nội dung chưa lưu', 'Tác giả sở hữu truyện 1; chương 2 có lịch sử và một bản chỉnh sửa mới.', ['Mở chương 2 trong Studio; nhập “Đoạn đang viết chưa lưu.”.', 'Mở Lịch sử phiên bản, xem bản “Khu vườn sau thư viện”.', 'Khôi phục thành bản nháp; kiểm tra editor và bảng chapter_versions.'], 'Editor nhận nội dung cũ chứa “Đoạn 1.”; phần đang viết chưa lưu vẫn được lưu trong lịch sử phiên bản.', 159),
        ('Diễn đàn báo lỗi mạng và tải lại thành công', 'Có bài “Kiểm tra phục hồi diễn đàn”; khách truy cập #/explore.', ['Chặn /api/forum/posts* rồi tải trang.', 'Kiểm tra thông báo lỗi, bỏ chặn và chọn thử lại.', 'Kiểm tra banner với washi, sepia, charcoal.'], 'Lỗi mạng có thông báo riêng, không bị hiểu là feed rỗng; thử lại phục hồi danh sách; banner có ảnh nền trong cả ba theme.', 177),
        ('Header thích ứng và truy cập web trên màn hình nhỏ', 'Hai trạng thái khách/đăng nhập; viewport rộng 320, 390, 768, 992, 1200, 1440 px.', ['Duyệt các kích thước, đo các nút/link và vùng header.', 'Ở 390 px mở mobile gate và chọn liên kết tiếp tục web.'], 'Không có nút/link bị cắt hoặc các nhóm header chồng nhau; liên kết bỏ qua mobile gate vẫn mở ứng dụng web.', 198),
        ('Mở liên kết trực tiếp và điều hướng bằng bàn phím', 'Khách mở #/about; tài khoản đã đăng nhập mở #/explore.', ['Mở trực tiếp trang Giới thiệu.', 'Focus huy hiệu hồ sơ, nhấn Enter rồi dùng browser Back.'], 'Trang Giới thiệu tải được; Enter mở hồ sơ; Back đưa về giao diện Khám phá.', 216),
        ('Không gian kiểm duyệt, phân quyền và giao diện thích ứng', 'Tài khoản chỉ có quyền moderator; 13 chương pending và 12 báo cáo; theme washi/sepia/charcoal.', ['Duyệt sidebar, tab bằng bàn phím và phân trang.', 'Mô phỏng lỗi xem trước rồi thử lại; duyệt/từ chối/ẩn chương; xử lý báo cáo có bộ lọc.', 'Kiểm tra quyền thao tác tài khoản và kích thước 1440, 992, 768, 390, 320 px.'], 'Các thao tác được kiểm tra cập nhật đúng dữ liệu; không có điều hướng admin hay nút khóa tài khoản cho moderator; các nút trong workspace không bị cắt ở ba theme.', 237),
        ('Quản trị tài khoản, sự kiện và phục hồi lỗi mạng', 'Tài khoản admin; dữ liệu người dùng và đơn tác giả trong DB tạm.', ['Duyệt tab bằng bàn phím, tìm kiếm/lọc/phân trang và duyệt tác giả.', 'Nhập sự kiện, đổi tab; chặn lưu để kiểm tra giữ bản nháp; bỏ chặn và lưu/phát hành/dừng.', 'Thêm thể loại, chọn truyện nổi bật; chặn dashboard rồi thử lại; kiểm tra theme/kích thước và URL trực tiếp.'], 'Tab/bộ lọc và bản nháp được giữ; chu trình sự kiện ghi DB đúng; thể loại và truyện nổi bật được cập nhật; lỗi mạng có thử lại; không cắt điều khiển ở các cấu hình được kiểm tra.', 357),
        ('Đăng ký bằng OTP và đăng nhập tự động trên web', 'Guest; email browser-signup@example.test; username browser_verified_reader; OTP ở chế độ DEV, SMTP tắt.', ['Mở Đăng ký, nhập tên tài khoản, tên hiển thị, mật khẩu hợp lệ và email thử nghiệm.', 'Gửi mã xác thực, dùng OTP dev rồi hoàn tất đăng ký.', 'Đối chiếu phiên localStorage và bản ghi users.'], 'Tự đăng nhập đúng tài khoản; email được lưu; mật khẩu lưu khác mật khẩu nhập. Chưa xác minh việc nhận email thật.', 475),
        ('Cookie HttpOnly và thu hồi phiên khi đăng xuất web', 'Tài khoản vừa đăng ký; cùng phiên cookie trên trình duyệt.', ['Kiểm tra mugen_token và document.cookie.', 'Gọi refresh trước logout; đăng xuất rồi gọi lại refresh.'], 'Không có token lưu ở localStorage; JS không đọc được mugen_auth; refresh trả 200 trước logout và 401 sau logout; localStorage phiên được xóa.', 488),
    ]
    assert len([line for line in log.splitlines() if line.startswith('PASS:')]) == len(records)
    return [dict(id=f'WEB-{i:02d}', group='Kiểm thử giao diện web', title=r[0], preconditions='Edge headless; fixture MySQL riêng; frontend đã build. ' + r[1], testData=r[1], steps=r[2], expected=r[3], actual='PASS; các assertion tương ứng đạt trong lần chạy trình duyệt. ' + r[3], status='PASS', source='tests/browser-upgrades.mjs', sourceLine=r[4], evidence='artifacts/test-report/evidence/web-browser.log', method='Tự động · Edge headless / CDP') for i, r in enumerate(records, 1)]

def normalize(case, category, method):
    case = dict(case)
    case['category'] = category
    case['group'] = case.get('group') or case.get('module') or category
    case['preconditions'] = case.get('preconditions') or case.get('precondition') or ''
    case['testData'] = case.get('testData') or case.get('data') or ''
    case['method'] = case.get('method') or method
    case['priority'] = case.get('priority') or 'Cao'
    return case

backend = read_json(OUT / 'backend-cases.json')
unit = read_json(OUT / 'frontend-mobile-cases.json')
manual = read_json(OUT / 'manual-cases.json')
if isinstance(manual, dict):
    manual = manual['cases']
cases = [normalize(c, 'Node · Backend', 'Tự động · Node / MySQL fixture') for c in backend['cases']]
cases += [normalize(c, 'Node · Web/Mobile', 'Tự động · Node') for c in unit]
web = web_cases()
cases += [normalize(c, 'UI · Web', c['method']) for c in web]
mobile_ui_path = OUT / 'mobile-browser-case.json'
if mobile_ui_path.exists():
    mobile_ui = read_json(mobile_ui_path)
    if isinstance(mobile_ui, dict):
        mobile_ui = [mobile_ui]
    cases += [normalize(c, 'UI · Mobile web', 'Tự động · Expo web / Edge') for c in mobile_ui]
else:
    raise RuntimeError('Awaiting mobile browser result; do not publish incomplete evidence')
cases += [normalize(c, 'Thủ công · Chưa chạy', 'Thủ công · Chưa thực thi') for c in manual]
assert len({c['id'] for c in cases}) == len(cases)
counts = Counter(c['status'] for c in cases)
executed = counts['PASS'] + counts['FAIL']
quality = [json.loads(line) for line in (EVIDENCE / 'quality-checks.jsonl').read_text(encoding='utf-8-sig').splitlines() if line.strip()]
quality.append(dict(name='backend-build', command=backend['builds'][0]['command'], exit_code=backend['builds'][0]['exitCode']))

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.55)
sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
sec.header_distance = sec.footer_distance = Cm(0.7)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(10)
normal.font.color.rgb = RGBColor.from_string('263445')
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 28), ('Heading 1', 17), ('Heading 2', 12), ('Heading 3', 10.5)]:
    style = doc.styles[name]
    style.font.name, style.font.size = 'Calibri', Pt(size)
    style.font.color.rgb = RGBColor.from_string(TEAL if name == 'Heading 2' else NAVY)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(12 if name == 'Heading 1' else 8)
    style.paragraph_format.space_after = Pt(5)

def shade(cell, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shd)

def compact_table(headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        shade(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if ri % 2 == 0:
                shade(cells[i], 'EFF5F8')
    for row in table.rows:
        cant = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant)
        for ci, cell in enumerate(row.cells):
            if widths:
                cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(font_size)
    return table

header = sec.header.paragraphs[0]
header.text = 'MUGENBUNKO  /  BÁO CÁO KIỂM THỬ'
for r in header.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run(TESTER + '  ·  22/10/2026  ·  Trang ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for r in footer.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)

doc.add_paragraph('MUGENBUNKO', 'Subtitle')
doc.add_heading('BÁO CÁO TEST CASE\nVÀ KẾT QUẢ KIỂM THỬ', 0)
doc.add_paragraph('Hệ thống đọc, xuất bản và cộng đồng tiểu thuyết đa nền tảng')
compact_table(['Thông tin', 'Nội dung'], [
    ('Người test', TESTER),
    ('Ngày kiểm thử / lập báo cáo', '22/10/2026 · Asia/Saigon (UTC+7)'),
    ('Mã nguồn được kiểm tra', 'MugenBunko · commit 2d7d95c · thư mục làm việc local'),
    ('Phiên bản tài liệu', '1.0'),
    ('Phạm vi', 'Web React, backend Express/MySQL, nội dung/API mobile, Expo web và danh sách nghiệm thu thủ công.'),
], [4.0, 13.4], 10)
doc.add_paragraph()
doc.add_heading('Kết quả ghi nhận', 1)
doc.add_paragraph(f'{len(cases)} test case: {counts["PASS"]} PASS, {counts["FAIL"]} FAIL, {counts["BLOCKED"]} BLOCKED, {counts["NOT_RUN"]} NOT RUN. {sum(q["exit_code"] == 0 for q in quality)}/{len(quality)} kiểm tra kiểu dữ liệu, lint và build thành công.')
doc.add_paragraph('Kết quả tự động được ghi từ log chạy thực tế trong đợt này. Người test được ghi theo thông tin cung cấp; các ca thủ công chưa thực hiện đều mang trạng thái NOT RUN. Các nhóm luồng trình duyệt được tính riêng với số test của Node.')
doc.add_page_break()

doc.add_heading('1. Tổng quan và môi trường kiểm thử', 1)
groups = ['Node · Backend', 'Node · Web/Mobile', 'UI · Web', 'UI · Mobile web', 'Thủ công · Chưa chạy']
summary_rows = []
for group in groups:
    subset = [c for c in cases if c['category'] == group]
    cc = Counter(c['status'] for c in subset)
    summary_rows.append((group, len(subset), cc['PASS'], cc['FAIL'], cc['BLOCKED'], cc['NOT_RUN']))
summary_rows.append(('Tổng cộng', len(cases), counts['PASS'], counts['FAIL'], counts['BLOCKED'], counts['NOT_RUN']))
compact_table(['Nhóm', 'Tổng', 'PASS', 'FAIL', 'BLOCKED', 'NOT RUN'], summary_rows, [7.0, 1.7, 1.7, 1.7, 2.5, 2.8])
doc.add_paragraph(f'Tỷ lệ PASS trên các ca đã thực thi: {counts["PASS"]}/{executed} ({100*counts["PASS"]/executed:.1f}%). Chỉ số này loại trừ NOT RUN và không chứng minh toàn bộ sản phẩm đã được kiểm thử.')
compact_table(['Thành phần', 'Cấu hình / giới hạn'], [
    ('Máy chạy', 'Windows; Node.js v24.13.0; npm 11.6.2; mã nguồn local.'),
    ('Backend', 'Express/Socket.IO; MySQL local; mỗi fixture dùng DB mugenbunko_test_<16 hex> và cổng loopback ngẫu nhiên; tự dọn sau chạy.'),
    ('Trình duyệt', 'Microsoft Edge headless 155.0.4283.62; browser context/profile tạm; web và Expo web được kiểm tra bằng CDP.'),
    ('Email / OTP', 'SMTP bị tắt trong fixture, DEV_OTP=true. Kết quả OTP dev không thay cho xác nhận gửi/nhận thư thật.'),
    ('Mobile', 'Unit/API và Expo web trên viewport mô phỏng. Chưa chạy native Android/iOS trên thiết bị thật.'),
], [3.4, 14.0])
doc.add_heading('Quy ước kết quả', 2)
doc.add_paragraph('PASS: các assertion đã chạy đạt. FAIL: đã chạy và có kết quả sai mong đợi. BLOCKED: không thể thực thi do điều kiện hoặc môi trường. NOT RUN: ca được soạn để kiểm tra tiếp, chưa có quan sát thực tế.')
doc.add_heading('Kiểm tra chất lượng mã và bản dựng', 2)
compact_table(['Lệnh', 'Kết quả', 'Bằng chứng'], [(q['command'], 'PASS' if q['exit_code'] == 0 else 'FAIL', 'evidence/' + q['name'] + '.log') for q in quality], [7.3, 2.0, 8.1])
doc.add_paragraph('Vite build hoàn tất với cảnh báo có chunk lớn hơn 500 kB; cảnh báo này không làm build thất bại. Đợt báo cáo không bao gồm đo tải production, kiểm tra SMTP thật hay kiểm chứng native device.')
doc.add_page_break()

def case_block(case):
    label = case['status'].replace('_', ' ')
    doc.add_heading(f'{case["id"]}  |  {case["title"]}', 3)
    details = f'{case["group"]}  ·  {case["method"]}  ·  {label}'
    p = doc.add_paragraph(details)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(3)
    for r in p.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(MUTED)
    steps = case.get('steps') or []
    if isinstance(steps, list):
        steps = '\n'.join(f'{i}. {step}' for i, step in enumerate(steps, 1))
    source = text(case.get('source'))
    if case.get('sourceLine'):
        source += ':' + str(case['sourceLine'])
    ev = text(case.get('evidence'))
    ev = ev.replace('artifacts/test-report/', '')
    condition = text(case['preconditions'])
    data = text(case['testData'])
    if data and data not in condition:
        condition += '\nDữ liệu: ' + data
    result = f'{label} — ' + text(case.get('actual'))
    table = compact_table(['Hạng mục', 'Chi tiết'], [
        ('Điều kiện / dữ liệu', condition),
        ('Bước thực hiện', text(steps)),
        ('Kết quả mong đợi', text(case.get('expected'))),
        ('Kết quả thực tế', result),
        ('Nguồn / bằng chứng', source + ('\nLog: ' + ev if ev else '\nChưa có log thực thi cho ca thủ công.')),
    ], [3.0, 14.4], 9)
    # Keep a case together where it fits, while allowing larger cases to flow.
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    for p in table.rows[-1].cells[-1].paragraphs:
        p.paragraph_format.keep_with_next = False
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

section_number = 2
for group in groups:
    group_cases = [c for c in cases if c['category'] == group]
    doc.add_heading(f'{section_number}. {group}', 1)
    if group == 'UI · Web':
        doc.add_paragraph('11 nhóm kịch bản được ánh xạ từ 11 checkpoint PASS của script browser-upgrades.mjs; đây không phải 11 test Node độc lập.')
    if group == 'UI · Mobile web':
        doc.add_paragraph('Kết quả ở môi trường Expo web/Edge; không suy diễn thành kiểm thử native Android hoặc iOS.')
    if group == 'Thủ công · Chưa chạy':
        doc.add_paragraph('Các ca sau được soạn từ chức năng hiện có để nghiệm thu tiếp. Cột kết quả thực tế ghi rõ chưa thực thi; cần bổ sung sau khi chạy trên môi trường phù hợp.')
    for case in group_cases:
        case_block(case)
    section_number += 1
    doc.add_page_break()

doc.add_heading(f'{section_number}. Nhận xét và công việc cần kiểm tra tiếp', 1)
if not counts['FAIL'] and not counts['BLOCKED']:
    doc.add_paragraph(f'Trong phạm vi đã chạy, {counts["PASS"]} ca/nhóm ca đạt; không ghi nhận ca FAIL hoặc BLOCKED. Kết luận chỉ áp dụng cho các assertion và dữ liệu thử đã nêu.')
else:
    doc.add_paragraph(f'Đợt chạy ghi nhận {counts["FAIL"]} FAIL và {counts["BLOCKED"]} BLOCKED; xem cột kết quả thực tế và log của từng ca để xử lý.')
doc.add_paragraph(f'Cần thực hiện {counts["NOT_RUN"]} ca còn lại, ưu tiên gửi/nhận OTP thật, trải nghiệm native Android/iOS, quyền hệ điều hành và các tình huống biên của luồng nghiệp vụ.')
doc.add_paragraph('Ngưỡng độ dài review cần được xác nhận thống nhất giữa web, mobile và backend. Đây là điểm cần nghiệm thu thêm từ rà soát mã nguồn, chưa được ghi thành lỗi đã tái hiện trong đợt này.')
doc.add_heading('Truy xuất và chạy lại', 2)
doc.add_paragraph('Log gốc và metadata nằm tại artifacts/test-report/evidence/. Danh mục đầy đủ, có thể xử lý bằng công cụ, nằm tại all-test-cases.json. Các lệnh sau tái chạy bộ kiểm tra; cần cài dependency và MySQL local với quyền tạo/xóa database thử nghiệm.')
commands = [
    'npm run typecheck', 'npm run lint', 'npm run build', 'npm --prefix server run build', 'npm --prefix mobile run typecheck',
    'node --test tests/readingProgress.test.ts tests/mobile-content.test.ts tests/mobile-dependencies.test.mjs',
    'node --test server/dist/tests/upgrades.test.js',
    'node --test server/dist/tests/security.test.js',
    'node --test server/dist/tests/hardening.test.js',
    'node --test tests/mobile-api.test.mjs',
    'node tests/browser-upgrades.mjs', 'node tests/mobile-browser.mjs',
]
for command in commands:
    p = doc.add_paragraph(command)
    p.paragraph_format.space_after = Pt(2)
    for r in p.runs:
        r.font.name, r.font.size = 'Consolas', Pt(8)
doc.add_paragraph('Ảnh chụp của đợt kiểm thử web: artifacts/upgrades/. Ảnh chụp mobile: xem danh sách ảnh hiện hành trong metadata của ca mobile; không dùng các ảnh cũ để thay kết quả lần chạy này.')
doc.add_paragraph('Người test: ' + TESTER)
doc.add_paragraph('Ngày lập báo cáo: 22/10/2026')

props = doc.core_properties
props.title = 'MugenBunko - Báo cáo test case và kết quả kiểm thử'
props.subject = 'Test case, kết quả thực thi và kế hoạch nghiệm thu'
props.author = TESTER
props.last_modified_by = TESTER
props.language = 'vi-VN'
props.keywords = 'MugenBunko, test case, kiểm thử, Nguyễn Hoàng Việt'
destination = OUT / 'MugenBunko_TestCase_KetQua_NguyenHoangViet.docx'
doc.save(destination)
payload = dict(tester=TESTER, reportDate='2026-10-22', timezone='Asia/Saigon', commit='2d7d95c', summary=dict(counts), qualityChecks=quality, cases=cases)
(OUT / 'all-test-cases.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
with zipfile.ZipFile(destination) as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith('.xml') or name.endswith('.rels'):
            ET.fromstring(archive.read(name))
print(json.dumps(dict(file=str(destination), size=destination.stat().st_size, cases=len(cases), counts=dict(counts), qualityChecks=len(quality)), ensure_ascii=False))
