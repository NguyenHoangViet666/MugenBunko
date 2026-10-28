"""Build the verified daily payment progress report with bundled Python."""
from pathlib import Path
import json
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
FILE = OUT / '2026-10-28_Bao_cao_tien_do_thanh_toan_MugenBunko.docx'
QA = ROOT / '.cache/payment-daily-report-2026-10-28'
FONT = 'Calibri'


def node(tag, **attrs):
    result = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        result.set(qn('w:' + key), str(value))
    return result


def font(run, size=11, bold=False, color='000000', name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    props = run._element.get_or_add_rPr()
    props.append(node('rFonts', ascii=name, hAnsi=name, eastAsia=name, cs=name))
    props.append(node('lang', val='vi-VN'))


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.7), Inches(.7)
section.left_margin, section.right_margin = Inches(.8), Inches(.8)
section.footer_distance = Inches(.3)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet'):
    style = doc.styles[name]
    style.font.name = FONT
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    style._element.get_or_add_rPr().append(node('rFonts', ascii=FONT, hAnsi=FONT, eastAsia=FONT, cs=FONT))
    for border in style._element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.08
    pf.widow_control = True
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
for name, size in [('Title', 21), ('Heading 1', 14), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(10 if name.startswith('Heading') else 0)
    style.paragraph_format.space_after = Pt(6)
doc.styles['List Bullet'].paragraph_format.space_after = Pt(4)


def paragraph(text, style=None, size=11, bold=False):
    p = doc.add_paragraph(style=style)
    for part in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text):
        if not part:
            continue
        if part.startswith('`'):
            font(p.add_run(part[1:-1]), min(size, 10), name='Consolas')
        elif part.startswith('**'):
            font(p.add_run(part[2:-2]), size, True)
        else:
            font(p.add_run(part), size, bold)
    return p


def table(headers, rows, widths, centered=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    for row_index, values in enumerate([headers, *rows]):
        row = t.rows[0] if row_index == 0 else t.add_row()
        row._tr.get_or_add_trPr().append(node('cantSplit'))
        if row_index == 0:
            row._tr.get_or_add_trPr().append(node('tblHeader'))
        for index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = node('tcMar')
            for edge, amount in [('top', 65), ('bottom', 65), ('left', 90), ('right', 90)]:
                margins.append(node(edge, w=amount, type='dxa'))
            props.append(margins)
            borders = node('tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                borders.append(node(edge, val='single', sz=4, color='D9D9D9'))
            props.append(borders)
            props.append(node('shd', fill='344256' if row_index == 0 else ('F1F5F8' if row_index % 2 else 'FFFFFF')))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.06
            if index in centered:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(str(value)), 10.5, row_index == 0, 'FFFFFF' if row_index == 0 else '000000')
    paragraph('', size=4).paragraph_format.space_after = Pt(0)
    return t


def link(p, label, url):
    from docx.opc.constants import RELATIONSHIP_TYPE
    hyperlink = node('hyperlink')
    hyperlink.set(qn('r:id'), p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = node('r')
    props = node('rPr')
    props.append(node('rFonts', ascii=FONT, hAnsi=FONT))
    props.append(node('sz', val=21))
    props.append(node('color', val='244F75'))
    props.append(node('u', val='single'))
    run.append(props)
    text = node('t')
    text.text = label
    run.append(text)
    hyperlink.append(run)
    p._p.append(hyperlink)


paragraph('Báo cáo tiến độ tích hợp thanh toán MugenBunko', 'Title', 21, True)
paragraph('Ngày 28 tháng 10 năm 2026', size=11)
paragraph('Nhánh phát triển `codex/payment-integration`', size=10)
paragraph('Hôm nay đã hoàn thành **T01** và phần triển khai, kiểm thử của **DB01**. '
          'Kế hoạch công việc ban đầu đã đạt; DB01 được thực hiện thêm. '
          'Bộ kiểm thử sau DB01 đạt **69/69 ca**, backend TypeScript build thành công và mã nguồn đã được đẩy lên GitHub.')

paragraph('Công việc đã hoàn thành', 'Heading 1', 14, True)
paragraph('T01 Chuẩn bị nền tảng thanh toán', 'Heading 2', 12, True)
paragraph('Đã kiểm kê các đường ghi xu trực tiếp và gián tiếp qua default, seed, CLI và fixture; '
          'chốt gói xu ban đầu, giới hạn, trạng thái thu tiền và cấp xu, cùng DTO dự kiến. '
          'Schema năm bảng được kiểm tra bằng database ngẫu nhiên riêng và có công cụ preflight chỉ đọc.')
paragraph('Đã lập checklist merchant VNPAY Sandbox và lưu các quyết định kỹ thuật trong `docs/payments/`. '
          'Bộ tài liệu nền gồm chín kế hoạch kỹ thuật và roadmap năm tháng đã được lưu trên nhánh phát triển.')

paragraph('DB01 Tích hợp schema và kiểm tra startup', 'Heading 2', 12, True)
for text in [
    'Tích hợp năm bảng thanh toán vào khởi tạo và migration ứng dụng; thêm journal lưu version, checksum và trạng thái applying hoặc applied.',
    'Dùng khóa theo database để tránh migration thanh toán chạy đồng thời; hỗ trợ chạy lại sau DDL bị ngắt, giữ giá gói đã chỉnh sửa và không tự tạo lại bảng bị mất sau khi đã hoàn tất.',
    'Startup production kiểm tra cột, index, FK, CHECK và version bằng tài khoản DML. Schema thiếu hoặc sai làm khởi động thất bại; connection và pool chưa sẵn sàng được đóng.',
]:
    paragraph(text, 'List Bullet')

paragraph('Kết quả kiểm thử và review', 'Heading 1', 14, True)
table(['Nhóm kiểm thử', 'Đạt', 'Không đạt'], [
    ['Foundation T01', '14/14', '0'],
    ['Migration và startup DB01', '20/20', '0'],
    ['Security OTP và mật khẩu', '35/35', '0'],
    ['Tổng cộng', '69/69', '0'],
], [4.5, 1.2, 1.2], centered=(1, 2))
paragraph('Đợt kiểm thử chạy trên **MySQL local 9.2.0**, không bỏ qua ca nào. '
          'Review bổ sung migration, startup, fixture và tài liệu chưa phát hiện lỗi mới cần sửa trong phạm vi T01 và DB01.', size=10.5)

doc.add_page_break()
paragraph('Phạm vi xác minh và dữ liệu', 'Heading 1', 14, True)
paragraph('Preflight chỉ đọc xác nhận các số liệu tổng hợp dưới đây vẫn khớp baseline T01. '
          'Chưa chạy migration thanh toán trên database ứng dụng đang dùng, chưa thay số dư hoặc backfill opening balance.')
table(['Chỉ số hoặc thuộc tính', 'Kết quả'], [
    ['Số tài khoản', '23'],
    ['Tổng số dư hiện tại', '26.000 xu'],
    ['Số dư NULL hoặc âm hoặc vượt trần', '0 ở từng nhóm'],
    ['Default và nullability của users.coins', '1000 và còn nullable'],
], [4.5, 2.4])
paragraph('MySQL **8.4** là phiên bản trong cấu hình Compose nhưng container này chưa được kiểm chứng. '
          'Kết quả trên local 9.2.0 chưa thay thế việc chạy thử migration và restore trên môi trường staging.')
paragraph('Chưa đăng ký merchant hoặc gọi VNPAY; merchant code, hash secret và URL callback public chưa có trong môi trường đã kiểm tra. '
          'Chưa triển khai checkout, service cấp xu đã xác minh, worker đối soát hoặc giao diện thanh toán. '
          '69 ca đã đạt không phải 32 ca nghiệm thu PAY dự kiến cho service, gateway và UI.')

paragraph('Công việc tiếp theo', 'Heading 1', 14, True)
table(['Mốc', 'Nội dung và điều kiện hoàn thành'], [
    ['DB02', 'Diễn tập backup và restore, đóng băng writer; chuẩn hóa users.coins về INT NOT NULL DEFAULT 0, giữ số dư và tạo một opening balance mỗi user. Chuyển đăng ký, admin và seed sang chính sách ledger; chạy lại không nhân đôi số dư.'],
    ['DB03', 'Đối chiếu ledger theo user và tổng hệ thống, chốt connection UTC, kiểm tra index bằng truy vấn thực tế. Hoàn tất nền dữ liệu trước service cấp xu BE02.'],
], [.75, 6.15])

paragraph('Bàn giao mã nguồn và tài liệu', 'Heading 1', 14, True)
paragraph('Các checkpoint dưới đây đã được đẩy lên nhánh phát triển. HEAD cuối ngày cho phần triển khai là `5bf0067`.', size=10.5)
for short, full, label in [
    ('bde6fa2', 'bde6fa29859a2d3aa7c448e0808cc001ade9337d', 'Bộ kế hoạch kỹ thuật và roadmap năm tháng'),
    ('a18b0c4', 'a18b0c43b6d2f4979d7addb4970f4c147a585987', 'T01 contract schema fixture và preflight'),
    ('5bf0067', '5bf0067bc9927cfd347940df8c535588275c39cd', 'DB01 migration theo version và kiểm tra startup'),
]:
    p = paragraph(label + '  ', size=10.5)
    link(p, short, 'https://github.com/NguyenHoangViet666/MugenBunko/commit/' + full)

paragraph('Tệp kiểm chứng', 'Heading 2', 12, True)
paragraph('Trạng thái công việc: `docs/payments/T01_STATUS.md` và `docs/payments/DB01_STATUS.md`. '
          'Quy tắc schema và migration: `docs/payments/SCHEMA_CONTRACT.md`.', size=10)
paragraph('Log kiểm thử local: `.cache/payment-db01/verification.log` (không đưa vào Git). '
          'Có thể chạy lại bằng các script `test:payments:foundation`, `test:payments:migration` và `test:security` trong `server/package.json`.', size=10)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(footer.add_run('Trang '), 9)
field = node('fldSimple', instr='PAGE')
run = node('r')
text = node('t')
text.text = '1'
run.append(text)
field.append(run)
footer._p.append(field)

props = doc.core_properties
props.title = 'Báo cáo tiến độ tích hợp thanh toán MugenBunko'
props.subject = 'Kết quả công việc ngày 28 tháng 10 năm 2026'
props.author = 'MugenBunko'
props.language = 'vi-VN'
props.keywords = 'MugenBunko, payment, T01, DB01, báo cáo tiến độ'
doc.save(FILE)
QA.mkdir(parents=True, exist_ok=True)
(QA / 'manifest.json').write_text(json.dumps({'file': str(FILE), 'expected_pages': 2}, ensure_ascii=False, indent=2), encoding='utf-8')
print(str(FILE))
